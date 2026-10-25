#!/usr/bin/env python3
"""Publish evidence-backed progress without changing the historical workbook."""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import tempfile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('source_tables_r23', ROOT / 'scripts/export-source-evidence-r21.py')
TABLES = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(TABLES)
CREATOR = 'BOW fidelity progress R23'


def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def local_file(root, value):
    if not isinstance(value, str) or not value or '\\' in value:
        raise ValueError('Invalid evidence path')
    relative = Path(value)
    if relative.is_absolute() or any(part in ('.', '..') for part in relative.parts):
        raise ValueError('Evidence must stay inside the checkout')
    path = (root / relative).resolve(strict=True)
    if not path.is_relative_to(root.resolve()) or not path.is_file():
        raise ValueError('Evidence must be a local file')
    return path


def validate(tracking, root=ROOT):
    if tracking.get('schema') != 're.bow.parallel-workstreams/1' or tracking.get('canonical_fidelity_certified') != 0:
        raise ValueError('Unexpected tracking schema or unsupported fidelity certification')
    baseline = read(root / 'docs/R20_TRACKING.json')
    expected = {task['id'] for task in baseline['tasks'] if task['status'] != 'VALIDÉ'}
    tasks = tracking.get('tasks', [])
    if len(tasks) != len(expected) or {task['id'] for task in tasks} != expected:
        raise ValueError('All historical open workstreams must be present exactly once')
    for task in tasks:
        if task.get('status') != 'EN COURS' or task.get('category') not in ('sources', 'functions', 'validation'):
            raise ValueError('No workstream may be silently closed or omitted')
        if not task.get('advances') or not task.get('remaining') or not task.get('evidence'):
            raise ValueError('Each workstream requires an advance, an open criterion and evidence')
        for item in task['evidence']:
            local_file(root, os.path.normpath(str(Path('modules/fidelity-r23') / item['href'])))
    for resource in tracking['resources']:
        local_file(root, os.path.normpath(str(Path('modules/fidelity-r23') / resource['href'])))
    for evidence in tracking['inputs']:
        path = local_file(root, evidence['path'])
        if digest(path) != evidence['sha256']:
            raise ValueError('Evidence hash mismatch: ' + evidence['path'])
    ids = {item['id'] for item in read(root / 'data/roster.json')['catalog']}
    updates = tracking.get('identity_updates', [])
    if len({item['id'] for item in updates}) != len(updates) or not {item['id'] for item in updates}.issubset(ids):
        raise ValueError('Unknown or duplicate catalogue identity in progress')
    return tracking


def atomic_write(path, content):
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=path.parent, delete=False) as handle:
        temporary = Path(handle.name)
        handle.write(content)
    try:
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def export(tracking_path, baseline_path, output_path, progress_path, root=ROOT):
    tracking = validate(read(tracking_path), root)
    book = load_workbook(baseline_path)
    if book.properties.creator != 'BOW fidelity progress R22':
        raise ValueError('Expected the historical R22 workbook')
    if output_path.exists():
        previous = load_workbook(output_path, read_only=True)
        previous_creator = previous.properties.creator
        previous.close()
        if previous_creator != CREATOR:
            raise ValueError('Refusing to overwrite a workbook from another author')
    book.properties.creator = CREATOR
    book.properties.title = 'B.O.W. Protocol — chantiers parallèles R23'
    TABLES.sheet(book, 'Synthèse R23', ['Indicateur', 'Valeur', 'Portée'],
                 [[item['label'], item['value'], item.get('scope', '')] for item in tracking['statistics']]
                 + [['Fidélité globale 1:1 certifiée', 0, 'Les critères non vérifiés restent ouverts.']])
    TABLES.sheet(book, 'Chantiers R23', ['ID', 'Chantier', 'Statut R23', 'Avancée vérifiée', 'Reste à réaliser', 'Preuves'],
                 [[task['id'], task['title'], task['status'], '\n'.join(task['advances']), '\n'.join(task['remaining']),
                   '\n'.join(item['label'] + ': ' + item['href'] for item in task['evidence'])] for task in tracking['tasks']])
    TABLES.sheet(book, 'Sources R23', ['Chantier', 'UID ou portée', 'Nature', 'Source ou fichier', 'Limite', 'Citation ou observation'], tracking['source_rows'])
    if tracking.get('motion_manifest'):
        clip = read(local_file(root, tracking['motion_manifest']))
        TABLES.sheet(book, 'Images Malacoda R23', ['Image', 'Durée source ms', 'Atlas', 'Recadrage', 'SHA-256 RGBA', 'Lecture'],
                     [[frame['sourceFrame'] + 1, frame['durationMs'], frame['atlas'], ','.join(map(str, frame['crop'])),
                       frame['rgbaSha256'], 'MANUELLE — cadence non définie'] for frame in clip['frames']])
    TABLES.sheet(book, 'Contrôles R23', ['Contrôle', 'Résultat', 'Portée', 'Preuve'],
                 [[item['label'], item['result'], item['scope'], item['path']] for item in tracking['validation']])
    updates = {item['id']: item for item in tracking.get('identity_updates', [])}
    historical = list(book['Catalogue 389'].values)
    TABLES.sheet(book, 'Catalogue R23', list(historical[0]) + ['Ressource visuelle R23', 'Images source R23',
                 'Citations documentaires R23', 'Sons natifs livrés R23', 'Limites R23'],
                 [list(row) + [updates.get(row[0], {}).get(key, '') for key in ('visual', 'motion_frames', 'claims', 'audio', 'limitations')]
                  for row in historical[1:]])
    book.active = book.sheetnames.index('Chantiers R23')
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=output_path.parent, suffix='.xlsx', delete=False) as handle:
        temporary = Path(handle.name)
    try:
        book.save(temporary)
        check = load_workbook(temporary, read_only=True, data_only=True)
        if check['Catalogue R23'].max_row != 390 or check['Chantiers R23'].max_row != 18:
            raise ValueError('Unexpected catalogue or workstream count after export')
        check.close()
        os.replace(temporary, output_path)
    finally:
        temporary.unlink(missing_ok=True)
    progress = {key: tracking[key] for key in ('summary', 'statistics', 'resources', 'tasks')}
    progress['schema'] = 're.bow.fidelity-progress/1'
    atomic_write(progress_path, (json.dumps(progress, ensure_ascii=False, indent=2) + '\n').encode('utf-8'))
    return {'workbook': str(output_path), 'progress': str(progress_path), 'open_workstreams': len(tracking['tasks']),
            'canonical_fidelity_certified': 0}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--tracking', type=Path, default=ROOT / 'docs/R23_TRACKING.json')
    parser.add_argument('--baseline', type=Path, default=ROOT / 'docs/BOW_FIDELITE_R22.xlsx')
    parser.add_argument('--output', type=Path, default=ROOT / 'docs/BOW_FIDELITE_R23.xlsx')
    parser.add_argument('--progress', type=Path, default=ROOT / 'modules/fidelity-r23/progress.json')
    args = parser.parse_args()
    print(json.dumps(export(args.tracking, args.baseline, args.output, args.progress), ensure_ascii=False))


if __name__ == '__main__':
    main()
