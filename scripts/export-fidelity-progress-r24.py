#!/usr/bin/env python3
"""Export R24 evidence while preserving every historical R23 worksheet byte."""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import math
import os
from pathlib import Path
import re
import tempfile
from urllib.parse import unquote, urlsplit
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('source_tables_r24', ROOT / 'scripts/export-source-evidence-r21.py')
TABLES = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(TABLES)
CREATOR = 'BOW fidelity progress R24'
BASE_CREATOR = 'BOW fidelity progress R23'
ADDED_SHEETS = ('Synthèse R24', 'Chantiers R24', 'Catalogue R24', 'Sources R24', 'Recette R24', 'Clips sources R24')
UPDATE_KEYS = {'id', 'visual', 'rig', 'motion_clips', 'audio', 'claims', 'limitations'}
POSE_LABEL = 'POSE FIXE APRÈS ZÉRO'
MOTION_LABEL = 'MOUVEMENT APRÈS ZÉRO'


def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def checked_path(root, value, exists=False):
    root = Path(root).resolve()
    candidate = Path(value)
    candidate = candidate if candidate.is_absolute() else root / candidate
    try:
        path = candidate.resolve(strict=False)
    except (FileNotFoundError, OSError, RuntimeError) as error:
        raise ValueError('Evidence must be an existing local file') from error
    if not path.is_relative_to(root) or path == root:
        raise ValueError('Paths must stay inside the checkout')
    if exists and not path.is_file():
        raise ValueError('Evidence must be a local file')
    return path


def local_file(root, value):
    if not isinstance(value, str) or not value or '\\' in value or Path(value).is_absolute():
        raise ValueError('Invalid evidence path inside the checkout')
    if any(part in ('.', '..') for part in Path(value).parts):
        raise ValueError('Evidence path must stay inside the checkout')
    return checked_path(root, value, exists=True)


def evidence_link(root, href):
    if not isinstance(href, str) or not href or '\\' in href:
        raise ValueError('Invalid evidence link')
    url = urlsplit(href)
    if url.scheme or url.netloc or not url.path or url.path.startswith('/'):
        raise ValueError('Evidence links must stay inside the checkout')
    return checked_path(root, Path('modules/fidelity-r24') / unquote(url.path), exists=True)


def safe_cell(value):
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, str):
        if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value):
            raise ValueError('Unsafe spreadsheet control character')
        if value.lstrip(' \t\r\n').startswith(('=', '+', '-', '@')):
            return "'" + value
    elif value is not None and not isinstance(value, (bool, int, float)):
        raise ValueError('Unsupported spreadsheet cell value')
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError('Non-finite spreadsheet number')
    return value


def sheet(book, name, headers, rows):
    return TABLES.sheet(book, name, [safe_cell(value) for value in headers],
                        [[safe_cell(value) for value in row] for row in rows])


def validate(tracking, root=ROOT):
    if tracking.get('schema') != 're.bow.parallel-workstreams/2' or type(tracking.get('canonical_fidelity_certified')) is not int or tracking['canonical_fidelity_certified'] != 0:
        raise ValueError('Unexpected tracking schema or unsupported global fidelity certification')
    expected = {task['id'] for task in read(root / 'docs/R20_TRACKING.json')['tasks'] if task['status'] != 'VALIDÉ'}
    tasks = tracking.get('tasks', [])
    if len(expected) != 17 or len(tasks) != 17 or {task['id'] for task in tasks} != expected:
        raise ValueError('All seventeen historical open workstreams must be present exactly once')
    for task in tasks:
        if task.get('status') != 'EN COURS' or task.get('category') not in ('sources', 'functions', 'validation'):
            raise ValueError('No workstream may be silently closed or omitted')
        if not task.get('title') or not task.get('advances') or not task.get('remaining') or not task.get('evidence'):
            raise ValueError('Each workstream requires an advance, an open criterion and evidence')
        if any(not isinstance(item, str) or not item for item in task['advances'] + task['remaining']):
            raise ValueError('Workstream advances and criteria must be text')
        for item in task['evidence']:
            evidence_link(root, item['href'])
    for resource in tracking.get('resources', []):
        evidence_link(root, resource['href'])
    inputs = tracking.get('inputs', [])
    if not inputs or len({item['path'] for item in inputs}) != len(inputs):
        raise ValueError('Evidence inputs must be present and unique')
    for evidence in inputs:
        path = local_file(root, evidence['path'])
        if not re.fullmatch(r'[0-9a-f]{64}', str(evidence.get('sha256', ''))) or digest(path) != evidence['sha256']:
            raise ValueError('Evidence hash mismatch: ' + evidence['path'])
    roster = read(root / 'data/roster.json')['catalog']
    ids = {item['id'] for item in roster}
    updates = tracking.get('identity_updates', [])
    if len(roster) != 389 or len(ids) != 389 or len(updates) != 389 or {item['id'] for item in updates} != ids:
        raise ValueError('Exactly all 389 catalogue identities require progress; unknown or duplicate identity')
    if any(set(item) - UPDATE_KEYS for item in updates):
        raise ValueError('Unknown identity progress field')
    for row in tracking.get('source_rows', []):
        if not isinstance(row, list) or len(row) != 6:
            raise ValueError('Source evidence rows require exactly six columns')
        for value in row:
            safe_cell(value)
    for item in tracking.get('validation', []):
        if not all(key in item for key in ('label', 'result', 'scope', 'path')):
            raise ValueError('Validation requires its result, scope and local evidence')
        local_file(root, item['path'])
    return tracking


def motion_rows(tracking, root):
    if not tracking.get('motion_manifest'):
        return None
    path = local_file(root, tracking['motion_manifest'])
    if not any(local_file(root, item['path']) == path for item in tracking['inputs']):
        raise ValueError('Motion manifest requires a recorded evidence hash')
    manifest = read(path)
    if manifest.get('schema') != 're.bow.rigged-motion-pack/1' or manifest.get('summary', {}).get('canonicalFidelityCertified') != 0:
        raise ValueError('Unexpected motion manifest or unsupported motion fidelity certification')
    catalogue = {item['id'] for item in read(root / 'data/roster.json')['catalog']}
    rows, clip_ids = [], set()
    motion_count, pose_count = 0, 0
    for entry in manifest['entries']:
        if entry['id'] not in catalogue or entry.get('canonicalFidelityCertified') is not False or entry.get('rigProductionFinalized') is not False:
            raise ValueError('Unknown motion identity or unsupported finalized rig')
        for clip in entry['clips']:
            key = (entry['id'], clip['id'])
            if key in clip_ids:
                raise ValueError('Duplicate source animation')
            clip_ids.add(key)
            try:
                start, end, ticks = int(clip['keyTimeStart']), int(clip['keyTimeEnd']), int(clip['ticksPerSecond'])
                if ticks <= 0 or start > end:
                    raise ValueError('Invalid native animation clock')
                source_span, preroll, last = (end - start) / ticks, max(0, -start / ticks), end / ticks
                for field, expected in (('sourceTimeSpanSeconds', source_span), ('nativePreRollSeconds', preroll), ('lastNativeTimeSeconds', last)):
                    if not math.isfinite(float(clip[field])) or not math.isclose(float(clip[field]), expected, rel_tol=1e-9, abs_tol=1e-9):
                        raise ValueError('Native animation times disagree: ' + field)
            except (KeyError, TypeError, OverflowError) as error:
                raise ValueError('Missing native animation clock evidence') from error
            source_varying = clip.get('varyingAnimationCurveCount')
            nonnegative_varying = clip.get('varyingNonnegativeAnimationCurveCount')
            pose_only = clip.get('positiveTimePoseOnly')
            if (type(source_varying) is not int or type(nonnegative_varying) is not int
                    or not 0 <= nonnegative_varying <= source_varying or type(pose_only) is not bool
                    or pose_only != (nonnegative_varying == 0)):
                raise ValueError('Inconsistent positive-time pose/motion classification')
            expected_scope = 'pose-only-after-time-zero' if pose_only else 'nonnegative-time-motion'
            if clip.get('playbackVariationScope') != expected_scope:
                raise ValueError('Playback variation scope disagrees with the positive-time curve count')
            motion_count += not pose_only
            pose_count += pose_only
            rows.append([entry['id'], clip['id'], entry.get('game', ''), clip.get('member', ''),
                         clip.get('path') or 'MEMBRE DE L’ARCHIVE SOURCE', clip['sha256'],
                         start / ticks, last, source_span, preroll, max(0, last),
                         source_varying, nonnegative_varying, POSE_LABEL if pose_only else MOTION_LABEL,
                         clip['playbackVariationScope'],
                         'ACTION NON ATTRIBUÉE' if clip.get('semanticActionVerified') is not True else 'ACTION DOCUMENTÉE',
                         ('Une pose après zéro reste une pose même si le pré-roll varie. ' if pose_only else '')
                         + 'Étendue min→max distincte de la lecture 0→max ; pas de rig ou combat final certifié.'])
    if manifest['summary'].get('sourceAnimationFiles') != len(rows):
        raise ValueError('Source animation count differs from the motion manifest')
    if (manifest['summary'].get('positiveTimeMotionFiles') != motion_count
            or manifest['summary'].get('positiveTimePoseFiles') != pose_count):
        raise ValueError('Positive-time motion/pose totals differ from the clip classifications')
    return rows


def preserve_worksheet_xml(baseline, generated):
    """Retain original worksheet XML and relationships, including literal cells."""
    with ZipFile(baseline) as source, ZipFile(generated) as result:
        historical = {name: (source.getinfo(name), source.read(name)) for name in source.namelist() if name.startswith('xl/worksheets/')}
        missing = set(historical) - set(result.namelist())
        if missing:
            raise ValueError('A historical worksheet part was lost')
        with tempfile.NamedTemporaryFile(dir=generated.parent, suffix='.xlsx', delete=False) as handle:
            replacement = Path(handle.name)
        try:
            with ZipFile(replacement, 'w') as target:
                for item in result.infolist():
                    info, content = historical.get(item.filename, (item, result.read(item.filename)))
                    target.writestr(info, content)
            os.replace(replacement, generated)
        finally:
            replacement.unlink(missing_ok=True)


def worksheet_state(page):
    return (page.title, page.freeze_panes, tuple(str(item) for item in page.merged_cells.ranges),
            tuple(tuple((cell.value, cell.data_type, tuple(cell._style) if cell.has_style else None,
                         cell.hyperlink.target if cell.hyperlink else None,
                         (cell.comment.text, cell.comment.author) if cell.comment else None)
                        for cell in row) for row in page.iter_rows()))


def atomic_write(path, content):
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=path.parent, delete=False) as handle:
        temporary = Path(handle.name)
        handle.write(content)
    try:
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def export(tracking_path, baseline_path, output_path, progress_path, root=ROOT):
    root = Path(root).resolve()
    tracking_path, baseline_path = (checked_path(root, path, exists=True) for path in (tracking_path, baseline_path))
    output_path, progress_path = (checked_path(root, path) for path in (output_path, progress_path))
    if len({tracking_path, baseline_path, output_path, progress_path}) != 4:
        raise ValueError('Output must not overwrite historical inputs')
    tracking = validate(read(tracking_path), root)
    if not any(local_file(root, item['path']) == baseline_path for item in tracking['inputs']):
        raise ValueError('Historical workbook requires a recorded evidence hash')
    clips = motion_rows(tracking, root)
    book = load_workbook(baseline_path)
    if book.properties.creator != BASE_CREATOR:
        raise ValueError('Expected the historical R23 workbook')
    if any(name in book.sheetnames for name in ADDED_SHEETS):
        raise ValueError('The baseline already contains R24 sheets')
    if output_path.exists():
        previous = load_workbook(output_path, read_only=True)
        creator = previous.properties.creator
        previous.close()
        if creator != CREATOR:
            raise ValueError('Refusing to overwrite a workbook from another author')
    original_names = list(book.sheetnames)
    original_states = {name: worksheet_state(book[name]) for name in original_names}
    rows = list(book['Catalogue R23'].values)
    roster_ids = {item['id'] for item in read(root / 'data/roster.json')['catalog']}
    if len(rows) != 390 or {row[0] for row in rows[1:]} != roster_ids:
        raise ValueError('Historical catalogue does not match all 389 identities')
    book.properties.creator = CREATOR
    book.properties.title = 'B.O.W. Protocol — chantiers parallèles R24'
    sheet(book, 'Synthèse R24', ['Indicateur', 'Valeur', 'Portée'],
          [[item['label'], item['value'], item.get('scope', '')] for item in tracking['statistics']]
          + [['Fidélité globale 1:1 certifiée', 0, 'Les critères non vérifiés restent ouverts.']])
    sheet(book, 'Chantiers R24', ['ID', 'Chantier', 'Statut R24', 'Avancée vérifiée', 'Reste à réaliser', 'Preuves'],
          [[task['id'], task['title'], task['status'], '\n'.join(task['advances']), '\n'.join(task['remaining']),
            '\n'.join(item['label'] + ': ' + item['href'] for item in task['evidence'])] for task in tracking['tasks']])
    updates = {item['id']: item for item in tracking['identity_updates']}
    sheet(book, 'Catalogue R24', list(rows[0]) + ['Ressources visuelles R24', 'Armature source R24', 'Clips sources R24',
          'Sons natifs livrés R24', 'Citations documentaires R24', 'Limites R24'],
          [list(row) + [updates[row[0]].get(key, '') for key in ('visual', 'rig', 'motion_clips', 'audio', 'claims', 'limitations')] for row in rows[1:]])
    sheet(book, 'Sources R24', ['Chantier', 'UID ou portée', 'Nature', 'Source ou fichier', 'Limite', 'Citation ou observation'], tracking['source_rows'])
    sheet(book, 'Recette R24', ['Contrôle', 'Résultat', 'Portée', 'Preuve'],
          [[item['label'], item['result'], item['scope'], item['path']] for item in tracking['validation']])
    if clips is not None:
        sheet(book, 'Clips sources R24', ['UID', 'Clip source', 'Jeu', 'Membre archive', 'Fichier direct', 'SHA-256',
              'Début natif s', 'Dernier temps natif s', 'Étendue source min→max s', 'Pré-roll natif avant zéro s',
              'Lecture 0→max s', 'Courbes variables source, pré-roll inclus', 'Courbes variables à t ≥ 0',
              'Classification lecture t ≥ 0', 'Portée de la variation', 'Attribution du geste', 'Limites'], clips)
    book.active = book.sheetnames.index('Chantiers R24')
    progress = {key: tracking[key] for key in ('summary', 'statistics', 'resources', 'tasks')}
    progress['schema'] = 're.bow.fidelity-progress/1'
    progress_bytes = (json.dumps(progress, ensure_ascii=False, indent=2) + '\n').encode('utf-8')
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=output_path.parent, suffix='.xlsx', delete=False) as handle:
        temporary = Path(handle.name)
    try:
        book.save(temporary)
        book.close()
        preserve_worksheet_xml(baseline_path, temporary)
        check = load_workbook(temporary, data_only=False)
        try:
            if check.sheetnames[:len(original_names)] != original_names or any(worksheet_state(check[name]) != original_states[name] for name in original_names):
                raise ValueError('A historical worksheet changed during export')
            if check['Catalogue R24'].max_row != 390 or check['Chantiers R24'].max_row != 18:
                raise ValueError('Unexpected catalogue or workstream count after export')
            for name in check.sheetnames[len(original_names):]:
                if any(cell.data_type == 'f' for row in check[name] for cell in row):
                    raise ValueError('Spreadsheet formula injection rejected')
        finally:
            check.close()
        os.replace(temporary, output_path)
        atomic_write(progress_path, progress_bytes)
    finally:
        temporary.unlink(missing_ok=True)
    return {'workbook': str(output_path), 'progress': str(progress_path), 'open_workstreams': 17,
            'catalogue_identities': 389, 'source_clips': len(clips) if clips is not None else 0,
            'positive_time_motion_files': sum(row[13] == MOTION_LABEL for row in clips) if clips is not None else 0,
            'positive_time_pose_files': sum(row[13] == POSE_LABEL for row in clips) if clips is not None else 0,
            'historical_worksheets_preserved': len(original_names), 'canonical_fidelity_certified': 0}


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--tracking', type=Path, default=ROOT / 'docs/R24_TRACKING.json')
    ap.add_argument('--baseline', type=Path, default=ROOT / 'docs/BOW_FIDELITE_R23.xlsx')
    ap.add_argument('--output', type=Path, default=ROOT / 'docs/BOW_FIDELITE_R24.xlsx')
    ap.add_argument('--progress', type=Path, default=ROOT / 'modules/fidelity-r24/progress.json')
    args = ap.parse_args()
    print(json.dumps(export(args.tracking, args.baseline, args.output, args.progress), ensure_ascii=False))


if __name__ == '__main__':
    main()
