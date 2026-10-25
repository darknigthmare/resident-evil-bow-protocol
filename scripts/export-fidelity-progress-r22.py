#!/usr/bin/env python3
"""Export the actual R22 source coverage and keep unfinished production explicit."""
import argparse
import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import os

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('source_evidence_tables', ROOT / 'scripts/export-source-evidence-r21.py')
tables = importlib.util.module_from_spec(spec)
spec.loader.exec_module(tables)
CREATOR = 'BOW fidelity progress R22'


def read(path):
    return json.loads(Path(path).read_text())


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--references', type=Path, default=ROOT / 'qa-source-references-r21/manifest.json')
    parser.add_argument('--reviewed', type=Path, default=ROOT / 'qa-source-references-r22/manifest.json')
    parser.add_argument('--output', type=Path, default=ROOT / 'docs/BOW_FIDELITE_R22.xlsx')
    args = parser.parse_args()
    roster = read(ROOT / 'data/roster.json')
    expansion = read(ROOT / 'data/catalog-expansion-v070.json')
    tracking = read(ROOT / 'docs/R20_TRACKING.json')
    references, reviewed = read(args.references), read(args.reviewed)
    if references['schema'] != 're.bow.source-reference-pack/1' or reviewed['schema'] != 're.bow.reviewed-source-reference-pack/1':
        raise ValueError('Unexpected reference pack schema')
    original = {row['id']: row for row in references['entries']}
    named = {row['id']: row for row in reviewed['entries']}
    if set(original) & set(named):
        raise ValueError('The two input packs unexpectedly overlap')
    combined = {**original, **named}
    module_root = ROOT / 'modules/source-motion-r22'
    live = read(module_root / 'native-portraits.json')
    delivered = {row['id']: row for row in live['entries']}
    if set(delivered) != set(named):
        raise ValueError('Delivered dossier images do not match the reviewed input')
    for uid, item in delivered.items():
        path = (module_root / item['image']).resolve(strict=True)
        if not path.is_relative_to(module_root.resolve()) or digest(path) != item['sha256'] or item['sha256'] != named[uid]['sha256']:
            raise ValueError('Delivered source pixels differ from the reviewed input: ' + uid)
    clip = read(module_root / 'media/nemesis-infection/manifest.json')
    if clip['frameCount'] != 150 or clip['totalDurationMs'] != 6000 or sum(row['durationMs'] for row in clip['frames']) != 6000:
        raise ValueError('Native motion count or durations changed')
    if digest(module_root / 'media/nemesis-infection' / clip['source']['path']) != clip['source']['sha256']:
        raise ValueError('Original motion source changed')
    for atlas in clip['atlases']:
        if digest(module_root / 'media/nemesis-infection' / atlas['path']) != atlas['sha256']:
            raise ValueError('Native atlas bytes changed')
    main_ids = {item['uid'] for item in roster['fighters']}
    contexts = {item['id'] for item in expansion['records'] if item['classification'] == 'contextual'}
    book = Workbook()
    book.remove(book.active)
    book.properties.creator = CREATOR
    summary = [
        ['Portée', 'Références, corrections des Dossiers et lecture documentaire', 'Ce suivi ne déclare pas tous les manquants réalisés.'],
        ['Identités du catalogue', len(roster['catalog']), '115 principales, 231 expérimentales, 43 contextes.'],
        ['Identités avec fichier lié à sa source', len(combined), 'La présence du fichier ne certifie pas une géométrie ou une animation finale.'],
        ['Références nommées revues', len(named), '12 images de corps, 5 portraits, 3 scènes ; corps complet non inféré.'],
        ['Dossiers avec image source intégrée', len(delivered), 'Remplacement de présentation uniquement ; données de combat inchangées.'],
        ['Vue d’origine manquante', 're7-jack-baker-scissor', 'Le composite connu est masqué dans le Dossier, sans substitution humaine.'],
        ['Séquence de Nemesis', 150, '150 images originales de 40 ms, six secondes ; scène muette avec décor.'],
        ['Images de mouvement inventées', 0, 'Aucune interpolation, aucun miroir ou découpage anatomique non sourcé.'],
        ['Rigs et clips de combat finalisés', 0, 'La scène documentaire ne clôture pas le rig ou les gestes du catalogue.'],
        ['Fidélité globale 1:1 certifiée', 0, 'Mesures, phases, vues opposées, mouvements et sons restent à vérifier.'],
        ['Travaux ouverts du suivi R20', 17, 'Les deux productions ouvertes et les quinze lots globaux restent ouverts.'],
    ]
    tables.sheet(book, 'Avancement R22', ['Indicateur', 'Valeur', 'Portée'], summary)
    rows = []
    for identity in roster['catalog']:
        uid = identity['id']
        record, update = combined.get(uid), delivered.get(uid)
        kind = 'Contexte' if uid in contexts else 'Principal' if uid in main_ids else 'Expérimental'
        rows.append([uid, identity['name'], identity['game'], kind,
                     'FICHIER LIÉ À SA SOURCE' if record else 'SOURCE À TROUVER',
                     record['source_page']['url'] if record else identity.get('sourcePage', ''),
                     record['source_image']['url'] if record else '', record['sha256'] if record else '',
                     update['view_scope'] if update else 'Référence non revue individuellement' if record else '',
                     'IMAGE SOURCE INTÉGRÉE' if update else 'COMPOSITE MASQUÉ — VUE MANQUANTE' if uid == 're7-jack-baker-scissor' else '',
                     'SCÈNE ORIGINALE — 150 IMAGES' if uid == 're3r-nemesis' else '',
                     'À PRODUIRE', 'À PRODUIRE', 'NON CERTIFIÉ'])
    tables.sheet(book, 'Catalogue 389', ['UID', 'Identité', 'Jeu', 'Parcours', 'Fichier source', 'Page',
                                       'Image source', 'SHA-256', 'Cadrage', 'Dossier R22', 'Mouvement documentaire',
                                       'Rig final', 'Clips de combat', 'Fidélité complète'], rows)
    tables.sheet(book, 'Vues intégrées', ['UID', 'Jeu', 'Forme revue', 'Cadrage', 'Fichier livré', 'SHA-256', 'Limites'],
                 [[uid, item['source_game'], item['curation']['specific_form'], item['view_scope'], item['image'],
                   item['sha256'], '\n'.join(item['visible_review']['limitations'])] for uid, item in delivered.items()])
    tables.sheet(book, 'Images source Nemesis', ['Image', 'Début ms', 'Durée ms', 'Atlas', 'Recadrage', 'SHA-256 RGBA'],
                 [[frame['sourceFrame'] + 1, sum(x['durationMs'] for x in clip['frames'][:index]),
                   frame['durationMs'], frame['atlas'], ','.join(map(str, frame['crop'])), frame['rgbaSha256']]
                  for index, frame in enumerate(clip['frames'])])
    tables.sheet(book, 'Travaux ouverts', ['ID', 'Chantier', 'Statut historique', 'Reste à produire'],
                 [[item['id'], item['title'], item['status'], item['remaining']]
                  for item in tracking['tasks'] if item['status'] != 'VALIDÉ'])
    output = args.output
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        previous = load_workbook(output, read_only=True)
        creator = previous.properties.creator
        previous.close()
        if creator != CREATOR:
            raise ValueError('Refusing to overwrite a workbook not generated by this exporter')
    with tempfile.NamedTemporaryFile(dir=output.parent, suffix='.xlsx', delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        book.save(temporary_path)
        check = load_workbook(temporary_path, read_only=True, data_only=True)
        assert check['Catalogue 389'].max_row == 390
        assert check['Vues intégrées'].max_row == len(delivered) + 1
        assert check['Images source Nemesis'].max_row == 151
        assert check['Travaux ouverts'].max_row == 18
        check.close()
        os.replace(temporary_path, output)
    finally:
        temporary_path.unlink(missing_ok=True)
    print(json.dumps({'output': str(output), 'source_bound_identities': len(combined),
                      'delivered_dossier_images': len(delivered), 'source_frames': 150,
                      'unfinished_tasks': 17, 'canonical_fidelity_certified': 0}, ensure_ascii=False))


if __name__ == '__main__':
    main()
