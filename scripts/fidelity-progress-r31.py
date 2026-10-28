#!/usr/bin/env python3
"""Bind separately owned R31 workstreams and append proof sheets to frozen R30."""
import argparse
from collections import Counter
import copy
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
from urllib.parse import parse_qs, urlencode, urlsplit
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MODULE = Path('modules/fidelity-r31')
BASELINE = 'docs/BOW_FIDELITE_R30.xlsx'
PRIOR = 'docs/R30_TRACKING.json'
REPORT_INDEX = 'docs/R31_REPORT_INDEX.json'
SHEETS = ('Synthèse R31', 'Chantiers R31', 'Catalogue R31', 'Lots R31', 'Ajouts R31', 'Recette R31', 'Sources R31')
FIELDS = ('r31_lots', 'r31_workstreams', 'r31_summary', 'r31_artifacts', 'r31_source_editions', 'r31_has_additions', 'r31_limits')
LOTS = (
    ('sprites', 'r31_sprites_generation', ('R20-10', 'G-RENDU'), 'sprites/delivery.json'),
    ('motion', 'r31_motion_commands', ('R20-09', 'G-COMMANDES', 'G-CINEMATIQUES', 'G-MANCHES'), 'motion/delivery.json'),
    ('canonical', 'r31_canonical_scale', ('G-PERIMETRE', 'G-PREUVES', 'G-ECHELLE', 'G-VICTIMES'), 'canonical/delivery.json'),
    ('persistence', 'r31_replays_saves', ('G-REPLAYS', 'G-SAUVES'), 'persistence/delivery.json'),
    ('audio', 'r31_audio_sources', ('G-AUDIO',), 'audio/delivery.json'),
    ('validation', 'root', ('G-PERF', 'G-ACCESS'), 'validation/delivery.json'),
)
TRACKING_TASKS = ('G-LIVRAISON', 'G-RECETTE')
SPEC = importlib.util.spec_from_file_location('r31_frozen_helpers', ROOT / 'scripts/fidelity-progress-r30.py')
PREVIOUS = importlib.util.module_from_spec(SPEC); SPEC.loader.exec_module(PREVIOUS)
UTIL = PREVIOUS.UTIL
GENERATED = {'docs/R31_TRACKING.json', 'docs/R31_VALIDATION.json', 'docs/R31_WORKSTREAMS.json',
             'docs/R31_LOTS.json', 'docs/FIDELITE_R31.md', 'docs/BOW_FIDELITE_R31.xlsx',
             (MODULE / 'progress.json').as_posix(), (MODULE / 'dossier-additions.json').as_posix()}


def read(name, root=ROOT):
    return json.loads((Path(root) / name).read_text())


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def raw(value):
    return (json.dumps(value, ensure_ascii=False, indent=2) + '\n').encode()


def write(path, value):
    UTIL.atomic_write(Path(path), value if isinstance(value, bytes) else raw(value))


def local(name, root=ROOT, generated=False):
    if (not isinstance(name, str) or not name or '\\' in name or name.startswith('/')
            or any(part in ('', '.', '..') for part in name.split('/'))
            or re.match(r'^[a-zA-Z][a-zA-Z0-9+.-]*:', name)):
        raise ValueError('Evidence must be a local checkout path')
    path = (Path(root) / name).resolve()
    if not path.is_relative_to(Path(root).resolve()) or not path.is_file() and not (generated and name in GENERATED):
        raise ValueError('Missing or external evidence: ' + name)
    return path


def binding(name, root=ROOT):
    path = local(name, root)
    return {'path': name, 'bytes': path.stat().st_size, 'sha256': sha(path)}


def checked(record, root=ROOT):
    path = local(record['path'], root)
    if (not re.fullmatch('[a-f0-9]{64}', str(record.get('sha256', '')))
            or sha(path) != record['sha256'] or 'bytes' in record and path.stat().st_size != record['bytes']):
        raise ValueError('Evidence hash or size mismatch: ' + record['path'])
    return path


def text_list(values, label):
    if not isinstance(values, list) or not values or len(values) > 256 or any(not isinstance(value, str) or not value.strip() or len(value) > 10000 for value in values):
        raise ValueError(label + ' must contain scoped, nonempty observations')
    return list(values)


def href(name):
    return os.path.relpath(name, MODULE).replace(os.sep, '/')


def collect_lots(root=ROOT, allow_pending=False):
    known = {row['id'] for row in read('data/roster.json', root)['catalog']}
    deliveries, inputs, pending = [], [], []
    for lot_id, owner, task_ids, tail in LOTS:
        name = (MODULE / tail).as_posix()
        shell = {'id': lot_id, 'owner': owner, 'taskIds': list(task_ids), 'path': name}
        if not (Path(root) / name).is_file():
            pending.append(name)
            deliveries.append({**shell, 'state': 'EN ATTENTE', 'summary': 'Lot attribué ; livraison vérifiable encore attendue.',
                'advances': ['Le lot a un responsable explicite.'], 'remaining': ['Livrer les fichiers et leurs preuves propres au chantier.'],
                'manifests': [], 'resources': [], 'identities': [], 'counts': {}})
            continue
        value = read(name, root)
        if (value.get('schema') != 're.bow.fidelity-lot-r31/1' or value.get('release') != 'R31'
                or value.get('id') != lot_id or set(value.get('taskIds', [])) != set(task_ids)
                or len(value.get('taskIds', [])) != len(task_ids) or value.get('canonicalFidelityCertified') is not False
                or not isinstance(value.get('summary'), str) or not value['summary'].strip()):
            raise ValueError('Invalid R31 lot identity, owner scope or certification: ' + name)
        advances = text_list(value.get('advances'), 'Lot advances')
        remaining = text_list(value.get('remaining'), 'Lot remaining criteria')
        task_updates = value.get('taskUpdates', {})
        if not isinstance(task_updates, dict) or set(task_updates) - set(task_ids):
            raise ValueError('Scoped task updates cannot escape their assigned lot')
        for task, update in task_updates.items():
            if not isinstance(update, dict):
                raise ValueError('A scoped task update needs actual observations')
            for key in ('advances', 'remaining'):
                if key in update:
                    text_list(update[key], 'Task ' + key)
        manifests = value.get('manifests')
        if not isinstance(manifests, list) or not manifests or len(manifests) > 64 or len({row['path'] for row in manifests}) != len(manifests):
            raise ValueError('A delivered lot needs unique bound manifests')
        for proof in manifests:
            checked(proof, root)
        if not any(row['path'].startswith(('modules/fidelity-r31/', 'docs/R31_')) for row in manifests):
            raise ValueError('A new lot cannot consist only of earlier evidence')
        resources = []
        for item in value.get('resources', []):
            local(item['path'], root)
            if not isinstance(item.get('label'), str) or not item['label'].strip():
                raise ValueError('A resource needs a human-readable label')
            resources.append({'path': item['path'], 'label': item['label']})
        identities, seen = [], set()
        for row in value.get('identities', []):
            if row.get('id') not in known or row['id'] in seen or not isinstance(row.get('summary'), str) or len(row['summary']) > 10000:
                raise ValueError('Lot additions need unique exact catalogue identities')
            seen.add(row['id']); artifacts = []
            if not isinstance(row.get('artifacts', []), list) or len(row.get('artifacts', [])) > 256:
                raise ValueError('Identity evidence must fit the bounded Dossier index')
            for item in row.get('artifacts', []):
                target = local(item['path'], root)
                if not isinstance(item.get('label'), str) or not isinstance(item.get('kind'), str):
                    raise ValueError('An identity artifact needs a kind and label')
                proof = binding(item['path'], root)
                if ('sha256' in item and item['sha256'] != proof['sha256']
                        or 'bytes' in item and item['bytes'] != proof['bytes']):
                    raise ValueError('Identity artifact bytes differ: ' + item['path'])
                if item.get('canonicalFidelityCertified') is True or item.get('originalGameFrameCertified') is True:
                    raise ValueError('A source observation cannot certify global fidelity or an original game frame')
                parameter = item.get('identityParameter')
                if parameter is not None and parameter not in ('id', 'uid', 'subjectId', 'identity'):
                    raise ValueError('Unknown exact-identity query parameter')
                default_route = item['path'] + ('?' + urlencode({parameter: row['id']}) if parameter else '')
                route = item.get('href', default_route)
                if not isinstance(route, str) or len(route) > 4096 or '\\' in route:
                    raise ValueError('An identity route must stay with its bound local artifact')
                parsed = urlsplit(route); query = parse_qs(parsed.query, keep_blank_values=True)
                if (parsed.scheme or parsed.netloc or parsed.path != item['path']
                        or any(any(uid != row['id'] for uid in query[key]) for key in ('id', 'uid', 'subjectId', 'identity') if key in query)):
                    raise ValueError('An identity route cannot substitute another incarnation or file')
                artifacts.append({**proof, 'label': item['label'], 'kind': item['kind'], 'sourceEdition': item.get('sourceEdition', ''), 'href': route})
            identities.append({'id': row['id'], 'summary': row['summary'], 'artifacts': artifacts})
        inputs += [name] + [row['path'] for row in manifests]
        deliveries.append({**shell, 'state': 'LIVRÉ PARTIEL', 'summary': value['summary'], 'advances': advances,
            'remaining': remaining, 'manifests': manifests, 'resources': resources, 'identities': identities,
            'taskUpdates': task_updates, 'counts': value.get('counts', {})})
    if pending and not allow_pending:
        raise ValueError('R31 lots are unfinished: ' + ', '.join(pending))
    return deliveries, inputs, pending


def validation(root=ROOT, allow_pending=False):
    index = read(REPORT_INDEX, root) if (Path(root) / REPORT_INDEX).is_file() else {'reports': []}
    checks, pending, failed, totals, names = [], [], [], Counter(), set()
    if not index.get('reports'):
        pending.append(REPORT_INDEX)
    for item in index.get('reports', []):
        name = item['path']
        if name in names:
            raise ValueError('An executed report cannot be counted twice')
        names.add(name)
        if not (Path(root) / name).is_file():
            pending.append(name); continue
        value = read(name, root); counts = value.get('counts', {})
        passed = value.get('passed', counts.get('passed')); failures = value.get('failed', counts.get('failed'))
        if type(passed) is not int or type(failures) is not int or passed < 0 or failures < 0 or passed + failures <= 0:
            raise ValueError('Executed report has no nonempty actual result: ' + name)
        current = failures == 0 and value.get('pass') is not False
        for key in ('sourceFiles', 'inputs', 'sources', 'executedAdapterInputs'):
            for proof in value.get(key, []) if isinstance(value.get(key), list) else []:
                if isinstance(proof, dict) and 'path' in proof and 'sha256' in proof:
                    try:
                        checked(proof, root)
                    except ValueError:
                        current = False
        if not current:
            failed.append(name)
        kind = item.get('kind', 'checks'); totals[kind] += passed if current else 0
        checks.append({'label': item.get('label', name), 'path': name, 'sha256': sha(Path(root) / name),
            'kind': kind, 'passed': passed, 'failed': failures, 'source_current': current,
            'result': f'{passed} réussis, {failures} échec(s)' if current else 'Résultat non courant ou en échec',
            'scope': item.get('scope', 'Ajouts R31 dans le périmètre de ce rapport exécuté.')})
    if not allow_pending and (pending or failed):
        raise ValueError('R31 validation is unfinished: ' + ', '.join(pending + failed))
    return {'schema': 're.bow.validation-r31/1', 'release': 'R31', 'complete': not pending and not failed,
        'checks': checks, 'totals': dict(totals), 'pending_reports': pending, 'failed_reports': failed,
        'canonical_fidelity_certified': 0,
        'historical_results': binding('docs/R30_VALIDATION.json', root)}


def build(root=ROOT, allow_pending=False):
    root = Path(root).resolve(); prior = read(PRIOR, root)
    known = {row['id'] for row in read('data/roster.json', root)['catalog']}
    if len(known) != 389 or {row['id'] for row in prior['identity_updates']} != known:
        raise ValueError('The frozen catalogue must retain389 exact identities')
    lots, lot_inputs, pending = collect_lots(root, allow_pending)
    checks = validation(root, allow_pending)
    checks['pending_reports'] += pending; checks['complete'] = not checks['pending_reports'] and not checks['failed_reports']
    by_uid = {uid: [] for uid in known}; artifacts, sources = [], []
    for lot in lots:
        for row in lot['identities']:
            by_uid[row['id']].append((lot, row))
            for item in row['artifacts']:
                artifacts.append([lot['id'], lot['owner'], row['id'], item['kind'], item['label'], item['sourceEdition'], item['path'], item['sha256'], item['href']])
        for proof in lot['manifests']:
            sources.append([lot['id'], ', '.join(lot['taskIds']), '', 'Manifeste du lot', proof['path'], lot['summary'], proof['sha256']])
        for row in lot['identities']:
            for item in row['artifacts']:
                sources.append([lot['id'], ', '.join(lot['taskIds']), row['id'], item['kind'], item['path'], item['sourceEdition'] or row['summary'], item['sha256']])
    updates = []
    for original in prior['identity_updates']:
        rows = by_uid[original['id']]; row = dict(original)
        row.update(r31_lots=[lot['id'] for lot, item in rows],
            r31_workstreams=sorted({task for lot, item in rows for task in lot['taskIds']}),
            r31_summary=[item['summary'] for lot, item in rows],
            r31_artifacts=[{**proof, 'lot': lot['id']} for lot, item in rows for proof in item['artifacts']],
            r31_source_editions=sorted({proof['sourceEdition'] for lot, item in rows for proof in item['artifacts'] if proof['sourceEdition']}),
            r31_has_additions=bool(rows),
            r31_limits='Les ajouts gardent leur lot, édition et périmètre. Un rendu généré, une source native, une référence documentaire, un geste observé et un idle certifié restent des critères distincts ; fidélité globale 1:1 non certifiée.')
        updates.append(row)
    owners = {task: lot for lot in lots for task in lot['taskIds']}
    expected = set(owners) | set(TRACKING_TASKS)
    if len(expected) != 17 or {row['id'] for row in prior['tasks']} != expected:
        raise ValueError('All17 open workstreams need their separate explicit owner')
    tasks = []
    for old in prior['tasks']:
        row = copy.deepcopy(old); task_id = row['id']; row['status'] = 'EN COURS'
        row['evidence'] = [{'href': '../fidelity-r30/index.html', 'label': 'Acquis antérieurs et critères R30'}]
        if task_id in owners:
            lot = owners[task_id]; scoped = lot.get('taskUpdates', {}).get(task_id, {})
            row.update(owner=lot['owner'], lotId=lot['id'], deliveryState=lot['state'], r31_new_implementation=lot['state'] != 'EN ATTENTE')
            row['advances'] = ['Lot ' + lot['id'] + ' · responsable : ' + lot['owner'] + '.', *scoped.get('advances', lot['advances'])]
            row['remaining'] = list(dict.fromkeys([*old['remaining'], *scoped.get('remaining', lot['remaining'])]))
            if lot['state'] != 'EN ATTENTE':
                row['evidence'].append({'href': href(lot['path']), 'label': 'Livraison et limites du lot ' + lot['id']})
                row['evidence'] += [{'href': href(item['path']), 'label': item['label']} for item in lot['resources'][:4]]
        else:
            row.update(owner='r31_excel_validation', lotId='tracking', deliveryState='LIVRÉ PARTIEL', r31_new_implementation=True)
            row['advances'] = ['Les 17 chantiers sont attribués à 7 lots distincts, avec état de livraison, preuves et critères restants.',
                'Les 56 feuilles R30 et leurs relations sont conservées ; sept feuilles R31 suivent les nouveaux lots et les 389 UID.',
                'La recette ne compte que les rapports réellement exécutés, sans échec ni empreinte périmée.']
            row['remaining'] = list(old['remaining'])
            row['evidence'].append({'href': '../../docs/R31_VALIDATION.json', 'label': 'Recette R31 par périmètre'})
        tasks.append(row)
    available = sum(lot['state'] != 'EN ATTENTE' for lot in lots)
    lot_rows = [[lot['id'], lot['owner'], lot['taskIds'], lot['state'], lot['summary'], lot['counts'], lot['remaining'], lot['path']] for lot in lots]
    lot_rows.append(['tracking', 'r31_excel_validation', list(TRACKING_TASKS), 'LIVRÉ PARTIEL', 'Attribution exhaustive, export conservateur et recette des lots.',
        {'identities': 389, 'workstreams': 17, 'historicalSheets': 56}, ['Revue humaine et complétude canonique restent distinctes de la recette technique.'], 'docs/R31_TRACKING.json'])
    resources = [{'title': 'Lots et responsables', 'description': 'Chaque chantier ouvert possède un responsable et une portée explicite.',
                  'href': '../../docs/R31_LOTS.json', 'label': 'Examiner les7 lots'},
                 {'title': 'Historique R30', 'description': 'Les références et qualifications antérieures sont conservées sans devenir des acquis nouveaux.',
                  'href': '../fidelity-r30/index.html', 'label': 'Consulter R30'}]
    resources += [{'title': lot['id'], 'description': lot['summary'], 'href': href(lot['resources'][0]['path'] if lot['resources'] else lot['path']),
                   'label': 'Examiner le lot ' + lot['id']} for lot in lots if lot['state'] != 'EN ATTENTE']
    count = sum(row['r31_has_additions'] for row in updates)
    stats = [{'label': 'chantiers ouverts attribués', 'value': 17, 'scope': 'Tous les chantiers R30 ont un responsable explicite.'},
             {'label': 'lots parallèles livrés dans leur périmètre', 'value': available + 1, 'scope': 'Six lots de ressources/fonctions et un lot de suivi ; complétude1:1 distincte.'},
             {'label': 'incarnations avec preuves ou audits R31', 'value': count, 'scope': 'Un audit des commandes adaptées ne constitue pas un nouvel actif original, un corps entier ou un idle.'},
             {'label': 'objets de preuve par incarnation', 'value': len(artifacts), 'scope': 'Fichiers locaux réellement présents et liés par SHA-256.'},
             {'label': 'feuilles historiques conservées', 'value': 56, 'scope': 'XML et relations R30 préservés byte à byte.'},
             {'label': 'refus initiaux conservés', 'value': 22, 'scope': 'Le refus de révision Tyrant reste conservé séparément.'},
             {'label': 'chantiers clos par fidélité globale 1:1', 'value': 0, 'scope': 'Les critères ouverts gardent leurs limites.'}]
    input_names = list(dict.fromkeys([BASELINE, PRIOR, 'docs/R30_VALIDATION.json', 'data/roster.json',
        'scripts/fidelity-progress-r31.py', (MODULE / 'index.html').as_posix(), (MODULE / 'dashboard.js').as_posix(), (MODULE / 'dossier-link.js').as_posix(),
        *lot_inputs, *([REPORT_INDEX] if (root / REPORT_INDEX).is_file() else []), *[item['path'] for item in checks['checks']]]))
    summary = f'{available + 1}/7 lots livrés dans leur périmètre ; 17 chantiers attribués et toujours ouverts ; {count} incarnations reliées aux preuves et audits R31. Les médias et leurs qualifications restent distincts ; fidélité globale 1:1 non certifiée.'
    value = {'schema': 're.bow.fidelity-tracking-r31/1', 'release': 'R31', 'canonical_fidelity_certified': 0,
        'summary': summary, 'statistics': stats, 'tasks': tasks, 'lots': lots, 'lot_rows': lot_rows,
        'identity_updates': updates, 'artifact_rows': artifacts, 'source_rows': sources, 'resources': resources,
        'validation': checks['checks'], 'validation_complete': checks['complete'],
        'validation_gaps': checks['pending_reports'] + checks['failed_reports'],
        'generation_preference': prior.get('generation_preference', 'Conserver les refus ; fidélité prioritaire.'),
        'inputs': [binding(name, root) for name in input_names]}
    return checks, value


def validate(value, root=ROOT):
    if value.get('release') != 'R31' or value.get('canonical_fidelity_certified') != 0:
        raise ValueError('Uncertified R31 release required')
    checks, actual = build(root, allow_pending=not value.get('validation_complete'))
    if actual != value:
        raise ValueError('R31 projection differs from current lot proofs or task ownership')
    for item in value['inputs']:
        checked(item, root)
    return value


def progress(value):
    return {**{key: value[key] for key in ('summary', 'statistics', 'resources', 'tasks')},
            'schema': 're.bow.fidelity-progress/1', 'release': 'R31'}


def dossier(value):
    result = {'schema': 're.bow.dossier-additions-r31/1', 'canonicalFidelityCertified': False,
        'identities': [{'id': row['id'], 'lots': row['r31_lots'], 'workstreams': row['r31_workstreams'],
                       'summaries': row['r31_summary'], 'artifacts': row['r31_artifacts'],
                       'sourceEditions': row['r31_source_editions'], 'hasAdditions': row['r31_has_additions']}
                      for row in value['identity_updates']]}
    if len(raw(result)) > 768 * 1024:
        raise ValueError('Dossier proof index exceeds its bounded transport')
    return result


def export(value, root=ROOT, destination=None, progress_path=None):
    root = Path(root).resolve(); validate(value, root)
    baseline = root / BASELINE
    destination = UTIL.checked_path(root, destination or root / 'docs/BOW_FIDELITE_R31.xlsx')
    progress_path = UTIL.checked_path(root, progress_path or root / MODULE / 'progress.json')
    protected = {checked(item, root) for item in value['inputs']}
    if destination == progress_path or destination in protected or progress_path in protected:
        raise ValueError('Outputs cannot overwrite historical inputs')
    for output in (destination, progress_path):
        relative = output.relative_to(root)
        if relative.parts[0] != 'docs' and not relative.parts[0].startswith('qa-') and output != root / MODULE / 'progress.json':
            raise ValueError('Outputs must stay in R31 publication or isolated QA locations')
    if destination.exists():
        check = load_workbook(destination, read_only=True); creator = check.properties.creator; check.close()
        if creator != 'BOW fidelity progress R31':
            raise ValueError('Refusing to overwrite another author’s workbook')
    if progress_path.exists():
        try:
            previous = json.loads(progress_path.read_text())
        except (UnicodeError, json.JSONDecodeError) as error:
            raise ValueError('Refusing to overwrite an unrelated progress file') from error
        if (not isinstance(previous, dict) or previous.get('schema') != 're.bow.fidelity-progress/1'
                or previous.get('release') != 'R31' or not isinstance(previous.get('tasks'), list)
                or len(previous['tasks']) != 17):
            raise ValueError('Refusing to overwrite another release’s progress file')
    book = load_workbook(baseline)
    if book.properties.creator != 'BOW fidelity progress R30' or len(book.sheetnames) != 56 or any(name in book.sheetnames for name in SHEETS):
        raise ValueError('Expected the preserved56-sheet R30 workbook')
    original = list(book.sheetnames); old_catalogue = list(book['Catalogue R30'].values)
    if len(old_catalogue) != 390 or {row[0] for row in old_catalogue[1:]} != {row['id'] for row in value['identity_updates']}:
        raise ValueError('Historical catalogue must preserve389 exact identities')
    book.properties.creator = 'BOW fidelity progress R31'; book.properties.title = 'B.O.W. Protocol — lots parallèles et preuves R31'
    UTIL.sheet(book, SHEETS[0], ['Indicateur', 'Valeur', 'Portée'], [[row['label'], row['value'], row['scope']] for row in value['statistics']]
        + [['Recette R31 complète', value['validation_complete'], 'Périmètre technique des lots, sans clôture canonique1:1.']])
    UTIL.sheet(book, SHEETS[1], ['ID', 'Chantier', 'Responsable', 'Lot', 'Livraison du lot', 'Statut du chantier', 'Acquis R31', 'Critères ouverts', 'Preuves'],
        [[row['id'], row['title'], row['owner'], row['lotId'], row['deliveryState'], row['status'], row['advances'], row['remaining'], row['evidence']] for row in value['tasks']])
    updates = {row['id']: row for row in value['identity_updates']}
    UTIL.sheet(book, SHEETS[2], list(old_catalogue[0]) + list(FIELDS), [list(row) + [updates[row[0]][key] for key in FIELDS] for row in old_catalogue[1:]])
    UTIL.sheet(book, SHEETS[3], ['Lot', 'Responsable', 'Chantiers', 'État de livraison', 'Résumé', 'Compteurs propres au lot', 'Critères ouverts', 'Preuve'], value['lot_rows'])
    UTIL.sheet(book, SHEETS[4], ['Lot', 'Responsable', 'UID', 'Nature', 'Description', 'Édition source', 'Fichier', 'SHA-256', 'Route par incarnation'], value['artifact_rows'])
    UTIL.sheet(book, SHEETS[5], ['Contrôle', 'Résultat', 'Portée', 'Rapport', 'SHA-256'],
        [[row['label'], row['result'], row['scope'], row['path'], row['sha256']] for row in value['validation']]
        + [['Rapport ou lot manquant/non courant', 'NON VALIDÉ', 'Résultat exclu du cumul R31', name, ''] for name in value['validation_gaps']])
    UTIL.sheet(book, SHEETS[6], ['Lot', 'Chantiers', 'UID', 'Nature', 'Fichier', 'Portée', 'SHA-256'], value['source_rows'])
    book.active = book.sheetnames.index('Chantiers R31'); destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=destination.parent, suffix='.xlsx', delete=False) as handle:
        temporary = Path(handle.name)
    try:
        book.save(temporary); book.close(); PREVIOUS.OLD.canonical_zip(baseline, temporary)
        with ZipFile(baseline) as before, ZipFile(temporary) as after:
            for name in before.namelist():
                if name.startswith('xl/worksheets/') and before.read(name) != after.read(name):
                    raise ValueError('Historical worksheet XML or relationships changed')
        check = load_workbook(temporary)
        try:
            if check.sheetnames[:56] != original or check['Catalogue R31'].max_row != 390 or check['Chantiers R31'].max_row != 18:
                raise ValueError('New sheets omitted a historical identity or workstream')
            if any(cell.data_type == 'f' for name in SHEETS for row in check[name] for cell in row):
                raise ValueError('New evidence sheets cannot contain executable formulas')
        finally:
            check.close()
        os.replace(temporary, destination); write(progress_path, progress(value))
    finally:
        temporary.unlink(missing_ok=True)
    return {'workbook': destination.relative_to(root).as_posix(), 'sha256': sha(destination), 'baselineSha256': sha(baseline),
        'historical_worksheets_preserved': 56, 'worksheets': 63, 'identities': 389, 'open_workstreams': 17,
        'lots': 7, 'artifact_rows': len(value['artifact_rows']), 'source_rows': len(value['source_rows']), 'canonical_fidelity_certified': 0}


def overview(value, checks):
    return f'''# R31 — Tous les chantiers ouverts, répartis en lots

{value['summary']}

Le [classeur R31](BOW_FIDELITE_R31.xlsx) conserve exactement les 56 feuilles R30 et leurs relations. Sept feuilles suivent les responsables, livraisons, nouveaux fichiers et critères encore ouverts. Le [suivi central](../modules/fidelity-r31/index.html) donne accès aux preuves de chaque lot.

Une livraison partielle ne clôt aucun chantier : rendu généré, source native, référence documentaire, geste observé, édition et cycle idle qualifié restent des preuves différentes. Les 22 refus initiaux et le refus de révision Tyrant sont conservés. Les acquis R30 ne sont pas recomptés comme nouveaux.

La recette est {'complète' if checks['complete'] else 'provisoire'} dans son périmètre : {json.dumps(checks['totals'], ensure_ascii=False)}. Les sorties après export restent séparées pour éviter les dépendances circulaires de SHA.

Preuves : [lots](R31_LOTS.json), [suivi structuré](R31_TRACKING.json), [recette](R31_VALIDATION.json), [export](R31_EXPORT.json), [reproduction](R31_REPRODUCTION.json).

```sh
python3 scripts/fidelity-progress-r31.py build
python3 scripts/fidelity-progress-r31.py check
python3 scripts/fidelity-progress-r31.py export
python3 scripts/fidelity-progress-r31.py reproduce
```
'''.encode()


def reproduce(value, root=ROOT):
    root = Path(root).resolve(); checks = []
    try:
        for seed in ('19', '79'):
            command = [sys.executable, 'scripts/fidelity-progress-r31.py', 'check']
            if not value['validation_complete']:
                command.append('--allow-pending')
            completed = subprocess.run(command, cwd=root, env={**os.environ, 'PYTHONHASHSEED': seed}, capture_output=True, text=True)
            if completed.returncode:
                raise ValueError(completed.stderr)
        checks.append({'name': 'Tracking is stable across independent Python hash seeds', 'pass': True})
        folder = root / 'qa-r31-reproduction'; folder.mkdir(exist_ok=True); baseline = sha(root / BASELINE)
        with tempfile.TemporaryDirectory(dir=folder) as name:
            temporary = Path(name); result = export(value, root, temporary / 'reproduced.xlsx', temporary / 'progress.json')
            if result['sha256'] != sha(root / 'docs/BOW_FIDELITE_R31.xlsx'):
                raise ValueError('Actual R31 workbook bytes do not reproduce')
            checks.append({'name': 'Actual R31 workbook reproduces exactly', 'pass': True, 'sha256': result['sha256']})
            if (temporary / 'progress.json').read_bytes() != (root / MODULE / 'progress.json').read_bytes() or sha(root / BASELINE) != baseline:
                raise ValueError('Progress or preserved R30 baseline changed')
            checks.append({'name': 'Progress matches and R30 history remains byte-exact', 'pass': True})
    except Exception as error:
        checks.append({'name': 'Reproduction failure', 'pass': False, 'error': str(error)})
    return {'schema': 're.bow.fidelity-reproduction-r31/1', 'passed': sum(row['pass'] for row in checks), 'failed': sum(not row['pass'] for row in checks), 'checks': checks}


def main():
    parser = argparse.ArgumentParser(description=__doc__); parser.add_argument('action', choices=('build', 'check', 'export', 'reproduce'))
    parser.add_argument('--allow-pending', action='store_true'); args = parser.parse_args()
    if args.action in ('build', 'check'):
        checks, value = build(allow_pending=args.allow_pending)
        outputs = {'docs/R31_VALIDATION.json': raw(checks), 'docs/R31_TRACKING.json': raw(value),
            'docs/R31_WORKSTREAMS.json': raw({key: value[key] for key in ('release', 'summary', 'tasks', 'generation_preference')}),
            'docs/R31_LOTS.json': raw({'schema': 're.bow.fidelity-lots-r31/1', 'release': 'R31', 'lots': value['lot_rows']}),
            (MODULE / 'progress.json').as_posix(): raw(progress(value)),
            (MODULE / 'dossier-additions.json').as_posix(): raw(dossier(value)), 'docs/FIDELITE_R31.md': overview(value, checks)}
        for name, content in outputs.items():
            if args.action == 'check':
                if (ROOT / name).read_bytes() != content:
                    raise ValueError('R31 output differs from current lot evidence: ' + name)
            else:
                write(ROOT / name, content)
        print(json.dumps({'complete': checks['complete'], 'totals': checks['totals'], 'lots': len(value['lots']) + 1,
                          'identitiesWithAdditions': sum(row['r31_has_additions'] for row in value['identity_updates'])}))
    elif args.action == 'export':
        result = export(read('docs/R31_TRACKING.json')); write(ROOT / 'docs/R31_EXPORT.json', result); print(json.dumps(result))
    else:
        result = reproduce(read('docs/R31_TRACKING.json')); write(ROOT / 'docs/R31_REPRODUCTION.json', result); print(json.dumps(result)); return int(bool(result['failed']))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
