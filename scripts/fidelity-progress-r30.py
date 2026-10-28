#!/usr/bin/env python3
"""Append exact-identity native sprite, motion and selection evidence to R29."""
import argparse
from collections import Counter
import hashlib
import importlib.util
import json
import math
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MODULE = Path('modules/fidelity-r30')
NATIVE = (MODULE / 'native/manifest.json').as_posix()
AVATARS = (MODULE / 'selection/manifest.json').as_posix()
CONTEXT_GRAPH = 'modules/fidelity-r28/evolutions/graph.json'
BASELINE = 'docs/BOW_FIDELITE_R29.xlsx'
SHEETS = ('Synthèse R30', 'Chantiers R30', 'Catalogue R30', 'Natif R30', 'Avatars R30', 'Recette R30', 'Sources R30')
FIELDS = ('r30_native_entries', 'r30_native_rendered_frames', 'r30_native_source_editions',
          'r30_native_static_views', 'r30_native_combat_stances', 'r30_native_smooth_sequences',
          'r30_native_verified_idle_cycles', 'r30_native_optional_neutral_variants', 'r30_native_playable_optional_neutral_variants',
          'r30_context_only', 'r30_native_edition_variants', 'r30_native_gameplay_eligible',
          'r30_avatar_image', 'r30_avatar_source_kind', 'r30_avatar_source_edition',
          'r30_native_combat_pending', 'r30_full_idle_pending', 'r30_avatar_pending', 'r30_limits')
SPEC = importlib.util.spec_from_file_location('r30_previous_export', ROOT / 'scripts/export-fidelity-progress-r28.py')
OLD = importlib.util.module_from_spec(SPEC); SPEC.loader.exec_module(OLD)
UTIL = OLD.OLD
# The delivery owner may add current executed reports without changing this script.
REPORT_INDEX = 'docs/R30_REPORT_INDEX.json'


def read(path, root=ROOT):
    return json.loads((Path(root) / path).read_text())


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def raw(value):
    return (json.dumps(value, ensure_ascii=False, indent=2) + '\n').encode()


def write(path, value):
    UTIL.atomic_write(Path(path), value if isinstance(value, bytes) else raw(value))


def checked_binding(record, root=ROOT):
    path = UTIL.local_file(root, record['path'])
    if (not re.fullmatch('[a-f0-9]{64}', str(record.get('sha256', '')))
            or sha(path) != record['sha256'] or ('bytes' in record and path.stat().st_size != record['bytes'])):
        raise ValueError('Evidence hash or size mismatch: ' + record['path'])
    return path


def smooth_sequence(entry):
    rate, duration = entry.get('sampleRate'), entry.get('durationSeconds')
    return (entry.get('kind') == 'native-sequence' and entry.get('fullNativeDuration') is True
            and type(rate) in (int, float) and math.isfinite(rate) and 30 <= rate <= 120
            and type(duration) in (int, float) and math.isfinite(duration) and 0 < duration <= 600
            and len(entry.get('frames', [])) >= 30
            and abs(len(entry['frames']) - math.ceil(rate * duration - 1e-7)) <= 1)


def admission(entry):
    """A static stance and a verified native idle are separate capabilities."""
    common = (entry.get('admittedTargetEdition') is True and entry.get('fullBodyReviewed') is True
              and entry.get('fightingPoseReviewed') is True and entry.get('sourceFacing') in ('left', 'right')
              and entry.get('canMirror') is False and not entry.get('knownAppearanceGaps')
              and entry.get('admittedCombatStates') == ['idle'])
    static = common and entry.get('kind') == 'native-static-pose' and entry.get('actionIdentity') == 'neutral-stance'
    animated = (common and entry.get('kind') == 'native-sequence' and entry.get('actionIdentity') == 'idle'
                and entry.get('actionVerified') is True and entry.get('continuityVerified') is True
                and smooth_sequence(entry))
    return static, animated


def observed_variant(entry):
    """An optional observed source-edition loop does not certify a native action."""
    return (entry.get('kind') == 'native-sequence' and entry.get('admittedTargetEdition') is False
            and entry.get('sourceActionCertified') is False and entry.get('observedNeutralLoopReviewed') is True
            and entry.get('continuityVerified') is True and entry.get('fullBodyReviewed') is True
            and entry.get('fightingPoseReviewed') is True and entry.get('sourceFacing') in ('left', 'right')
            and entry.get('canMirror') is False and not entry.get('knownAppearanceGaps')
            and entry.get('selectablePresentationVariant') is True and entry.get('variantCombatStates') == ['idle']
            and isinstance(entry.get('variantFamilyId'), str) and bool(entry['variantFamilyId'])
            and entry.get('admittedCombatStates') == [] and smooth_sequence(entry))


def native_entries(root=ROOT, allow_pending=False):
    root = Path(root).resolve()
    if not (root / NATIVE).is_file():
        if allow_pending:
            return [], [], [NATIVE]
        raise ValueError('Native index is missing')
    index = read(NATIVE, root)
    if index.get('schema') != 're.bow.native-2d-index-r30/1' or index.get('release') != 'R30':
        raise ValueError('Unsupported native index')
    entries, inputs, seen = [], [NATIVE], set()
    known = {item['id'] for item in read('data/roster.json', root)['catalog']}
    lots = index.get('lots', [])
    if not isinstance(lots, list) or len(lots) > 16 or len({lot['path'] for lot in lots}) != len(lots):
        raise ValueError('Native lots must be unique and bounded')
    for lot in lots:
        path = checked_binding(lot, root)
        if not lot['path'].startswith('modules/fidelity-r30/native/') or path.name != 'manifest.json':
            raise ValueError('Native lot leaves its release')
        pack = read(lot['path'], root); inputs.append(lot['path'])
        if pack.get('schema') != 're.bow.native-2d-lot-r30/1' or pack.get('release') != 'R30':
            raise ValueError('Unsupported native lot')
        for original in pack.get('entries', []):
            entry = dict(original); entry['manifest'] = lot['path']
            if (entry.get('subjectId') not in known or not isinstance(entry.get('id'), str)
                    or entry['id'] in seen or entry.get('canonicalFidelityCertified') is not False
                    or entry.get('canMirror') is not False or not entry.get('sourceEdition')
                    or entry.get('kind') not in ('native-static-pose', 'native-sequence')):
                raise ValueError('Native entry identity or scope is invalid')
            seen.add(entry['id'])
            frames = entry.get('frames', [])
            if not frames or len(frames) > 36000 or entry.get('frameCount', len(frames)) != len(frames):
                raise ValueError('Native frames must be delivered, bounded and counted exactly')
            paths = set()
            for frame in frames:
                if not frame.get('path', '').startswith('modules/fidelity-r30/native/') or not frame['path'].endswith('.png'):
                    raise ValueError('Native frame path leaves its release')
                checked_binding(frame, root)
                if frame['path'] in paths:
                    raise ValueError('Native frame paths must be unique within a sequence')
                paths.add(frame['path'])
            static, animated = admission(entry)
            if entry.get('admittedCombatStates') and not (static or animated):
                raise ValueError('A documentary, unverified or sparse scene cannot acquire combat idle')
            if entry.get('variantCombatStates') and not observed_variant(entry):
                raise ValueError('An optional source-edition loop requires its complete observed qualification')
            if entry.get('originalGameFrames') is True:
                raise ValueError('Derived native rendering cannot certify original game frames')
            entries.append(entry)
    pending = index.get('pendingLots', [])
    if not isinstance(pending, list) or any(not isinstance(path, str) or not path.startswith('modules/fidelity-r30/native/') for path in pending):
        raise ValueError('Pending native lots must remain explicit local paths')
    if pending and not allow_pending:
        raise ValueError('Native lots are unfinished: ' + ', '.join(pending))
    return entries, inputs, pending


def avatar_entries(root=ROOT, allow_pending=False):
    root = Path(root).resolve()
    if not (root / AVATARS).is_file():
        if allow_pending:
            return [], [], [AVATARS]
        raise ValueError('Avatar manifest is missing')
    pack = read(AVATARS, root)
    if pack.get('schema') != 're.bow.selection-avatars-r30/1':
        raise ValueError('Unsupported avatar manifest')
    known = {row['id'] for row in read('data/roster.json', root)['catalog']}; seen = set()
    for entry in pack.get('entries', []):
        if (entry.get('id') not in known or entry['id'] in seen or not entry.get('sourceKind')
                or entry.get('canonicalFidelityCertified') is not False or not entry.get('sourceEdition')):
            raise ValueError('An avatar needs a unique exact identity and source edition')
        seen.add(entry['id'])
        checked_binding({'path': entry['image'], 'sha256': entry['sha256'], 'bytes': entry['bytes']}, root)
    missing = {entry['id'] if isinstance(entry, dict) else entry for entry in pack.get('missing', [])}
    if missing & seen or missing | seen != known or len(missing) != len(pack.get('missing', [])):
        raise ValueError('All 389 identities require an avatar or an explicit missing entry')
    return pack.get('entries', []), [AVATARS], []


def validation(root=ROOT, allow_pending=False):
    root = Path(root).resolve(); pending, failed, checks = [], [], []
    totals = Counter()
    index = read(REPORT_INDEX, root) if (root / REPORT_INDEX).is_file() else {'reports': []}
    if not index.get('reports'):
        pending.append(REPORT_INDEX)
    for binding in index.get('reports', []):
        path = binding['path']
        if not (root / path).is_file():
            pending.append(path); continue
        report = read(path, root); counts = report.get('counts', {})
        passed = report.get('passed', counts.get('passed')); failures = report.get('failed', counts.get('failed'))
        if type(passed) is not int or type(failures) is not int or passed + failures <= 0:
            raise ValueError('Executed report has no real result: ' + path)
        current = failures == 0 and report.get('pass') is not False
        for key in ('sourceFiles', 'executedAdapterInputs', 'sources', 'inputs'):
            for source in report.get(key, []) if isinstance(report.get(key), list) else []:
                if isinstance(source, dict) and isinstance(source.get('path'), str) and isinstance(source.get('sha256'), str):
                    try:
                        checked_binding(source, root)
                    except ValueError:
                        current = False
        if not current:
            failed.append(path)
        kind = binding.get('kind', 'checks'); totals[kind] += passed if current else 0
        checks.append({'label': binding.get('label', path), 'path': path, 'sha256': sha(root / path),
            'kind': kind, 'passed': passed, 'failed': failures, 'source_current': current,
            'result': f'{passed} réussis, {failures} échec(s)' if current else 'Résultat non courant ou en échec',
            'scope': binding.get('scope', 'Ajouts R30 dans le périmètre du rapport exécuté.')})
    if not allow_pending and (pending or failed):
        raise ValueError('R30 validation is unfinished: ' + ', '.join(pending + failed))
    return {'schema': 're.bow.validation-r30/1', 'release': 'R30', 'complete': not pending and not failed,
        'checks': checks, 'totals': dict(totals), 'pending_reports': pending, 'failed_reports': failed,
        'canonical_fidelity_certified': 0,
        'historical_results': {'path': 'docs/R29_VALIDATION.json', 'sha256': sha(root / 'docs/R29_VALIDATION.json'),
            'scope': 'Conservé sans recompter les résultats historiques comme exécutés en R30.'}}


def build(root=ROOT, allow_pending=False):
    root = Path(root).resolve(); prior = read('docs/R29_TRACKING.json', root)
    contexts = set(read(CONTEXT_GRAPH, root)['contextOnlyUids'])
    native, native_inputs, pending_native = native_entries(root, allow_pending)
    avatars, avatar_inputs, pending_avatars = avatar_entries(root, allow_pending)
    checks = validation(root, allow_pending)
    checks['pending_reports'] += pending_native + pending_avatars
    checks['complete'] = not checks['pending_reports'] and not checks['failed_reports']
    by_uid = {row['id']: [] for row in read('data/roster.json', root)['catalog']}
    for entry in native:
        by_uid[entry['subjectId']].append(entry)
    portraits = {row['id']: row for row in avatars}
    updates, native_rows, avatar_rows, sources = [], [], [], []
    for old in prior['identity_updates']:
        uid = old['id']; entries = by_uid[uid]; portrait = portraits.get(uid)
        static = sum(admission(row)[0] for row in entries); animated = sum(admission(row)[1] for row in entries)
        smooth = sum(smooth_sequence(row) for row in entries)
        update = dict(old)
        update.update(r30_native_entries=[row['id'] for row in entries],
            r30_native_rendered_frames=sum(len(row['frames']) for row in entries),
            r30_native_source_editions=sorted({row['sourceEdition'] for row in entries}),
            r30_native_static_views=sum(row['kind'] == 'native-static-pose' for row in entries),
            r30_native_combat_stances=static, r30_native_smooth_sequences=smooth,
            r30_native_verified_idle_cycles=animated,
            r30_native_optional_neutral_variants=len({row['variantFamilyId'] for row in entries if observed_variant(row)}),
            r30_native_playable_optional_neutral_variants=0 if uid in contexts else len({row['variantFamilyId'] for row in entries if observed_variant(row)}),
            r30_context_only=uid in contexts,
            r30_native_edition_variants=sum(row.get('admittedTargetEdition') is not True for row in entries),
            r30_native_gameplay_eligible=bool(static or animated),
            r30_avatar_image=portrait['image'] if portrait else '',
            r30_avatar_source_kind=portrait['sourceKind'] if portrait else '',
            r30_avatar_source_edition=portrait['sourceEdition'] if portrait else '',
            r30_native_combat_pending=not bool(static or animated), r30_full_idle_pending=not bool(animated),
            r30_avatar_pending=portrait is None,
            r30_limits='Les portraits prouvent seulement leur incarnation visible. Une pose source, une séquence fluide non attribuée, une posture de combat et un idle natif bouclé restent des critères distincts ; fidélité globale 1:1 non certifiée.')
        updates.append(update)
    for entry in native:
        static, animated = admission(entry)
        native_rows.append([entry['subjectId'], entry['id'], entry['kind'], entry['sourceEdition'],
            entry.get('admittedTargetEdition') is True, len(entry['frames']), entry.get('sampleRate', ''),
            entry.get('durationSeconds', ''), entry.get('fullNativeDuration', False), entry.get('actionIdentity', 'unassigned'),
            entry.get('actionVerified', False), entry.get('continuityVerified', False), static, animated,
            observed_variant(entry), observed_variant(entry) and entry['subjectId'] not in contexts,
            entry['subjectId'] in contexts, entry.get('variantFamilyId', ''), entry.get('sourceActionCertified', False),
            entry.get('sourceFacing', 'unknown'), entry.get('knownAppearanceGaps', []), entry['manifest']])
        sources.append(['R20-09' if entry['kind'] == 'native-sequence' else 'R20-10', entry['subjectId'],
            'Séquence native dérivée' if entry['kind'] == 'native-sequence' else 'Pose native dérivée',
            entry['manifest'], entry['sourceEdition'] + '; ' + entry.get('actionIdentity', 'unassigned'), sha(root / entry['manifest'])])
    for entry in avatars:
        avatar_rows.append([entry['id'], entry.get('label', ''), entry.get('game', ''), entry['image'],
            entry['sourceKind'], entry['sourceEdition'], entry.get('sourcePage', ''), entry['sha256'], False])
        sources.append(['G-PREUVES', entry['id'], 'Référence d’avatar de sélection', entry['image'],
            entry['sourceKind'] + '; ' + entry['sourceEdition'] + '; portrait uniquement', entry['sha256']])
    totals = Counter()
    for row in updates:
        for key in ('r30_native_rendered_frames', 'r30_native_static_views', 'r30_native_combat_stances',
                    'r30_native_smooth_sequences', 'r30_native_verified_idle_cycles', 'r30_native_edition_variants',
                    'r30_native_optional_neutral_variants', 'r30_native_playable_optional_neutral_variants',
                    'r30_native_combat_pending', 'r30_full_idle_pending', 'r30_avatar_pending'):
            totals[key] += row[key]
    avatar_kinds = Counter(row['sourceKind'] for row in avatars)
    statistics = [
        {'label': 'avatars attribués à leur UID', 'value': len(avatars), 'scope': 'Images sources de sélection, distinctes des corps et animations.'},
        {'label': 'vignettes de menu natives extraites', 'value': avatar_kinds['native-menu-thumbnail'], 'scope': 'Sous-ensemble des avatars ; les références secondaires restent distinguées.'},
        {'label': 'avatars dérivés de modèles natifs exacts', 'value': avatar_kinds['native-model-render'], 'scope': 'Projection nouvelle du modèle source, distincte d’une image 2D originale extraite du jeu.'},
        {'label': 'avatars encore ouverts', 'value': totals['r30_avatar_pending'], 'scope': 'Chaque manque reste explicite par incarnation.'},
        {'label': 'vues statiques natives dérivées', 'value': totals['r30_native_static_views'], 'scope': 'Les poses de base non combattantes restent des références.'},
        {'label': 'postures natives de combat admissibles', 'value': totals['r30_native_combat_stances'], 'scope': 'Édition, corps entier, pose combattante et côté qualifiés.'},
        {'label': 'séquences natives fluides complètes', 'value': totals['r30_native_smooth_sequences'], 'scope': 'Durée source complète à 30 images/s au moins ; geste et boucle qualifiés séparément.'},
        {'label': 'images natives dérivées livrées', 'value': totals['r30_native_rendered_frames'], 'scope': 'Rendus des sources, distincts des images originales du moteur du jeu.'},
        {'label': 'idles natifs complets vérifiés', 'value': totals['r30_native_verified_idle_cycles'], 'scope': 'Geste idle et continuité vérifiés, admissibles pour la bonne édition.'},
        {'label': 'familles de boucles neutres sources observées', 'value': totals['r30_native_optional_neutral_variants'], 'scope': 'Boucles revues de l’édition source ; attribution originale du geste non certifiée, édition cible non admise.'},
        {'label': 'variantes neutres optionnelles jouables', 'value': totals['r30_native_playable_optional_neutral_variants'], 'scope': 'Activation volontaire sur les acteurs existants ; les rencontres contextuelles ne deviennent pas des combattants.'},
        {'label': 'incarnations sans idle natif qualifié', 'value': totals['r30_full_idle_pending'], 'scope': 'Une pose statique ne réduit pas ce manque.'},
        {'label': 'incarnations sans posture native admissible', 'value': totals['r30_native_combat_pending'], 'scope': 'Les variantes d’une autre édition restent séparées.'},
        {'label': 'entrées d’édition distincte non admises', 'value': totals['r30_native_edition_variants'], 'scope': 'Aucune promotion silencieuse vers l’édition cible.'},
        {'label': 'chantiers ouverts', 'value': 17, 'scope': 'Aucune clôture automatique par présence de fichiers.'}]
    facts = {
        'R20-09': [f"{totals['r30_native_smooth_sequences']} séquences natives complètes et fluides, {totals['r30_native_rendered_frames']} images dérivées. {totals['r30_native_verified_idle_cycles']} idles natifs complets qualifiés.",
                   f"{totals['r30_native_optional_neutral_variants']} familles de boucles neutres sont observées dans l’édition source, dont {totals['r30_native_playable_optional_neutral_variants']} variantes activables volontairement sur des acteurs existants. Les contextes ne deviennent pas jouables ; l’action originale et l’édition cible restent non certifiées.",
                   f"{totals['r30_full_idle_pending']} incarnations sans idle natif complet qualifié restent explicitement ouvertes."],
        'R20-10': [f"{totals['r30_native_static_views']} vues natives dérivées ; {totals['r30_native_combat_stances']} postures de combat admissibles. {len(avatars)} avatars exact-UID documentés.",
                   'Les 363 images générées R29, les 22 refus historiques et le refus de révision Tyrant restent conservés séparément des nouvelles sources natives.'],
        'G-PREUVES': ['Les nouveaux manifestes relient chaque UID aux fichiers, empreintes, édition et observations réellement fournis.'],
        'G-RENDU': ['La projection 2D des modèles et courbes natifs garde un cadrage fixe. Les séquences non attribuées et variantes d’édition restent documentaires.'],
        'G-ACCESS': ['Les avatars de sélection conservent leur identité et leur rôle de portrait ; les groupes et commandes existants restent soumis à la recette navigateur.'],
        'G-PERF': ['Les images et animations sont chargées dans les bornes du lecteur ; les résultats appareils physiques gardent leurs critères propres.'],
        'G-RECETTE': ['Les recettes R30 distinguent fichiers, admission au combat, lecture fluide, sélection et conservation historique.'],
        'G-LIVRAISON': ['Les 49 feuilles R29 et leurs relations sont conservées ; sept feuilles R30 distinguent les avatars, poses natives et vrais cycles idle.']}
    tasks = []
    task_links = {
        'R20-09': [('native/index.html', 'Séquences natives et qualifications')],
        'R20-10': [('native/manifest.json', 'Registre natif R30'), ('selection/manifest.json', 'Avatars par incarnation')],
        'G-PREUVES': [('native/manifest.json', 'Fichiers et périmètres des lots natifs'), ('selection/manifest.json', 'Origines des portraits')],
        'G-RENDU': [('native/index.html', 'Examiner les postures et mouvements'), ('combat/native-runtime.js', 'Contrat du rendu natif')],
        'G-ACCESS': [('selection/manifest.json', 'Identités et portraits de sélection')],
        'G-PERF': [('combat/native-runtime.js', 'Chargement borné des images et animations')],
        'G-RECETTE': [('../../docs/R30_VALIDATION.json', 'Recette R30 par périmètre')],
        'G-LIVRAISON': [('index.html', 'Suivi des 17 chantiers R30')]}
    for old in prior['tasks']:
        task = {key: old[key] for key in ('id', 'title', 'category')}
        links = [{'href': href, 'label': label} for href, label in task_links.get(task['id'], [])
                 if (root / MODULE / href).resolve().is_file()]
        if not links:
            links = ([{'href': '../fidelity-r29/' + proof['href'], 'label': proof['label'] + ' · historique R29'}
                      for proof in old['evidence']] if task['id'] not in facts
                     else [{'href': 'index.html', 'label': 'Suivi central R30'}])
        task.update(status='EN COURS', advances=facts.get(task['id'], ['Acquis antérieurs conservés dans leur périmètre documenté.']),
            remaining=list(old['remaining']), evidence=links,
            r30_new_implementation=task['id'] in facts)
        if task['id'] == 'R20-09':
            task['remaining'] = [f"{totals['r30_full_idle_pending']} incarnations sans idle natif complet vérifié.",
                                 'Attribution des gestes, continuité native, équipements, orientations et édition cible restent vérifiés séparément.']
        tasks.append(task)
    resources = [{'href': 'native/index.html', 'title': 'Sprites et mouvements natifs', 'label': 'Comparer les sources et leurs qualifications'},
        {'href': 'selection/manifest.json', 'title': 'Avatars de sélection', 'label': 'Consulter les portraits attribués par UID'},
        {'href': '../../docs/BOW_FIDELITE_R29.xlsx', 'title': 'Historique R29', 'label': 'Conserver les 49 feuilles antérieures'}]
    resources = [item for item in resources if (root / MODULE / item['href']).resolve().is_file()]
    inputs = [BASELINE, 'docs/R29_TRACKING.json', 'data/roster.json', CONTEXT_GRAPH, 'modules/fidelity-r29/sprites/manifest.json',
              'scripts/fidelity-progress-r30.py', (MODULE / 'index.html').as_posix(), (MODULE / 'dossier-link.js').as_posix()]
    inputs += native_inputs + avatar_inputs + [row['path'] for row in checks['checks']]
    if (root / REPORT_INDEX).is_file():
        inputs.append(REPORT_INDEX)
    value = {'schema': 're.bow.parallel-workstreams/3', 'release': 'R30', 'canonical_fidelity_certified': 0,
        'summary': f"{len(avatars)} avatars exact-UID, {totals['r30_native_static_views']} vues natives et {totals['r30_native_smooth_sequences']} séquences fluides complètes ; {totals['r30_native_optional_neutral_variants']} familles de boucles neutres sources observées, dont {totals['r30_native_playable_optional_neutral_variants']} variantes optionnelles jouables. {totals['r30_native_verified_idle_cycles']} idles natifs qualifiés pour la cible ; 17 chantiers ouverts.",
        'statistics': statistics, 'tasks': tasks, 'identity_updates': updates, 'native_rows': native_rows,
        'avatar_rows': avatar_rows, 'source_rows': sources, 'resources': resources,
        'validation': checks['checks'], 'validation_complete': checks['complete'],
        'validation_gaps': checks['pending_reports'] + checks['failed_reports'],
        'native_index': NATIVE if native_inputs else None, 'avatar_manifest': AVATARS if avatar_inputs else None,
        'avatar_source_kind_counts': dict(avatar_kinds),
        'generation_preference': prior['generation_preference'],
        'inputs': [{'path': path, 'sha256': sha(root / path)} for path in dict.fromkeys(inputs)]}
    return checks, value


def validate(value, root=ROOT):
    root = Path(root).resolve()
    if value.get('release') != 'R30' or value.get('schema') != 're.bow.parallel-workstreams/3' or value.get('canonical_fidelity_certified') != 0:
        raise ValueError('Unsupported R30 tracking scope')
    known = {item['id'] for item in read('data/roster.json', root)['catalog']}
    updates = value.get('identity_updates', [])
    if len(known) != 389 or len(updates) != 389 or {item['id'] for item in updates} != known or any(any(key not in item for key in FIELDS) for item in updates):
        raise ValueError('All 389 identities require unique complete R30 progress')
    expected = {task['id'] for task in read('docs/R29_TRACKING.json', root)['tasks']}
    tasks = value.get('tasks', [])
    if len(tasks) != 17 or {task['id'] for task in tasks} != expected or any(task.get('status') != 'EN COURS' for task in tasks):
        raise ValueError('All seventeen open workstreams must remain explicit')
    for task in tasks:
        if not all(task.get(key) for key in ('title', 'advances', 'remaining', 'evidence')):
            raise ValueError('An open workstream needs actual progress, limits and proof')
    for link in [proof for task in tasks for proof in task['evidence']] + value.get('resources', []):
        href = link.get('href')
        if not isinstance(href, str) or not href or '\\' in href or href.startswith('/') or re.match(r'^[a-zA-Z][a-zA-Z0-9+.-]*:', href):
            raise ValueError('Evidence links must stay local')
        UTIL.checked_path(root, MODULE / href, exists=True)
    inputs = value.get('inputs', [])
    if not inputs or len({item['path'] for item in inputs}) != len(inputs) or BASELINE not in {item['path'] for item in inputs}:
        raise ValueError('Unique evidence inputs must bind the historical workbook')
    for item in inputs:
        checked_binding(item, root)
    # Rebuild from current source contracts, rather than trusting workbook counters.
    entries, _, _ = native_entries(root, allow_pending=not value.get('validation_complete'))
    avatars, _, _ = avatar_entries(root, allow_pending=not value.get('validation_complete'))
    contexts = set(read(CONTEXT_GRAPH, root)['contextOnlyUids'])
    per_uid = {uid: [] for uid in known}
    for entry in entries:
        per_uid[entry['subjectId']].append(entry)
    portraits = {entry['id']: entry for entry in avatars}
    for item in updates:
        native = per_uid[item['id']]; portrait = portraits.get(item['id'])
        static = sum(admission(entry)[0] for entry in native); idle = sum(admission(entry)[1] for entry in native)
        projection = {'r30_native_entries': [entry['id'] for entry in native],
            'r30_native_rendered_frames': sum(len(entry['frames']) for entry in native),
            'r30_native_source_editions': sorted({entry['sourceEdition'] for entry in native}),
            'r30_native_static_views': sum(entry['kind'] == 'native-static-pose' for entry in native),
            'r30_native_smooth_sequences': sum(smooth_sequence(entry) for entry in native),
            'r30_native_edition_variants': sum(entry.get('admittedTargetEdition') is not True for entry in native),
            'r30_native_combat_stances': static, 'r30_native_verified_idle_cycles': idle,
            'r30_native_optional_neutral_variants': len({entry['variantFamilyId'] for entry in native if observed_variant(entry)}),
            'r30_native_playable_optional_neutral_variants': 0 if item['id'] in contexts else len({entry['variantFamilyId'] for entry in native if observed_variant(entry)}),
            'r30_context_only': item['id'] in contexts,
            'r30_native_gameplay_eligible': bool(static or idle), 'r30_native_combat_pending': not bool(static or idle),
            'r30_full_idle_pending': not bool(idle), 'r30_avatar_pending': portrait is None,
            'r30_avatar_image': portrait['image'] if portrait else '',
            'r30_avatar_source_kind': portrait['sourceKind'] if portrait else '',
            'r30_avatar_source_edition': portrait['sourceEdition'] if portrait else ''}
        if any(item[key] != expected_value for key, expected_value in projection.items()):
            raise ValueError('Native/idle/avatar coverage differs from exact-UID evidence')
    return value


def progress(value):
    return {**{key: value[key] for key in ('summary', 'statistics', 'resources', 'tasks')},
            'schema': 're.bow.fidelity-progress/1', 'release': 'R30'}


def dossier(value):
    return {'schema': 're.bow.dossier-additions-r30/1', 'canonicalFidelityCertified': False,
        'identities': [{'id': row['id'], 'nativeEntries': row['r30_native_entries'],
            'nativeFrames': row['r30_native_rendered_frames'], 'combatStances': row['r30_native_combat_stances'],
            'smoothSequences': row['r30_native_smooth_sequences'], 'verifiedIdleCycles': row['r30_native_verified_idle_cycles'],
            'optionalObservedVariants': row['r30_native_optional_neutral_variants'],
            'optionalPlayableVariants': row['r30_native_playable_optional_neutral_variants'], 'contextOnly': row['r30_context_only'],
            'sourceEditions': row['r30_native_source_editions'], 'avatarImage': row['r30_avatar_image'],
            'avatarSourceKind': row['r30_avatar_source_kind']} for row in value['identity_updates']]}


def export(value, root=ROOT, destination=None, progress_path=None):
    root = Path(root).resolve(); validate(value, root)
    baseline = root / BASELINE
    destination = UTIL.checked_path(root, destination or root / 'docs/BOW_FIDELITE_R30.xlsx')
    progress_path = UTIL.checked_path(root, progress_path or root / MODULE / 'progress.json')
    if destination in (baseline, progress_path) or progress_path == baseline:
        raise ValueError('Outputs cannot overwrite historical inputs')
    if destination.exists():
        book = load_workbook(destination, read_only=True); creator = book.properties.creator; book.close()
        if creator != 'BOW fidelity progress R30':
            raise ValueError('Refusing to overwrite another author’s workbook')
    book = load_workbook(baseline)
    if book.properties.creator != 'BOW fidelity progress R29' or any(name in book.sheetnames for name in SHEETS):
        raise ValueError('Expected the preserved R29 workbook')
    original = list(book.sheetnames); old_catalogue = list(book['Catalogue R29'].values)
    if len(old_catalogue) != 390 or {row[0] for row in old_catalogue[1:]} != {row['id'] for row in value['identity_updates']}:
        raise ValueError('Historical catalogue must preserve all 389 identities')
    book.properties.creator = 'BOW fidelity progress R30'; book.properties.title = 'B.O.W. Protocol — sprites natifs, avatars et mouvements R30'
    UTIL.sheet(book, SHEETS[0], ['Indicateur', 'Valeur', 'Portée'], [[row['label'], row['value'], row['scope']] for row in value['statistics']]
        + [['Recette R30 complète', value['validation_complete'], 'Résultats manquants ou périmés explicitement exclus.'], ['Fidélité globale 1:1 certifiée', 0, 'Critères ouverts.']])
    UTIL.sheet(book, SHEETS[1], ['ID', 'Chantier', 'Statut', 'Acquis R30', 'Critères ouverts', 'Preuves'],
        [[row['id'], row['title'], row['status'], '\n'.join(row['advances']), '\n'.join(row['remaining']), row['evidence']] for row in value['tasks']])
    updates = {row['id']: row for row in value['identity_updates']}
    UTIL.sheet(book, SHEETS[2], list(old_catalogue[0]) + list(FIELDS),
        [list(row) + [updates[row[0]][key] for key in FIELDS] for row in old_catalogue[1:]])
    UTIL.sheet(book, SHEETS[3], ['UID', 'Entrée', 'Type', 'Édition source', 'Édition cible admise', 'Images livrées',
        'Échantillonnage images/s', 'Durée secondes', 'Durée native complète', 'Geste', 'Geste vérifié', 'Continuité vérifiée',
        'Pose combattante admise', 'Idle animé complet admis', 'Boucle neutre source observée',
        'Variante optionnelle jouable sur acteur existant', 'Contexte uniquement', 'Famille de variante',
        'Geste original certifié', 'Côté', 'Lacunes d’apparence', 'Manifeste'], value['native_rows'])
    UTIL.sheet(book, SHEETS[4], ['UID', 'Nom', 'Jeu cible', 'Image', 'Nature de source', 'Édition source', 'Publication', 'SHA-256', 'Fidélité globale certifiée'], value['avatar_rows'])
    UTIL.sheet(book, SHEETS[5], ['Contrôle', 'Résultat', 'Portée', 'Rapport', 'SHA-256'],
        [[row['label'], row['result'], row['scope'], row['path'], row['sha256']] for row in value['validation']]
        + [['Rapport manquant ou non courant', 'NON VALIDÉ', 'Résultat exclu du cumul R30', path, ''] for path in value['validation_gaps']])
    UTIL.sheet(book, SHEETS[6], ['Chantier', 'UID', 'Nature', 'Fichier', 'Portée', 'SHA-256'], value['source_rows'])
    book.active = book.sheetnames.index('Chantiers R30'); destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=destination.parent, suffix='.xlsx', delete=False) as handle:
        temporary = Path(handle.name)
    try:
        book.save(temporary); book.close(); OLD.canonical_zip(baseline, temporary)
        with ZipFile(baseline) as before, ZipFile(temporary) as after:
            if any(before.read(name) != after.read(name) for name in before.namelist() if name.startswith('xl/worksheets/')):
                raise ValueError('Historical worksheet bytes or relationships changed')
        check = load_workbook(temporary)
        try:
            if check.sheetnames[:len(original)] != original or check['Catalogue R30'].max_row != 390 or check['Chantiers R30'].max_row != 18:
                raise ValueError('New sheets omitted a historical identity or workstream')
            if any(cell.data_type == 'f' for title in SHEETS for row in check[title] for cell in row):
                raise ValueError('New evidence sheets cannot contain executable formulas')
        finally:
            check.close()
        os.replace(temporary, destination); write(progress_path, progress(value))
    finally:
        temporary.unlink(missing_ok=True)
    return {'workbook': destination.relative_to(root).as_posix(), 'sha256': sha(destination),
        'baselineSha256': sha(baseline), 'historical_worksheets_preserved': len(original),
        'worksheets': len(original) + len(SHEETS), 'identities': 389, 'open_workstreams': 17,
        'native_rows': len(value['native_rows']), 'avatar_rows': len(value['avatar_rows']),
        'source_rows': len(value['source_rows']), 'canonical_fidelity_certified': 0}


def overview(value, checks):
    return f'''# R30 — Sprites natifs, avatars et mouvements complets

{value['summary']}

Le [classeur R30](BOW_FIDELITE_R30.xlsx) ajoute sept feuilles aux 49 feuilles R29, dont les XML et relations restent identiques. Le [suivi central](../modules/fidelity-r30/index.html) expose les 17 chantiers ouverts et le catalogue complet de 389 UID.

Les poses natives rendues, les postures de combat admissibles et les vrais cycles idle sont comptés séparément. Une séquence native entièrement échantillonnée à 30 images/s au moins peut être fluide tout en gardant un geste non attribué ou une boucle non vérifiée. Les variantes d’une autre édition restent hors de l’admission cible. Les portraits de sélection ne certifient pas de géométrie, d’animation ou de corps complet.

Les 363 images générées R29, les 22 refus historiques et le refus de révision Tyrant restent conservés. Aucun refus ne devient un nouvel essai de génération par l’ajout de références natives. Aucun chantier n’est déclaré clos ni aucune fidélité globale 1:1 certifiée.

La recette R30 est {'complète' if checks['complete'] else 'provisoire, avec rapports manquants ou non courants explicités'} dans son périmètre : {json.dumps(checks['totals'], ensure_ascii=False)}. Les recettes historiques ne sont pas recomptées comme de nouvelles exécutions.

Preuves : [suivi structuré](R30_TRACKING.json), [recette](R30_VALIDATION.json), [export](R30_EXPORT.json), [reproduction](R30_REPRODUCTION.json).

```sh
python3 scripts/fidelity-progress-r30.py build
python3 scripts/fidelity-progress-r30.py check
python3 scripts/fidelity-progress-r30.py export
python3 scripts/fidelity-progress-r30.py reproduce
```
'''.encode()


def reproduce(value, root=ROOT):
    root = Path(root).resolve(); checks = []
    try:
        for seed in ('17', '73'):
            command = [sys.executable, 'scripts/fidelity-progress-r30.py', 'check']
            if not value['validation_complete']:
                command.append('--allow-pending')
            result = subprocess.run(command, cwd=root, env={**os.environ, 'PYTHONHASHSEED': seed}, capture_output=True, text=True)
            if result.returncode:
                raise ValueError(result.stderr)
        checks.append({'name': 'Tracking is stable across independent Python hash seeds', 'pass': True})
        folder = root / 'qa-r30-reproduction'; folder.mkdir(exist_ok=True); baseline = sha(root / BASELINE)
        with tempfile.TemporaryDirectory(dir=folder) as name:
            temporary = Path(name); result = export(value, root, temporary / 'reproduced.xlsx', temporary / 'progress.json')
            if result['sha256'] != sha(root / 'docs/BOW_FIDELITE_R30.xlsx'):
                raise ValueError('Actual R30 workbook bytes do not reproduce')
            checks.append({'name': 'Actual R30 workbook reproduces exactly', 'pass': True, 'sha256': result['sha256']})
            if (temporary / 'progress.json').read_bytes() != (root / MODULE / 'progress.json').read_bytes() or sha(root / BASELINE) != baseline:
                raise ValueError('Progress or preserved R29 baseline changed')
            checks.append({'name': 'Progress matches and R29 baseline remains byte-exact', 'pass': True})
    except Exception as error:
        checks.append({'name': 'Reproduction failure', 'pass': False, 'error': str(error)})
    return {'schema': 're.bow.fidelity-reproduction-r30/1', 'passed': sum(row['pass'] for row in checks),
            'failed': sum(not row['pass'] for row in checks), 'checks': checks}


def main():
    parser = argparse.ArgumentParser(description=__doc__); parser.add_argument('action', choices=('build', 'check', 'export', 'reproduce'))
    parser.add_argument('--allow-pending', action='store_true'); args = parser.parse_args()
    if args.action in ('build', 'check'):
        checks, value = build(allow_pending=args.allow_pending)
        outputs = {'docs/R30_VALIDATION.json': raw(checks), 'docs/R30_TRACKING.json': raw(value),
            'docs/R30_WORKSTREAMS.json': raw({key: value[key] for key in ('release', 'summary', 'tasks', 'generation_preference')}),
            (MODULE / 'progress.json').as_posix(): raw(progress(value)),
            (MODULE / 'dossier-additions.json').as_posix(): raw(dossier(value)),
            'docs/FIDELITE_R30.md': overview(value, checks)}
        for path, content in outputs.items():
            if args.action == 'check':
                if (ROOT / path).read_bytes() != content:
                    raise ValueError('R30 output differs from its evidence: ' + path)
            else:
                write(ROOT / path, content)
        print(json.dumps({'complete': checks['complete'], 'totals': checks['totals'], 'nativeRows': len(value['native_rows']), 'avatars': len(value['avatar_rows'])}))
    elif args.action == 'export':
        result = export(read('docs/R30_TRACKING.json')); write(ROOT / 'docs/R30_EXPORT.json', result); print(json.dumps(result))
    else:
        result = reproduce(read('docs/R30_TRACKING.json')); write(ROOT / 'docs/R30_REPRODUCTION.json', result); print(json.dumps(result)); return 1 if result['failed'] else 0
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
