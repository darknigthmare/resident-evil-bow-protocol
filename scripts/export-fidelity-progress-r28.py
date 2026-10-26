#!/usr/bin/env python3
"""Append R28 evidence while preserving all R27 worksheet XML and relationships."""
import argparse
from collections import Counter
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import re
import tempfile
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('r27_export_helpers', ROOT / 'scripts/export-fidelity-progress-r27.py')
OLD = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(OLD)
CREATOR = 'BOW fidelity progress R28'
SHEETS = ('Synthèse R28', 'Chantiers R28', 'Catalogue R28', 'Sources R28', 'Recette R28', 'Évolutions R28', 'Corps sources R28')
FIELDS = ('visual', 'sprite_review_status', 'sprite_full_body_reviewed', 'sprite_reference_full_body_verified',
          'sprite_scope', 'sprite_facing', 'sprite_pose_kind', 'sprite_available_states', 'sprite_combat_eligible',
          'sprite_generation_status', 'sprite_generation_reason', 'r28_replaced', 'r28_family_id', 'r28_family_base',
          'r28_evolution_targets', 'r28_held_targets', 'r28_context_only', 'r28_source_models', 'r28_source_views',
          'r28_body_jointly_qualified', 'r28_equipment_complete', 'r28_known_appearance_gaps',
          'r28_canonical_claims', 'r28_documentary_reviewed', 'r28_audio_exact_clips', 'r28_audio_candidate_clips', 'limitations')


def read(path):
    return json.loads(Path(path).read_text())


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def evidence_link(root, href):
    if not isinstance(href, str) or not href or '\\' in href or re.match(r'^[a-zA-Z][a-zA-Z0-9+.-]*:', href) or href.startswith('/'):
        raise ValueError('Evidence links must remain inside the checkout')
    return OLD.checked_path(root, Path('modules/fidelity-r28') / href, exists=True)


def validate(tracking, root=ROOT):
    root = Path(root).resolve()
    if tracking.get('schema') != 're.bow.parallel-workstreams/3' or tracking.get('release') != 'R28' or tracking.get('canonical_fidelity_certified') != 0:
        raise ValueError('Unsupported R28 tracking schema or certification')
    expected = {item['id'] for item in read(root / 'docs/R27_TRACKING.json')['tasks']}
    tasks = tracking.get('tasks', [])
    if len(expected) != 17 or len(tasks) != 17 or {item['id'] for item in tasks} != expected:
        raise ValueError('All seventeen open workstreams must remain explicit')
    for task in tasks:
        if task.get('status') != 'EN COURS' or not all(task.get(key) for key in ('title', 'advances', 'remaining', 'evidence')):
            raise ValueError('Open workstream requires progress, remaining criteria and evidence')
        for proof in task['evidence']:
            evidence_link(root, proof['href'])
    for item in tracking.get('resources', []):
        evidence_link(root, item['href'])
    inputs = tracking.get('inputs', [])
    if not inputs or len({item['path'] for item in inputs}) != len(inputs):
        raise ValueError('Evidence inputs must be nonempty and unique')
    for item in inputs:
        path = OLD.local_file(root, item['path'])
        if not re.fullmatch('[0-9a-f]{64}', str(item.get('sha256', ''))) or digest(path) != item['sha256']:
            raise ValueError('Evidence hash mismatch: ' + item['path'])
    roster = read(root / 'data/roster.json')['catalog']; ids = {item['id'] for item in roster}
    updates = tracking.get('identity_updates', [])
    if len(roster) != 389 or len(ids) != 389 or len(updates) != 389 or {item['id'] for item in updates} != ids:
        raise ValueError('All 389 identities must have unique R28 progress')
    if any(any(key not in update for key in FIELDS) for update in updates):
        raise ValueError('R28 identity scope is incomplete')
    for row in tracking.get('source_rows', []):
        if not isinstance(row, list) or len(row) != 6:
            raise ValueError('Evidence requires exactly six columns')
        for value in row:
            OLD.safe_cell(value)
    for item in tracking.get('validation', []):
        OLD.local_file(root, item['path'])
    if tracking.get('sprite_manifest'):
        validate_sprites(tracking, root)
    by_uid = {item['id']: item for item in updates}
    if tracking.get('evolution_graph'):
        graph_path = OLD.local_file(root, tracking['evolution_graph'])
        if not any(item['path'] == tracking['evolution_graph'] for item in inputs):
            raise ValueError('Evolution graph requires an input binding')
        graph = read(graph_path)
        families = {uid: family for family in graph['families'] for uid in family['forms']}
        for uid, update in by_uid.items():
            family = families.get(uid)
            expected_fields = {'r28_family_id': family['id'] if family else '',
                'r28_family_base': family['baseUid'] if family else '', 'r28_context_only': uid in graph['contextOnlyUids'],
                'r28_evolution_targets': [edge['to'] for edge in graph['edges'] if edge['from'] == uid and edge['playable']],
                'r28_held_targets': [edge['to'] for edge in graph['edges'] if edge['from'] == uid and not edge['playable']]}
            if any(update[key] != value for key, value in expected_fields.items()):
                raise ValueError('Evolution family or encounter scope differs from the source graph')
    if tracking.get('canonical_manifest'):
        canonical_path = OLD.local_file(root, tracking['canonical_manifest'])
        if not any(item['path'] == tracking['canonical_manifest'] for item in inputs):
            raise ValueError('Documentary manifest requires an input binding')
        canonical = read(canonical_path)
        reviewed = {entry['id'] for entry in canonical['entries'] if entry.get('new_claim_ids') or entry.get('previous_reviewed_claim_count', 0)}
        for uid, update in by_uid.items():
            if (update['r28_documentary_reviewed'] != (uid in reviewed)
                    or update['r28_canonical_claims'] != sum(claim['uid'] == uid for claim in canonical['claims'])):
                raise ValueError('Documentary coverage differs from the actual assertions')
    if tracking.get('audio_manifest'):
        audio_path = OLD.local_file(root, tracking['audio_manifest'])
        if not any(item['path'] == tracking['audio_manifest'] for item in inputs):
            raise ValueError('Audio manifest requires an input binding')
        audio = read(audio_path)
        exact = Counter(entry['identity_id'] for entry in audio['entries'] if entry.get('identity_id'))
        candidates = Counter(uid for entry in audio['entries'] for uid in entry.get('candidate_identity_ids', []))
        for uid, update in by_uid.items():
            if update['r28_audio_exact_clips'] != exact[uid] or update['r28_audio_candidate_clips'] != candidates[uid]:
                raise ValueError('Audio candidates cannot become exact subject assignments')
    return tracking


def validate_sprites(tracking, root):
    path = OLD.local_file(root, tracking['sprite_manifest'])
    if not any(item['path'] == tracking['sprite_manifest'] for item in tracking['inputs']):
        raise ValueError('Current sprite manifest requires an input binding')
    pack = read(path); sprites = pack['sprites']; pending = pack['pending']; counts = pack['counts']
    if pack.get('schema') != 're.bow.generated-combat-sprites/1':
        raise ValueError('Unexpected sprite manifest')
    jobs = read(root / 'modules/fidelity-r27/sprites/jobs.json')['jobs']
    target_ids = {item['id'] for item in jobs}; sprite_ids = {item['id'] for item in sprites}; pending_ids = {item['id'] for item in pending}
    if len(jobs) != 385 or len(sprite_ids) != len(sprites) or len(pending_ids) != len(pending) or sprite_ids & pending_ids or sprite_ids | pending_ids != target_ids:
        raise ValueError('The same 385 target identities must be resolved exactly once')
    eligible = lambda item: item.get('reviewStatus') == 'visually-reviewed' and item.get('fullBodyReviewed') is True and item.get('poseKind') == 'neutral-base' and item.get('scope') == 'combatant-neutral-pose' and item.get('sourceFacing') in ('left', 'right') and item.get('availableStates') == ['idle']
    computed = {'targetIdentities': 385, 'generatedSprites': len(sprites), 'pendingSprites': len(pending),
        'visuallyReviewedSprites': sum(item.get('reviewStatus') == 'visually-reviewed' for item in sprites),
        'fullBodyReviewedSprites': sum(item.get('fullBodyReviewed') is True for item in sprites),
        'neutralBaseCombatEligibleSprites': sum(eligible(item) for item in sprites),
        'generationBlockedSprites': sum(item.get('generationStatus') == 'BLOCKED_BY_TOOL' for item in pending),
        'originalGameSprites': 0, 'completeCombatAnimationSets': 0, 'canonicalFidelityCertified': 0}
    if any(type(counts.get(key)) is not int or counts[key] != value for key, value in computed.items()):
        raise ValueError('R28 counts must match actual files, reviews and refusals')
    updates = {item['id']: item for item in tracking['identity_updates']}
    for sprite in sprites:
        image = OLD.local_file(root, sprite['image'])
        if image.stat().st_size != sprite['bytes'] or digest(image) != sprite['sha256']:
            raise ValueError('Active sprite bytes differ from their manifest')
        if sprite.get('canMirror') is not False or sprite.get('canonicalFidelityCertified') is not False or sprite.get('combatIntegrated') is not False:
            raise ValueError('Unsupported art or combat certification')
        update = updates[sprite['id']]
        expected = {'visual': sprite['image'], 'sprite_review_status': sprite['reviewStatus'],
            'sprite_full_body_reviewed': sprite['fullBodyReviewed'], 'sprite_reference_full_body_verified': sprite['completeBodyReferenceVerified'],
            'sprite_scope': sprite['scope'], 'sprite_facing': sprite['sourceFacing'], 'sprite_pose_kind': sprite['poseKind'],
            'sprite_available_states': sprite['availableStates'], 'sprite_combat_eligible': eligible(sprite), 'sprite_generation_status': 'GENERATED',
            'r28_equipment_complete': sprite.get('equipmentCompletenessReviewed') is True,
            'r28_known_appearance_gaps': sprite.get('knownAppearanceGaps', [])}
        if any(update[key] != value for key, value in expected.items()):
            raise ValueError('Identity progress does not match active sprite qualification')
    for item in pending:
        update = updates[item['id']]
        if (update.get('sprite_generated') is not False or update['sprite_combat_eligible'] is not False
                or update['sprite_full_body_reviewed'] is not False or update['sprite_available_states'] != []
                or update['sprite_generation_status'] != item['generationStatus']
                or update['sprite_generation_reason'] != item.get('reason', '')):
            raise ValueError('A refusal cannot become an image, reviewed body or combat state')


def canonical_zip(baseline, generated):
    """Keep original worksheet parts byte-exact and make new ZIP metadata stable."""
    with ZipFile(baseline) as source, ZipFile(generated) as current:
        historical = {name: (source.getinfo(name), source.read(name)) for name in source.namelist() if name.startswith('xl/worksheets/')}
        if not set(historical) <= set(current.namelist()):
            raise ValueError('Historical worksheet part lost')
        with tempfile.NamedTemporaryFile(dir=generated.parent, suffix='.xlsx', delete=False) as handle:
            replacement = Path(handle.name)
        try:
            with ZipFile(replacement, 'w') as output:
                for entry in current.infolist():
                    if entry.filename in historical:
                        info, content = historical[entry.filename]
                    else:
                        info, content = entry, current.read(entry.filename)
                        info.date_time = (2026, 10, 26, 0, 0, 0)
                        if entry.filename == 'docProps/core.xml':
                            content = re.sub(rb'(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)', rb'\g<1>2026-10-26T00:00:00Z\g<2>', content)
                    output.writestr(info, content)
            os.replace(replacement, generated)
        finally:
            replacement.unlink(missing_ok=True)


def export(tracking_path, baseline_path, output_path, progress_path, root=ROOT):
    root = Path(root).resolve()
    tracking_path, baseline_path = (OLD.checked_path(root, value, exists=True) for value in (tracking_path, baseline_path))
    output_path, progress_path = (OLD.checked_path(root, value) for value in (output_path, progress_path))
    if len({tracking_path, baseline_path, output_path, progress_path}) != 4:
        raise ValueError('Outputs must not overwrite historical evidence')
    tracking = validate(read(tracking_path), root)
    if not any((root / item['path']).resolve() == baseline_path for item in tracking['inputs']):
        raise ValueError('Historical workbook requires a recorded evidence hash')
    if output_path.exists():
        old = load_workbook(output_path, read_only=True); creator = old.properties.creator; old.close()
        if creator != CREATOR:
            raise ValueError('Refusing to overwrite a workbook from another author')
    book = load_workbook(baseline_path)
    if book.properties.creator != 'BOW fidelity progress R27' or any(title in book.sheetnames for title in SHEETS):
        raise ValueError('Expected the preserved R27 baseline workbook')
    original_names = list(book.sheetnames)
    rows = list(book['Catalogue R27'].values)
    ids = {item['id'] for item in read(root / 'data/roster.json')['catalog']}
    if len(rows) != 390 or {row[0] for row in rows[1:]} != ids:
        raise ValueError('Historical catalogue must retain 389 identities')
    book.properties.creator = CREATOR; book.properties.title = 'B.O.W. Protocol — chantiers et évolutions R28'
    OLD.sheet(book, SHEETS[0], ['Indicateur', 'Valeur', 'Portée'], [[item['label'], item['value'], item.get('scope', '')] for item in tracking['statistics']] + [['Fidélité globale 1:1 certifiée', 0, 'Critères non vérifiés maintenus ouverts.']])
    OLD.sheet(book, SHEETS[1], ['ID', 'Chantier', 'Statut R28', 'Éléments R28 vérifiés', 'Critères ouverts', 'Preuves', 'Nouvelle implémentation R28'],
              [[task['id'], task['title'], task['status'], '\n'.join(task['advances']), '\n'.join(task['remaining']), '\n'.join(item['label'] + ': ' + item['href'] for item in task['evidence']), task.get('r28_new_implementation', False)] for task in tracking['tasks']])
    updates = {item['id']: item for item in tracking['identity_updates']}
    headers = ['Pose active R28', 'Revue R28', 'Corps anatomique entier revu', 'Référence complète vérifiée', 'Portée R28', 'Orientation observée', 'Type de pose', 'États disponibles', 'Admissible au seul idle', 'Statut génération', 'Refus ou attente', 'Version remplacée en R28', 'Famille R28', 'UID de base', 'Évolutions jouables directes', 'Transitions de rencontre retenues', 'Rencontre uniquement en R28', 'Modèles sources R28', 'Vues sources R28', 'Source et corps conjointement qualifiés', 'Équipement complet certifié par revue', 'Lacunes d’apparence connues', 'Nouvelles citations documentaires R28', 'Revue documentaire cumulée', 'Clips du groupe audio publié exact R28', 'Clips candidats liés sans attribution certaine', 'Limites R28']
    OLD.sheet(book, SHEETS[2], list(rows[0]) + headers, [list(row) + [updates[row[0]][key] for key in FIELDS] for row in rows[1:]])
    OLD.sheet(book, SHEETS[3], ['Chantier', 'UID ou portée', 'Nature', 'Source ou fichier', 'Limite', 'Citation ou SHA'], tracking['source_rows'])
    OLD.sheet(book, SHEETS[4], ['Contrôle', 'Résultat', 'Portée', 'Rapport', 'SHA-256'], [[item['label'], item['result'], item['scope'], item['path'], item['sha256']] for item in tracking['validation']])
    graph = read(OLD.local_file(root, tracking['evolution_graph'])) if tracking.get('evolution_graph') else {'edges': []}
    OLD.sheet(book, SHEETS[5], ['Lien', 'Famille', 'UID départ', 'UID arrivée', 'Jouable', 'Conditions originales R28', 'Citations sources', 'Limite'],
              [[edge['id'], edge['familyId'], edge['from'], edge['to'], edge['playable'], edge['conditions'], edge['sourceIds'], edge.get('boundary', '')] for edge in graph['edges']])
    body = read(OLD.local_file(root, tracking['body_manifest'])) if tracking.get('body_manifest') else {'entries': []}
    OLD.sheet(book, SHEETS[6], ['UID', 'Modèle source', 'Édition', 'Publication', 'Vues', 'Portée de pose', 'Équipement et phase', 'Image originale certifiée'],
              [[entry['id'], entry['modelId'], entry.get('sourceEdition', ''), entry['sourcePage'], len(entry['views']), entry['sourcePoseScope'], entry['equipmentScope'], False] for entry in body['entries']])
    book.active = book.sheetnames.index('Chantiers R28')
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=output_path.parent, suffix='.xlsx', delete=False) as handle:
        temporary = Path(handle.name)
    try:
        book.save(temporary); book.close(); canonical_zip(baseline_path, temporary)
        with ZipFile(baseline_path) as original, ZipFile(temporary) as result:
            if any(original.read(name) != result.read(name) for name in original.namelist() if name.startswith('xl/worksheets/')):
                raise ValueError('Historical worksheet bytes changed')
        check = load_workbook(temporary)
        try:
            if check.sheetnames[:len(original_names)] != original_names or check['Catalogue R28'].max_row != 390 or check['Chantiers R28'].max_row != 18:
                raise ValueError('Exported identity or workstream count changed')
            if any(cell.data_type == 'f' for title in SHEETS for row in check[title] for cell in row):
                raise ValueError('New spreadsheet formula rejected')
        finally:
            check.close()
        os.replace(temporary, output_path)
        progress = {key: tracking[key] for key in ('summary', 'statistics', 'resources', 'tasks')}
        progress['schema'] = 're.bow.fidelity-progress/1'; progress['release'] = 'R28'
        OLD.atomic_write(progress_path, (json.dumps(progress, ensure_ascii=False, indent=2) + '\n').encode())
    finally:
        temporary.unlink(missing_ok=True)
    return {'workbook': output_path.relative_to(root).as_posix(), 'sha256': digest(output_path),
            'baseline': {'path': baseline_path.relative_to(root).as_posix(), 'sha256': digest(baseline_path)},
            'tracking': {'path': tracking_path.relative_to(root).as_posix(), 'sha256': digest(tracking_path)},
            'progress': {'path': progress_path.relative_to(root).as_posix(), 'sha256': digest(progress_path)},
            'historical_worksheets_preserved': len(original_names), 'new_worksheets': len(SHEETS),
            'total_worksheets': len(original_names) + len(SHEETS), 'identities': 389,
            'open_workstreams': 17, 'source_rows': len(tracking['source_rows']), 'canonical_fidelity_certified': 0}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--tracking', type=Path, default=ROOT / 'docs/R28_TRACKING.json')
    parser.add_argument('--baseline', type=Path, default=ROOT / 'docs/BOW_FIDELITE_R27.xlsx')
    parser.add_argument('--output', type=Path, default=ROOT / 'docs/BOW_FIDELITE_R28.xlsx')
    parser.add_argument('--progress', type=Path, default=ROOT / 'modules/fidelity-r28/progress.json')
    args = parser.parse_args()
    result = export(args.tracking, args.baseline, args.output, args.progress)
    OLD.atomic_write(ROOT / 'docs/R28_EXPORT.json', (json.dumps(result, ensure_ascii=False, indent=2) + '\n').encode())
    print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    main()
