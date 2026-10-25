#!/usr/bin/env python3
"""Export R27 evidence while preserving every historical R26 worksheet byte."""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import math
import os
from pathlib import Path
import re
import tempfile
from urllib.parse import unquote, urlsplit
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('source_tables_r27', ROOT / 'scripts/export-source-evidence-r21.py')
TABLES = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(TABLES)
CREATOR = 'BOW fidelity progress R27'
BASE_CREATOR = 'BOW fidelity progress R26'
ADDED_SHEETS = ('Synthèse R27', 'Chantiers R27', 'Catalogue R27', 'Sources R27', 'Recette R27', 'Clips sources R27')
SPRITE_KEYS = ('sprite_generated', 'sprite_review_status', 'sprite_full_body_reviewed',
               'sprite_reference_full_body_verified', 'sprite_scope', 'sprite_facing', 'sprite_pose_kind',
               'sprite_available_states', 'sprite_combat_eligible', 'sprite_combat_integrated',
               'sprite_generation_status', 'sprite_generation_reason')
UPDATE_KEYS = {'id', 'visual', 'rig', 'motion_clips', 'audio', 'claims', 'limitations', *SPRITE_KEYS}
POSE_LABEL = 'POSE FIXE APRÈS ZÉRO'
MOTION_LABEL = 'MOUVEMENT APRÈS ZÉRO'


def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def checked_path(root, value, exists=False):
    root = Path(root).resolve()
    candidate = Path(value)
    candidate = candidate if candidate.is_absolute() else root / candidate
    try:
        path = candidate.resolve(strict=False)
    except (FileNotFoundError, OSError, RuntimeError) as error:
        raise ValueError('Evidence must be an existing local file') from error
    if not path.is_relative_to(root) or path == root:
        raise ValueError('Paths must stay inside the checkout')
    if exists and not path.is_file():
        raise ValueError('Evidence must be a local file')
    return path


def local_file(root, value):
    if not isinstance(value, str) or not value or '\\' in value or Path(value).is_absolute():
        raise ValueError('Invalid evidence path inside the checkout')
    if any(part in ('.', '..') for part in Path(value).parts):
        raise ValueError('Evidence path must stay inside the checkout')
    return checked_path(root, value, exists=True)


def evidence_link(root, href):
    if not isinstance(href, str) or not href or '\\' in href:
        raise ValueError('Invalid evidence link')
    url = urlsplit(href)
    if url.scheme or url.netloc or not url.path or url.path.startswith('/'):
        raise ValueError('Evidence links must stay inside the checkout')
    target = checked_path(root, Path('modules/fidelity-r27') / unquote(url.path), exists=True)
    if target.relative_to(Path(root).resolve()).parts[0] not in ('modules', 'docs'):
        raise ValueError('Evidence links must target reviewed modules or docs')
    return target


def safe_cell(value):
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, str):
        if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value):
            raise ValueError('Unsafe spreadsheet control character')
        if value.lstrip(' \t\r\n').startswith(('=', '+', '-', '@')):
            return "'" + value
    elif value is not None and not isinstance(value, (bool, int, float)):
        raise ValueError('Unsupported spreadsheet cell value')
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError('Non-finite spreadsheet number')
    return value


def sheet(book, name, headers, rows):
    return TABLES.sheet(book, name, [safe_cell(value) for value in headers],
                        [[safe_cell(value) for value in row] for row in rows])


def validate(tracking, root=ROOT):
    if tracking.get('schema') != 're.bow.parallel-workstreams/2' or type(tracking.get('canonical_fidelity_certified')) is not int or tracking['canonical_fidelity_certified'] != 0:
        raise ValueError('Unexpected tracking schema or unsupported global fidelity certification')
    expected = {task['id'] for task in read(root / 'docs/R20_TRACKING.json')['tasks'] if task['status'] != 'VALIDÉ'}
    tasks = tracking.get('tasks', [])
    if len(expected) != 17 or len(tasks) != 17 or {task['id'] for task in tasks} != expected:
        raise ValueError('All seventeen historical open workstreams must be present exactly once')
    for task in tasks:
        if task.get('status') != 'EN COURS' or task.get('category') not in ('sources', 'functions', 'validation'):
            raise ValueError('No workstream may be silently closed or omitted')
        if not task.get('title') or not task.get('advances') or not task.get('remaining') or not task.get('evidence'):
            raise ValueError('Each workstream requires an advance, an open criterion and evidence')
        if any(not isinstance(item, str) or not item for item in task['advances'] + task['remaining']):
            raise ValueError('Workstream advances and criteria must be text')
        for item in task['evidence']:
            evidence_link(root, item['href'])
    for resource in tracking.get('resources', []):
        evidence_link(root, resource['href'])
    inputs = tracking.get('inputs', [])
    if not inputs or len({item['path'] for item in inputs}) != len(inputs):
        raise ValueError('Evidence inputs must be present and unique')
    for evidence in inputs:
        path = local_file(root, evidence['path'])
        if not re.fullmatch(r'[0-9a-f]{64}', str(evidence.get('sha256', ''))) or digest(path) != evidence['sha256']:
            raise ValueError('Evidence hash mismatch: ' + evidence['path'])
    roster = read(root / 'data/roster.json')['catalog']
    ids = {item['id'] for item in roster}
    updates = tracking.get('identity_updates', [])
    if len(roster) != 389 or len(ids) != 389 or len(updates) != 389 or {item['id'] for item in updates} != ids:
        raise ValueError('Exactly all 389 catalogue identities require progress; unknown or duplicate identity')
    if any(set(item) - UPDATE_KEYS for item in updates):
        raise ValueError('Unknown identity progress field')
    for row in tracking.get('source_rows', []):
        if not isinstance(row, list) or len(row) != 6:
            raise ValueError('Source evidence rows require exactly six columns')
        for value in row:
            safe_cell(value)
    for item in tracking.get('validation', []):
        if not all(key in item for key in ('label', 'result', 'scope', 'path')):
            raise ValueError('Validation requires its result, scope and local evidence')
        local_file(root, item['path'])
    validate_sprites(tracking, root)
    return tracking


def validate_sprites(tracking, root):
    """Do not turn a queued job or a portrait into a reviewed combat body."""
    if not tracking.get('sprite_manifest'):
        return
    path = local_file(root, tracking['sprite_manifest'])
    if not any(local_file(root, item['path']) == path for item in tracking['inputs']):
        raise ValueError('Sprite manifest requires a recorded evidence hash')
    pack = read(path)
    if pack.get('schema') != 're.bow.generated-combat-sprites/1':
        raise ValueError('Unexpected generated sprite schema')
    roster = {item['id'] for item in read(root / 'data/roster.json')['catalog']}
    historical = {item['id'] for item in read(root / 'assets/sprites/bow/manifest.json')['sprites']}
    if len(historical) != 4 or not historical <= roster:
        raise ValueError('Expected the four historical R20 sprite identities')
    expected = roster - historical
    jobs_path = local_file(root, 'modules/fidelity-r27/sprites/jobs.json')
    if not any(local_file(root, item['path']) == jobs_path for item in tracking['inputs']):
        raise ValueError('Sprite jobs require a recorded evidence hash')
    jobs_pack = read(jobs_path)
    jobs = {item['id']: item for item in jobs_pack.get('jobs', [])}
    if (jobs_pack.get('schema') != 're.bow.sprite-generation-jobs/1'
            or len(jobs) != len(jobs_pack.get('jobs', [])) or set(jobs) != expected):
        raise ValueError('Sprite jobs must cover the same 385 target identities')
    sprites, pending = pack.get('sprites', []), pack.get('pending', [])
    actual_ids, pending_ids = [item['id'] for item in sprites], [item['id'] for item in pending]
    if (len(set(actual_ids)) != len(actual_ids) or len(set(pending_ids)) != len(pending_ids)
            or set(actual_ids) & set(pending_ids) or set(actual_ids) | set(pending_ids) != expected):
        raise ValueError('Exactly the 385 missing identities require a generated asset or an explicit pending state')
    counts = pack['counts']
    expected_counts = {'targetIdentities': 385, 'generatedSprites': len(sprites), 'pendingSprites': len(pending),
                       'visuallyReviewedSprites': sum(item.get('reviewStatus') == 'visually-reviewed' for item in sprites),
                       'fullBodyReviewedSprites': sum(item.get('fullBodyReviewed') is True for item in sprites),
                       'originalGameSprites': 0, 'completeCombatAnimationSets': 0, 'canonicalFidelityCertified': 0}
    for pose, key in (('neutral-base', 'neutralBaseSprites'), ('source-bind-pose', 'sourceBindPoseSprites'),
                      ('source-presentation', 'sourcePresentationSprites'), ('unknown', 'unknownPoseSprites')):
        expected_counts[key] = sum(item.get('poseKind') == pose for item in sprites)
    expected_counts['neutralBaseCombatEligibleSprites'] = sum(
        item.get('reviewStatus') == 'visually-reviewed' and item.get('fullBodyReviewed') is True
        and item.get('poseKind') == 'neutral-base' and item.get('scope') == 'combatant-neutral-pose'
        and item.get('sourceFacing') in ('left', 'right') and item.get('availableStates') == ['idle'] for item in sprites)
    if any(type(counts.get(key)) is not int or counts[key] != value for key, value in expected_counts.items()):
        raise ValueError('Generated, pending, reviewed and complete-cycle sprite counts disagree')
    updates = {item['id']: item for item in tracking['identity_updates']}
    actual = {item['id']: item for item in sprites}
    awaiting = {item['id']: item for item in pending}
    for uid in roster:
        update = updates[uid]
        if any(key not in update for key in SPRITE_KEYS):
            raise ValueError('All identities require explicit R27 sprite scope')
        for key in ('sprite_generated', 'sprite_full_body_reviewed', 'sprite_reference_full_body_verified',
                    'sprite_combat_eligible', 'sprite_combat_integrated'):
            if type(update[key]) is not bool:
                raise ValueError('Sprite progress flags must be booleans')
        if uid not in actual:
            if update['sprite_generated'] or update['sprite_combat_eligible'] or update['sprite_combat_integrated']:
                raise ValueError('A pending or historical identity cannot gain a new generated sprite')
            expected_status = 'HISTORICAL_R20' if uid in historical else awaiting[uid].get('generationStatus', 'AWAITING_OUTPUT')
            expected_reason = awaiting[uid].get('reason', '') if uid in awaiting else ''
            if update['sprite_generation_status'] != expected_status or update['sprite_generation_reason'] != expected_reason:
                raise ValueError('Pending generation statuses and tool refusal reasons must remain explicit')
            if update['sprite_pose_kind'] != ('neutral-base' if uid in historical else 'unknown'):
                raise ValueError('Pending poses stay unknown; historical neutral poses remain separately scoped')
            continue
        sprite = actual[uid]
        job = jobs[uid]
        if sprite.get('scope') != job.get('scope') or sprite['image'] != job.get('outputPath'):
            raise ValueError('Sprite scope and output path must preserve the frozen generation job')
        image = local_file(root, sprite['image'])
        if image.stat().st_size != sprite['bytes'] or digest(image) != sprite['sha256']:
            raise ValueError('Generated sprite bytes differ from their manifest')
        eligible = (sprite.get('reviewStatus') == 'visually-reviewed' and sprite.get('fullBodyReviewed') is True
                    and sprite.get('scope') == 'combatant-neutral-pose'
                    and sprite.get('sourceFacing') in ('left', 'right') and sprite.get('poseKind') == 'neutral-base')
        if (sprite.get('poseKind') not in ('neutral-base', 'source-bind-pose', 'source-presentation', 'unknown')
                or sprite['poseKind'] == 'neutral-base' and sprite.get('reviewStatus') != 'visually-reviewed'):
            raise ValueError('Pose classification must be explicit and neutral-base requires a visual review')
        if (sprite.get('combatIntegrated') is not False or sprite.get('canonicalFidelityCertified') is not False
                or sprite.get('canMirror') is not False or sprite.get('availableStates') != (['idle'] if eligible else [])):
            raise ValueError('Generated neutral poses cannot certify complete combat animation or original art')
        projected = dict(zip(SPRITE_KEYS, (True, sprite['reviewStatus'], sprite['fullBodyReviewed'],
                         sprite['completeBodyReferenceVerified'], sprite['scope'], sprite['sourceFacing'], sprite['poseKind'],
                         sprite['availableStates'], eligible, False, 'GENERATED', '')))
        if any(update[key] != value for key, value in projected.items()):
            raise ValueError('Sprite identity progress differs from the generated manifest')


def motion_rows(tracking, root):
    if not tracking.get('motion_manifest'):
        return None
    path = local_file(root, tracking['motion_manifest'])
    if not any(local_file(root, item['path']) == path for item in tracking['inputs']):
        raise ValueError('Motion manifest requires a recorded evidence hash')
    manifest = read(path)
    if manifest.get('schema') != 're.bow.rigged-motion-pack/1' or manifest.get('summary', {}).get('canonicalFidelityCertified') != 0:
        raise ValueError('Unexpected motion manifest or unsupported motion fidelity certification')
    catalogue = {item['id'] for item in read(root / 'data/roster.json')['catalog']}
    rows, clip_ids = [], set()
    motion_count, pose_count = 0, 0
    for entry in manifest['entries']:
        subject = entry.get('subjectId', entry['id'])
        role = entry.get('modelRole', 'body')
        if (subject not in catalogue or entry.get('canonicalFidelityCertified') is not False
                or entry.get('rigProductionFinalized') is not False
                or role not in ('body', 'body-variant', 'associated-object')
                or (entry['id'] != subject and role == 'body')):
            raise ValueError('Unknown motion identity or unsupported finalized rig')
        for clip in entry['clips']:
            key = (entry['id'], clip['id'])
            if key in clip_ids:
                raise ValueError('Duplicate source animation')
            clip_ids.add(key)
            try:
                start, end, ticks = int(clip['keyTimeStart']), int(clip['keyTimeEnd']), int(clip['ticksPerSecond'])
                if ticks <= 0 or start > end:
                    raise ValueError('Invalid native animation clock')
                source_span, preroll, last = (end - start) / ticks, max(0, -start / ticks), end / ticks
                for field, expected in (('sourceTimeSpanSeconds', source_span), ('nativePreRollSeconds', preroll), ('lastNativeTimeSeconds', last)):
                    if not math.isfinite(float(clip[field])) or not math.isclose(float(clip[field]), expected, rel_tol=1e-9, abs_tol=1e-9):
                        raise ValueError('Native animation times disagree: ' + field)
            except (KeyError, TypeError, OverflowError) as error:
                raise ValueError('Missing native animation clock evidence') from error
            source_varying = clip.get('varyingAnimationCurveCount')
            nonnegative_varying = clip.get('varyingNonnegativeAnimationCurveCount')
            pose_only = clip.get('positiveTimePoseOnly')
            if (type(source_varying) is not int or type(nonnegative_varying) is not int
                    or not 0 <= nonnegative_varying <= source_varying or type(pose_only) is not bool
                    or pose_only != (nonnegative_varying == 0)):
                raise ValueError('Inconsistent positive-time pose/motion classification')
            expected_scope = 'pose-only-after-time-zero' if pose_only else 'nonnegative-time-motion'
            if clip.get('playbackVariationScope') != expected_scope:
                raise ValueError('Playback variation scope disagrees with the positive-time curve count')
            motion_count += not pose_only
            pose_count += pose_only
            rows.append([subject, clip['id'], entry.get('game', ''), clip.get('member', ''),
                         clip.get('path') or 'MEMBRE DE L’ARCHIVE SOURCE', clip['sha256'],
                         start / ticks, last, source_span, preroll, max(0, last),
                         source_varying, nonnegative_varying, POSE_LABEL if pose_only else MOTION_LABEL,
                         clip['playbackVariationScope'],
                         'ACTION NON ATTRIBUÉE' if clip.get('semanticActionVerified') is not True else 'ACTION DOCUMENTÉE',
                         ('Une pose après zéro reste une pose même si le pré-roll varie. ' if pose_only else '')
                         + 'Étendue min→max distincte de la lecture 0→max ; pas de rig ou combat final certifié.',
                         entry['id'], role])
    if manifest['summary'].get('sourceAnimationFiles') != len(rows):
        raise ValueError('Source animation count differs from the motion manifest')
    if (manifest['summary'].get('positiveTimeMotionFiles') != motion_count
            or manifest['summary'].get('positiveTimePoseFiles') != pose_count):
        raise ValueError('Positive-time motion/pose totals differ from the clip classifications')
    return rows


def preserve_worksheet_xml(baseline, generated):
    """Retain original worksheet XML and relationships, including literal cells."""
    with ZipFile(baseline) as source, ZipFile(generated) as result:
        historical = {name: (source.getinfo(name), source.read(name)) for name in source.namelist() if name.startswith('xl/worksheets/')}
        missing = set(historical) - set(result.namelist())
        if missing:
            raise ValueError('A historical worksheet part was lost')
        with tempfile.NamedTemporaryFile(dir=generated.parent, suffix='.xlsx', delete=False) as handle:
            replacement = Path(handle.name)
        try:
            with ZipFile(replacement, 'w') as target:
                for item in result.infolist():
                    info, content = historical.get(item.filename, (item, result.read(item.filename)))
                    target.writestr(info, content)
            os.replace(replacement, generated)
        finally:
            replacement.unlink(missing_ok=True)


def worksheet_state(page):
    return (page.title, page.freeze_panes, tuple(str(item) for item in page.merged_cells.ranges),
            tuple(tuple((cell.value, cell.data_type, tuple(cell._style) if cell.has_style else None,
                         cell.hyperlink.target if cell.hyperlink else None,
                         (cell.comment.text, cell.comment.author) if cell.comment else None)
                        for cell in row) for row in page.iter_rows()))


def atomic_write(path, content):
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=path.parent, delete=False) as handle:
        temporary = Path(handle.name)
        handle.write(content)
    try:
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def export(tracking_path, baseline_path, output_path, progress_path, root=ROOT):
    root = Path(root).resolve()
    tracking_path, baseline_path = (checked_path(root, path, exists=True) for path in (tracking_path, baseline_path))
    output_path, progress_path = (checked_path(root, path) for path in (output_path, progress_path))
    if len({tracking_path, baseline_path, output_path, progress_path}) != 4:
        raise ValueError('Output must not overwrite historical inputs')
    tracking = validate(read(tracking_path), root)
    if not any(local_file(root, item['path']) == baseline_path for item in tracking['inputs']):
        raise ValueError('Historical workbook requires a recorded evidence hash')
    clips = motion_rows(tracking, root) or []
    book = load_workbook(baseline_path)
    if book.properties.creator != BASE_CREATOR:
        raise ValueError('Expected the historical R26 workbook')
    if any(name in book.sheetnames for name in ADDED_SHEETS):
        raise ValueError('The baseline already contains R27 sheets')
    if output_path.exists():
        previous = load_workbook(output_path, read_only=True)
        creator = previous.properties.creator
        previous.close()
        if creator != CREATOR:
            raise ValueError('Refusing to overwrite a workbook from another author')
    original_names = list(book.sheetnames)
    original_states = {name: worksheet_state(book[name]) for name in original_names}
    rows = list(book['Catalogue R26'].values)
    roster_ids = {item['id'] for item in read(root / 'data/roster.json')['catalog']}
    if len(rows) != 390 or {row[0] for row in rows[1:]} != roster_ids:
        raise ValueError('Historical catalogue does not match all 389 identities')
    book.properties.creator = CREATOR
    book.properties.title = 'B.O.W. Protocol — chantiers parallèles R27'
    sheet(book, 'Synthèse R27', ['Indicateur', 'Valeur', 'Portée'],
          [[item['label'], item['value'], item.get('scope', '')] for item in tracking['statistics']]
          + [['Fidélité globale 1:1 certifiée', 0, 'Les critères non vérifiés restent ouverts.']])
    sheet(book, 'Chantiers R27', ['ID', 'Chantier', 'Statut R27', 'Avancée vérifiée', 'Reste à réaliser', 'Preuves'],
          [[task['id'], task['title'], task['status'], '\n'.join(task['advances']), '\n'.join(task['remaining']),
            '\n'.join(item['label'] + ': ' + item['href'] for item in task['evidence'])] for task in tracking['tasks']])
    updates = {item['id']: item for item in tracking['identity_updates']}
    sheet(book, 'Catalogue R27', list(rows[0]) + ['Ressources visuelles R27', 'Armature source R27', 'Clips sources R27',
          'Clips audio documentaires R27', 'Citations documentaires R27', 'Limites R27',
          'Nouveau sprite généré', 'Revue du sprite', 'Corps entier revu', 'Référence corporelle complète déclarée dans le job figé',
          'Portée du sprite', 'Orientation disponible', 'Type de pose observé', 'États disponibles', 'Admissible au rendu idle de combat',
          'Intégration de combat certifiée', 'État de la génération', 'Motif d’attente ou de refus'],
          [list(row) + [updates[row[0]].get(key, '') for key in ('visual', 'rig', 'motion_clips', 'audio', 'claims', 'limitations', *SPRITE_KEYS)] for row in rows[1:]])
    sheet(book, 'Sources R27', ['Chantier', 'UID ou portée', 'Nature', 'Source ou fichier', 'Limite', 'Citation ou observation'], tracking['source_rows'])
    sheet(book, 'Recette R27', ['Contrôle', 'Résultat', 'Portée', 'Preuve'],
          [[item['label'], item['result'], item['scope'], item['path']] for item in tracking['validation']])
    if clips is not None:
        sheet(book, 'Clips sources R27', ['UID', 'Clip source', 'Jeu', 'Membre archive', 'Fichier direct', 'SHA-256',
              'Début natif s', 'Dernier temps natif s', 'Étendue source min→max s', 'Pré-roll natif avant zéro s',
              'Lecture 0→max s', 'Courbes variables source, pré-roll inclus', 'Courbes variables à t ≥ 0',
              'Classification lecture t ≥ 0', 'Portée de la variation', 'Attribution du geste', 'Limites',
              'Référence de modèle', 'Rôle : corps, variante ou objet associé'], clips)
    book.active = book.sheetnames.index('Chantiers R27')
    progress = {key: tracking[key] for key in ('summary', 'statistics', 'resources', 'tasks')}
    progress['schema'] = 're.bow.fidelity-progress/1'
    progress_bytes = (json.dumps(progress, ensure_ascii=False, indent=2) + '\n').encode('utf-8')
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=output_path.parent, suffix='.xlsx', delete=False) as handle:
        temporary = Path(handle.name)
    try:
        book.save(temporary)
        book.close()
        preserve_worksheet_xml(baseline_path, temporary)
        check = load_workbook(temporary, data_only=False)
        try:
            if check.sheetnames[:len(original_names)] != original_names or any(worksheet_state(check[name]) != original_states[name] for name in original_names):
                raise ValueError('A historical worksheet changed during export')
            if check['Catalogue R27'].max_row != 390 or check['Chantiers R27'].max_row != 18:
                raise ValueError('Unexpected catalogue or workstream count after export')
            for name in check.sheetnames[len(original_names):]:
                if any(cell.data_type == 'f' for row in check[name] for cell in row):
                    raise ValueError('Spreadsheet formula injection rejected')
        finally:
            check.close()
        os.replace(temporary, output_path)
        atomic_write(progress_path, progress_bytes)
    finally:
        temporary.unlink(missing_ok=True)
    return {'workbook': str(output_path), 'progress': str(progress_path), 'open_workstreams': 17,
            'catalogue_identities': 389, 'source_clips': len(clips) if clips is not None else 0,
            'positive_time_motion_files': sum(row[13] == MOTION_LABEL for row in clips) if clips is not None else 0,
            'positive_time_pose_files': sum(row[13] == POSE_LABEL for row in clips) if clips is not None else 0,
            'historical_worksheets_preserved': len(original_names), 'canonical_fidelity_certified': 0}


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--tracking', type=Path, default=ROOT / 'docs/R27_TRACKING.json')
    ap.add_argument('--baseline', type=Path, default=ROOT / 'docs/BOW_FIDELITE_R26.xlsx')
    ap.add_argument('--output', type=Path, default=ROOT / 'docs/BOW_FIDELITE_R27.xlsx')
    ap.add_argument('--progress', type=Path, default=ROOT / 'modules/fidelity-r27/progress.json')
    args = ap.parse_args()
    print(json.dumps(export(args.tracking, args.baseline, args.output, args.progress), ensure_ascii=False))


if __name__ == '__main__':
    main()
