#!/usr/bin/env python3
"""Build, export and reproduce R29 progress without copying historical assets."""
import argparse
from collections import Counter
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MODULE = Path('modules/fidelity-r29')
SPEC = importlib.util.spec_from_file_location('release_r28_helpers', ROOT / 'scripts/export-fidelity-progress-r28.py')
OLD = importlib.util.module_from_spec(SPEC); SPEC.loader.exec_module(OLD)
UTIL = OLD.OLD
DOSSIER_SPEC = importlib.util.spec_from_file_location('r29_source_observation_helpers', ROOT / 'scripts/build-dossier-additions-r29.py')
DOSSIER = importlib.util.module_from_spec(DOSSIER_SPEC); DOSSIER_SPEC.loader.exec_module(DOSSIER)
SHEETS = ('Synthèse R29', 'Chantiers R29', 'Catalogue R29', 'Sources R29', 'Recette R29', 'Images natives R29', 'Corps sources R29')
FIELDS = ('visual', 'sprite_full_body_reviewed', 'sprite_reference_full_body_verified', 'sprite_facing',
          'sprite_pose_kind', 'sprite_available_states', 'sprite_combat_eligible', 'sprite_scope',
          'r29_pixel_revision', 'r29_pose_requalification', 'r29_metadata_update', 'r29_source_models', 'r29_source_views',
          'r29_native_clips', 'r29_native_frames', 'r29_native_combat_states', 'r29_equipment_complete',
          'r29_known_appearance_gaps', 'r29_body_jointly_qualified', 'r29_documentary_variants',
          'r29_documentary_source_editions', 'r29_source_body_observed', 'r29_active_source_edition',
          'r29_active_source_variant_label', 'limitations')
REPORTS = [
    ('Export et périmètres', 'R29_EXPORT_TESTS.json', 'python_tests'),
    ('Registre et poses', 'R29_SPRITES_TESTS.json', 'python_tests'),
    ('Admission au rendu statique', 'R29_COMBAT_SPRITE_TESTS.json', 'targeted_node_groups'),
    ('Galerie des poses', 'R29_SPRITES_BROWSER.json', 'browser_groups'),
    ('Références corporelles', 'R29_BODY_SOURCES_TESTS.json', 'python_tests'),
    ('Références du second lot', 'R29_REVELATIONS_BODY_TESTS.json', 'python_tests'),
    ('Fichiers des images natives', 'R29_NATIVE_FRAMES_TESTS.json', 'python_tests'),
    ('Déformation native', 'R29_NATIVE_FRAME_RIG_TESTS.json', 'targeted_node_groups'),
    ('Lecteur des images natives', 'R29_NATIVE_FRAMES_BROWSER.json', 'browser_groups'),
    ('Sélection et contexte opaque', 'R29_SELECTION_BROWSER.json', 'browser_groups'),
    ('Dossiers par incarnation', 'R29_DOSSIER_BROWSER.json', 'browser_groups'),
    ('Suivi des dix-sept chantiers', 'R29_DASHBOARD_BROWSER.json', 'browser_groups'),
]


def read(path, root=ROOT):
    return json.loads((Path(root) / path).read_text())


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def raw(value):
    return (json.dumps(value, ensure_ascii=False, indent=2) + '\n').encode()


def write(path, content):
    UTIL.atomic_write(Path(path), content if isinstance(content, bytes) else raw(content))


def eligible(row, root=ROOT):
    return (row.get('reviewStatus') == 'visually-reviewed' and row.get('fullBodyReviewed') is True
            and DOSSIER.source_body_observed(row, root) and row.get('poseKind') == 'neutral-base'
            and row.get('sourceFacing') in ('left', 'right') and row.get('scope') == 'combatant-neutral-pose'
            and not row.get('knownAppearanceGaps') and row.get('availableStates') == ['idle'])


def validation(root=ROOT, allow_pending=False):
    root = Path(root).resolve(); checks, pending, failed = [], [], []
    totals = {'python_tests': 0, 'targeted_node_groups': 0, 'browser_groups': 0}
    for label, name, kind in REPORTS:
        path = root / 'docs' / name
        if not path.is_file():
            pending.append('docs/' + name); continue
        report = read(path); counts = report.get('counts', {})
        passed, failures = report.get('passed', counts.get('passed')), report.get('failed', counts.get('failed'))
        if type(passed) is not int or type(failures) is not int or passed + failures <= 0:
            raise ValueError('Report has no executed checks: ' + name)
        if failures or report.get('pass') is False:
            failed.append('docs/' + name)
        source_current = True
        for key in ('sourceFiles', 'executedAdapterInputs', 'sources', 'inputs'):
            bindings = report.get(key, [])
            for binding in bindings if isinstance(bindings, list) else []:
                if not isinstance(binding, dict) or not isinstance(binding.get('path'), str) or not isinstance(binding.get('sha256'), str):
                    continue
                source = (root / binding['path']).resolve()
                if not source.is_relative_to(root) or not source.is_file() or sha(source) != binding['sha256']:
                    failed.append('docs/' + name + ': stale executed source ' + binding['path'])
                    source_current = False
        totals[kind] += passed if source_current else 0
        checks.append({'label': label, 'path': 'docs/' + name, 'sha256': sha(path), 'kind': kind,
                       'passed': passed, 'failed': failures, 'source_current': source_current,
                       'result': f'{passed} réussis, {failures} échec(s)' if source_current else 'Sources modifiées depuis l’exécution ; résultat non compté comme R29 courant.',
                       'scope': 'Ajouts R29 dans la portée du rapport exécuté.'})
    if not allow_pending and (pending or failed):
        raise ValueError('R29 reports are unfinished: ' + ', '.join(pending + failed))
    return {'schema': 're.bow.validation-r29/1', 'release': 'R29', 'complete': not pending and not failed,
            'provisional': bool(pending or failed), 'checks': checks, 'totals': totals,
            'pending_reports': pending, 'failed_reports': failed, 'canonical_fidelity_certified': 0,
            'historical_results': {'path': 'docs/R28_VALIDATION.json', 'sha256': sha(root / 'docs/R28_VALIDATION.json'),
                'scope': 'Historique conservé, non recompté comme une exécution R29.'},
            'post_export_evidence': ['docs/R29_EXPORT.json', 'docs/R29_REPRODUCTION.json']}


def evidence_link(root, href):
    if not isinstance(href, str) or not href or '\\' in href or href.startswith('/') or re.match(r'^[a-zA-Z][a-zA-Z0-9+.-]*:', href):
        raise ValueError('Evidence link must remain local')
    return UTIL.checked_path(root, MODULE / href, exists=True)


def validate(value, root=ROOT):
    root = Path(root).resolve()
    if value.get('schema') != 're.bow.parallel-workstreams/3' or value.get('release') != 'R29' or value.get('canonical_fidelity_certified') != 0:
        raise ValueError('Unsupported R29 progress or certification')
    tasks = value['tasks']; old = read('docs/R28_TRACKING.json', root)['tasks']
    if len(tasks) != 17 or {row['id'] for row in tasks} != {row['id'] for row in old} or any(row['status'] != 'EN COURS' for row in tasks):
        raise ValueError('All seventeen open workstreams must remain explicit')
    for task in tasks:
        if not all(task.get(key) for key in ('title', 'advances', 'remaining', 'evidence')):
            raise ValueError('An open workstream needs progress, limits and proof')
        for link in task['evidence']:
            evidence_link(root, link['href'])
    for link in value['resources']:
        evidence_link(root, link['href'])
    inputs = value['inputs']
    if len({row['path'] for row in inputs}) != len(inputs) or not inputs:
        raise ValueError('Progress inputs must be unique')
    if not {'docs/BOW_FIDELITE_R28.xlsx', 'docs/R28_TRACKING.json'} <= {row['path'] for row in inputs}:
        raise ValueError('Historical workbook and progress require exact input bindings')
    for row in inputs:
        path = UTIL.local_file(root, row['path'])
        if not re.fullmatch('[a-f0-9]{64}', row['sha256']) or sha(path) != row['sha256']:
            raise ValueError('Evidence hash mismatch: ' + row['path'])
    known = {row['id'] for row in read('data/roster.json', root)['catalog']}
    updates = value['identity_updates']; by_uid = {row['id']: row for row in updates}
    if len(known) != 389 or len(updates) != 389 or set(by_uid) != known or any(any(key not in row for key in FIELDS) for row in updates):
        raise ValueError('All 389 identities require exact progress')
    for row in value['source_rows']:
        if not isinstance(row, list) or len(row) != 6:
            raise ValueError('Source evidence requires six columns')
        for cell in row:
            UTIL.safe_cell(cell)
    if value.get('sprite_manifest'):
        pack = read(value['sprite_manifest'], root); before = read('modules/fidelity-r28/sprites/manifest.json', root)
        if pack['pending'] != before['pending'] or len(pack['pending']) != 22 or {row['id'] for row in pack['sprites']} != {row['id'] for row in before['sprites']}:
            raise ValueError('All 22 refusals and active identities must stay preserved')
        counts = pack['counts']; sprites = pack['sprites']
        if (counts['generatedSprites'] != len(sprites) or counts['fullBodyReviewedSprites'] != sum(row['fullBodyReviewed'] for row in sprites)
                or counts['neutralBaseCombatEligibleSprites'] != sum(eligible(row, root) for row in sprites)
                or counts.get('originalGameSprites') != 0 or counts.get('completeCombatAnimationSets') != 0):
            raise ValueError('Files, body reviews and source cycles cannot be conflated')
        for sprite in sprites:
            image = UTIL.local_file(root, sprite['image']); update = by_uid[sprite['id']]
            if sprite.get('knownAppearanceGaps') and sprite.get('availableStates'):
                raise ValueError('A known appearance gap cannot acquire a combat state')
            if sha(image) != sprite['sha256'] or image.stat().st_size != sprite['bytes']:
                raise ValueError('Active sprite bytes changed')
            projection = {'visual': sprite['image'], 'sprite_full_body_reviewed': sprite['fullBodyReviewed'],
                'sprite_reference_full_body_verified': sprite['completeBodyReferenceVerified'], 'sprite_facing': sprite['sourceFacing'],
                'sprite_pose_kind': sprite['poseKind'], 'sprite_available_states': sprite['availableStates'],
                'sprite_combat_eligible': eligible(sprite, root), 'sprite_scope': sprite['scope'],
                'r29_source_body_observed': DOSSIER.source_body_observed(sprite, root),
                'r29_active_source_edition': sprite.get('sourceEdition') or '',
                'r29_active_source_variant_label': sprite.get('sourceVariantLabel') or ''}
            if any(update[key] != expected for key, expected in projection.items()):
                raise ValueError('Sprite qualification differs from actual evidence')
        before_sprites = {row['id']: row for row in before['sprites']}
        current_sprites = {row['id']: row for row in sprites}
        previous_gaps = {row['id'] for row in read('modules/fidelity-r28/sprites/body-gaps.json', root)['entries']}
        current_gaps = {row['id'] for row in read(MODULE / 'sprites/body-gaps.json', root)['entries']}
        for variant in pack.get('documentaryVariants', []):
            uid = variant['id']
            if (variant.get('admittedTargetEdition') is not False or variant.get('targetFullBodyReviewed') is not False
                    or variant.get('availableStates') != [] or current_sprites[uid]['sha256'] != before_sprites[uid]['sha256']
                    or uid in previous_gaps and uid not in current_gaps):
                raise ValueError('An edition variant cannot replace or complete its target body')
    if value.get('dossier_index'):
        if value['dossier_index'] not in {row['path'] for row in inputs}:
            raise ValueError('Dossier index requires an input binding')
        for entry in read(value['dossier_index'], root)['identities']:
            row = by_uid[entry['id']]
            if (row['r29_source_models'] != entry['bodyModels'] or row['r29_source_views'] != entry['bodyViews']
                    or row['r29_documentary_variants'] != [variant['variantId'] for variant in entry['documentaryVariants']]
                    or row['r29_documentary_source_editions'] != [variant['sourceEdition'] for variant in entry['documentaryVariants']]):
                raise ValueError('Dossier references or edition variants differ from their exact UID')
    if value.get('native_manifest'):
        if value['native_manifest'] not in {row['path'] for row in inputs}:
            raise ValueError('Native evidence requires an input binding')
        native = read(value['native_manifest'], root)
        clip_counts, frame_counts, rendered = Counter(), Counter(), {}
        for entry in native['entries']:
            if (entry['subjectId'] not in known or entry['actionIdentity'] != 'unassigned' or entry['availableCombatStates'] != []
                    or entry['originalGameFrames'] is not False or entry['canonicalFidelityCertified'] is not False
                    or entry['frameCount'] != len(entry['frames'])):
                raise ValueError('Source frames cannot become assigned combat cycles')
            clip_counts[entry['subjectId']] += 1; frame_counts[entry['subjectId']] += entry['frameCount']
            for frame in entry['frames'] + entry.get('poseViews', []):
                path = (MODULE / 'motion' / frame['path']).as_posix(); image = UTIL.local_file(root, path)
                if sha(image) != frame['sha256'] or image.stat().st_size != frame['bytes']:
                    raise ValueError('A native image differs from its source proof')
                rendered[path] = frame
        for uid, row in by_uid.items():
            if row['r29_native_clips'] != clip_counts[uid] or row['r29_native_frames'] != frame_counts[uid] or row['r29_native_combat_states'] != []:
                raise ValueError('Per-UID native counts or combat states were invented')
        if len(value['native_rows']) != len(rendered):
            raise ValueError('Native worksheet must expose the delivered images exactly once')
        for row in value['native_rows']:
            frame = rendered.get(row[5])
            if (not frame or row[6:9] != [frame['sha256'], frame.get('ticks', ''), frame.get('seconds', '')]
                    or row[-3:] != ['NON ATTRIBUÉ', False, False]):
                raise ValueError('Native worksheet changed an observation or assigned a cycle')
    return value


def build(root=ROOT, allow_pending=False):
    root = Path(root).resolve(); prior = read('docs/R28_TRACKING.json', root)
    pack = read(MODULE / 'sprites/manifest.json', root); old_pack = read('modules/fidelity-r28/sprites/manifest.json', root)
    dossier = read(MODULE / 'dossier-additions.json', root); checks = validation(root, allow_pending)
    rows, frames, body_rows = [], [], []
    inputs = ['docs/BOW_FIDELITE_R28.xlsx', 'docs/R28_TRACKING.json', 'data/roster.json',
              'modules/fidelity-r28/sprites/manifest.json', (MODULE / 'sprites/manifest.json').as_posix(),
              (MODULE / 'sprites/body-gaps.json').as_posix(), (MODULE / 'dossier-additions.json').as_posix(),
              (MODULE / 'dossier-link.js').as_posix(), (MODULE / 'index.html').as_posix(),
              'scripts/fidelity-progress-r29.py', 'scripts/build-dossier-additions-r29.py'] + [row['path'] for row in checks['checks']]
    for guide in ('docs/R29_SPRITES.md', 'docs/R29_BODY_SOURCES.md', 'docs/R29_REVELATIONS_BODIES.md',
                  'docs/R29_NATIVE_FRAMES.md', 'docs/R29_DOSSIER.md'):
        if (root / guide).is_file():
            inputs.append(guide)
    before = {row['id']: row for row in old_pack['sprites']}; sprites = {row['id']: row for row in pack['sprites']}
    deltas = {row['id']: row for row in dossier['identities']}
    gaps = read(MODULE / 'sprites/body-gaps.json', root); gap_ids = {row['id'] for row in gaps['entries']}
    counts = pack['counts']; native_path = MODULE / 'motion/manifest.json'
    native = read(native_path, root) if (root / native_path).is_file() else {'entries': []}
    if native['entries']:
        inputs.append(native_path.as_posix())
    body_paths = [MODULE / 'sources/body-rasters/manifest.json', MODULE / 'sources/revelations-body-rasters/manifest.json']
    for path in body_paths:
        if not (root / path).is_file():
            continue
        body = read(path, root); inputs.append(path.as_posix())
        for model in body['entries']:
            body_rows.append([model['id'], model.get('modelId', model['id']), model.get('sourceEdition', ''),
                model.get('sourcePage', ''), len(model.get('views', [])), model.get('sourceVariantScope', model.get('sourcePoseScope', '')),
                model.get('equipmentScope', ''), False])
    for row in dossier['identities']:
        for model in row['bodySources']:
            for view in model['views']:
                path = (MODULE / view['path']).as_posix(); inputs.append(path)
                rows.append(['G-PREUVES', row['id'], 'Vue dérivée du modèle source', path,
                             model['sourceEdition'] + '; ' + model['variant'] + '; ' + model['equipmentScope'], view['sha256']])
    for entry in native['entries']:
        for kind, images in (('image de clés natives', entry['frames']), ('vue de pose source', entry.get('poseViews', []))):
            for frame in images:
                path = (MODULE / 'motion' / frame['path']).as_posix(); inputs.append(path)
                rows.append(['R20-09', entry['subjectId'], kind, path,
                             entry['sourceEdition'] + '; geste non attribué, aucun cycle de combat admis.', frame['sha256']])
                frames.append([entry['subjectId'], entry['id'], entry['clipId'], entry.get('modelRole', 'body'), entry['sourceEdition'],
                    path, frame['sha256'], frame.get('ticks', ''), frame.get('seconds', ''), frame.get('angle', ''), kind,
                    'NON ATTRIBUÉ', False, False])
    review_path = pack['reviews']['path']; pose_review_ids = set()
    if (root / review_path).is_file():
        review = read(review_path, root); inputs.append(review_path)
        for item in review.get('entries', []):
            pose_review_ids.add(item.get('uid', item.get('id')))
            if item.get('path'):
                inputs.append(item['path']); rows.append(['G-RENDU', item.get('uid', item.get('id', '')), 'Requalification de pose R29', item['path'],
                                                        'Observation de la pose existante ; aucun nouveau PNG déduit.', item.get('sha256', '')])
    for uid in sorted(pack['replacements']):
        sprite = sprites[uid]; inputs.append(sprite['image'])
        rows.append(['R20-10', uid, 'Version corporelle remplacée R29', sprite['image'],
                     'Anatomie, accessoires, phase et état de combat restent qualifiés séparément.', sprite['sha256']])
        proof = sprite.get('sourceCompletenessEvidence')
        if proof:
            inputs.append(proof['path'])
    for sprite in pack['sprites']:
        source_note = sprite.get('r29LegacySourceNote')
        if source_note:
            inputs.append(source_note['path'])
            rows.append(['G-PREUVES', sprite['id'], 'Qualification d’édition héritée documentée en R29', source_note['path'],
                         sprite['sourceEdition'] + '; ' + sprite['sourceVariantLabel'] + '; aucun pixel, corps ou pose ajouté.', source_note['sha256']])
        if sprite.get('completeBodyReferenceVerified') is not True and sprite.get('availableStates') == ['idle']:
            files = DOSSIER.retained_source_observation_files(sprite, root)
            if not files:
                raise ValueError('A legacy idle needs its retained source observation')
            inputs.extend(files)
            rows.append(['G-RENDU', sprite['id'], 'Qualification source héritée revérifiée', files[0],
                         'Déclaration du job figée inchangée ; corps source effectivement observé dans une preuve manuelle SHA-liée.', sha(root / files[0])])
    for variant in pack.get('documentaryVariants', []):
        inputs.append(variant['image'])
        rows.append(['G-PREUVES', variant['id'], 'Interprétation documentaire d’une autre édition', variant['image'],
                     variant['sourceEdition'] + '; cible ' + variant['targetIncarnation'] + '; aucune admission ni fermeture du corps cible.', variant['sha256']])
    screening = MODULE / 'research/appearance-source-screening.json'
    if (root / screening).is_file():
        inputs.append(screening.as_posix()); rows.append(['G-PREUVES', 'CANDIDATS RETENUS HORS PRODUCTION', 'Périmètre des sources candidates',
            screening.as_posix(), 'Mauvaise phase ou édition : aucune admission corporelle ni génération déduite.', sha(root / screening)])
    updates = []
    for previous in prior['identity_updates']:
        uid = previous['id']; sprite = sprites.get(uid); delta = deltas[uid]
        update = dict(previous)
        if sprite:
            update.update(visual=sprite['image'], sprite_full_body_reviewed=sprite['fullBodyReviewed'],
                sprite_reference_full_body_verified=sprite['completeBodyReferenceVerified'], sprite_facing=sprite['sourceFacing'],
                sprite_pose_kind=sprite['poseKind'], sprite_available_states=sprite['availableStates'],
                sprite_combat_eligible=eligible(sprite, root), sprite_scope=sprite['scope'])
        update.update(r29_pixel_revision=bool(sprite and sprite['sha256'] != before[uid]['sha256']),
            r29_pose_requalification=uid in pose_review_ids,
            r29_metadata_update=bool(delta['spriteChange'] and not delta['spriteChange']['pixelRevision']),
            r29_source_models=delta['bodyModels'], r29_source_views=delta['bodyViews'],
            r29_native_clips=len(delta['nativeClips']), r29_native_frames=sum(clip['frameCount'] for clip in delta['nativeClips']),
            r29_native_combat_states=[], r29_equipment_complete=bool(sprite and sprite.get('equipmentCompletenessReviewed')),
            r29_known_appearance_gaps=sprite.get('knownAppearanceGaps', []) if sprite else [],
            r29_body_jointly_qualified=bool(sprite and uid not in gap_ids),
            r29_documentary_variants=[variant['variantId'] for variant in delta['documentaryVariants']],
            r29_documentary_source_editions=[variant['sourceEdition'] for variant in delta['documentaryVariants']],
            r29_source_body_observed=bool(sprite and DOSSIER.source_body_observed(sprite, root)),
            r29_active_source_edition=sprite.get('sourceEdition', '') if sprite else '',
            r29_active_source_variant_label=sprite.get('sourceVariantLabel', '') if sprite else '',
            limitations='La vue du modèle et les images de clés natives restent des références. Geste, équipements, autres phases, cycle jouable, rendu et dimensions 1:1 ne sont pas certifiés par ces seuls fichiers.')
        updates.append(update)
    stats = dossier['summary']; native_frames = sum(entry['frameCount'] for entry in native['entries'])
    pose_views = sum(len(entry.get('poseViews', [])) for entry in native['entries'])
    facts = {
        'R20-09': [f"{len(native['entries'])} scènes sources donnent {native_frames} images de clés natives et {pose_views} vues de pose. Les gestes restent non attribués et aucun cycle de combat n’est admis."],
        'R20-10': [f"{len(pack['replacements'])} versions corporelles sont remplacées en R29 ; {counts.get('newPoseReviews', 0)} poses existantes sont requalifiées. Les 363 identités actives et 22 refus restent distincts.",
                   f"Le cumul compte {counts['fullBodyReviewedSprites']} corps anatomiques revus et {counts['neutralBaseCombatEligibleSprites']} attentes statiques admissibles.",
                   f"{len(pack.get('documentaryVariants', []))} variantes documentaires d’une autre édition restent séparées des PNG actifs et des corps cibles à compléter.",
                   f"{counts.get('recordedRevisionFailures', 0)} refus de révision sont conservés séparément des 22 cibles historiquement refusées ; ils ne produisent aucun nouveau PNG."],
        'G-PREUVES': [f"Les Dossiers exposent les seules vues locales du sujet : {stats['bodySourceModels']} modèles attribués. Les empreintes, éditions, variantes et équipements restent lisibles séparément.",
                      'Les sources candidates de mauvaise phase ou édition gardent leur statut hors production.'],
        'G-RENDU': ['Les requalifications de pose et les nouvelles versions conservent leurs corps, côté et état réellement produits. Les images natives ne sont pas intégrées comme attaques ou déplacements.'],
        'G-ECHELLE': ['Les caméras et unités des modèles sources restent documentées ; les limites de rendu ne deviennent pas des dimensions physiques de combat.'],
        'G-COMMANDES': ['Les commandes et évolutions R28 restent dans leur contrat conservé. Les clips BA/DM et autres codes sources ne sont pas promus en gestes officiels à partir de leur nom.'],
        'G-PERF': ['Le registre de poses et les nouvelles preuves natives gardent leurs bornes de transport et de décodage. Le Codex ne charge aucune vue corporelle PNG automatiquement.'],
        'G-ACCESS': ['Le suivi et les Dossiers R29 exposent les identités exactes, liens locaux et critères ouverts ; la recette distingue les parcours navigateur des appareils physiques.'],
        'G-RECETTE': ['Les rapports exécutés R29 vérifient les nouveaux corps, requalifications, clés et déformations natives, sélecteurs, Dossiers et export. Les recettes historiques restent séparées.'],
        'G-LIVRAISON': ['Les 42 feuilles R28 et leurs relations sont conservées. Sept nouvelles feuilles suivent les 17 chantiers, les 389 identités et les images natives ; aucun chantier n’est clôturé automatiquement.'],
    }
    links = {'R20-09': [('motion/index.html', 'Images des clés natives')], 'R20-10': [('sprites/index.html', 'Poses et versions R29')],
        'G-PREUVES': [('dossier-additions.json', 'Références par UID')], 'G-RENDU': [('sprites/manifest.json', 'Qualifications actuelles')],
        'G-ECHELLE': [('dossier-additions.json', 'Vues et périmètres des modèles')], 'G-COMMANDES': [('../fidelity-r28/evolutions/graph.json', 'Évolutions conservées')],
        'G-PERF': [('motion/index.html', 'Lecteur des images sources')], 'G-ACCESS': [('index.html', 'Suivi central R29')],
        'G-RECETTE': [('../../docs/R29_VALIDATION.json', 'Rapports exécutés')], 'G-LIVRAISON': [('../../docs/R29_TRACKING.json', 'Suivi structuré R29')]}
    if (root / screening).is_file():
        links['G-PREUVES'].append(('research/appearance-source-screening.json', 'Sources hors production : phase et édition'))
    tasks = []
    for old in prior['tasks']:
        task = {key: old[key] for key in ('id', 'title', 'category')}; tid = task['id']
        task.update(status='EN COURS', advances=facts.get(tid, ['Acquis R28 conservés dans leur recette d’origine.']),
            remaining=list(old['remaining']), evidence=[{'href': href, 'label': label} for href, label in links.get(tid, [])]
                or [{'href': '../fidelity-r28/' + proof['href'], 'label': proof['label'] + ' · R28'} for proof in old['evidence']],
            r29_new_implementation=tid in facts)
        if tid == 'R20-10':
            task['remaining'] = [f"{gaps['combatantBodiesPending']} corps de combattants et {gaps['contextBodiesNotQualified']} contextes restent conjointement non qualifiés ; les 22 refus sont conservés.",
                                 'Équipements, autres phases, deux orientations et cycles complets de combat encore non vérifiés.']
        tasks.append(task)
    statistics = [
        {'label': 'poses actives revues', 'value': counts['visuallyReviewedSprites'], 'scope': 'Même périmètre de 385 cibles ; 22 refus conservés.'},
        {'label': 'versions remplacées en R29', 'value': len(pack['replacements']), 'scope': 'Versions, distinctes des nouvelles identités.'},
        {'label': 'variantes documentaires d’autre édition', 'value': len(pack.get('documentaryVariants', [])), 'scope': 'Aucun remplacement du PNG actif ou corps cible complété.'},
        {'label': 'poses requalifiées en R29', 'value': counts.get('newPoseReviews', 0), 'scope': 'Lecture de pixels existants, sans nouveau fichier déduit.'},
        {'label': 'corps anatomiques entiers revus', 'value': counts['fullBodyReviewedSprites'], 'scope': 'Équipements et phases restent des critères distincts.'},
        {'label': 'attentes statiques admissibles', 'value': counts['neutralBaseCombatEligibleSprites'], 'scope': 'Corps, source, pose neutre et côté disponibles vérifiés.'},
        {'label': 'images des clés natives', 'value': native_frames, 'scope': f'{len(native["entries"])} scènes ; geste non attribué, zéro cycle de combat admis.'},
        {'label': 'vues de pose native', 'value': pose_views, 'scope': 'Vues dérivées, distinctes des images du jeu original.'},
        {'label': 'corps de combattants encore ouverts', 'value': gaps['combatantBodiesPending'], 'scope': f"{gaps['contextBodiesNotQualified']} contextes et 22 refus conservés séparément."},
        {'label': 'chantiers ouverts', 'value': 17, 'scope': 'Critères manquants maintenus explicitement.'}]
    value = {'schema': 're.bow.parallel-workstreams/3', 'release': 'R29', 'canonical_fidelity_certified': 0,
        'summary': f"{len(pack['replacements'])} versions corporelles R29, {counts.get('newPoseReviews', 0)} poses requalifiées et {native_frames} images natives documentées ; 17 chantiers ouverts.",
        'statistics': statistics, 'tasks': tasks, 'identity_updates': updates, 'source_rows': rows, 'native_rows': frames, 'body_rows': body_rows,
        'resources': [{'href': 'sprites/index.html', 'title': 'Poses R29', 'label': 'Examiner les versions et qualifications'},
            {'href': 'motion/index.html', 'title': 'Images natives', 'label': 'Examiner les clés et éditions sources'},
            {'href': 'dossier-additions.json', 'title': 'Références par incarnation', 'label': 'Lire l’index des preuves des Dossiers'},
            {'href': 'sprites/body-gaps.json', 'title': 'Corps encore ouverts', 'label': 'Lire les critères restants'},
            {'href': 'evolutions/index.html', 'title': 'Évolutions volontaires', 'label': 'Jouer avec les règles conservées et la présentation R29'},
            {'href': '../fidelity-r28/audio/index.html', 'title': 'Audio documentaire', 'label': 'Écouter les groupes et candidats conservés'}],
        'validation': checks['checks'], 'validation_complete': checks['complete'],
        'validation_gaps': checks['pending_reports'] + checks['failed_reports'],
        'sprite_manifest': (MODULE / 'sprites/manifest.json').as_posix(),
        'dossier_index': (MODULE / 'dossier-additions.json').as_posix(),
        'native_manifest': native_path.as_posix() if native['entries'] else None,
        'generation_preference': prior['generation_preference'],
        'inputs': [{'path': path, 'sha256': sha(root / path)} for path in dict.fromkeys(inputs)]}
    return checks, value


def progress(value):
    return {**{key: value[key] for key in ('summary', 'statistics', 'resources', 'tasks')}, 'schema': 're.bow.fidelity-progress/1', 'release': 'R29'}


def export(value, root=ROOT, destination=None, progress_path=None):
    root = Path(root).resolve(); validate(value, root)
    baseline = root / 'docs/BOW_FIDELITE_R28.xlsx'; destination = Path(destination or root / 'docs/BOW_FIDELITE_R29.xlsx')
    progress_path = Path(progress_path or root / MODULE / 'progress.json')
    destination = UTIL.checked_path(root, destination); progress_path = UTIL.checked_path(root, progress_path)
    if destination == baseline or progress_path == baseline or destination == progress_path:
        raise ValueError('Output cannot overwrite historical inputs')
    if destination.exists():
        previous = load_workbook(destination, read_only=True); creator = previous.properties.creator; previous.close()
        if creator != 'BOW fidelity progress R29':
            raise ValueError('Refusing to overwrite another author’s workbook')
    book = load_workbook(baseline)
    if book.properties.creator != 'BOW fidelity progress R28' or any(name in book.sheetnames for name in SHEETS):
        raise ValueError('Expected the preserved R28 baseline')
    original_names = list(book.sheetnames); catalogue = list(book['Catalogue R28'].values)
    if len(catalogue) != 390 or {row[0] for row in catalogue[1:]} != {row['id'] for row in value['identity_updates']}:
        raise ValueError('Baseline catalogue does not preserve all 389 identities')
    book.properties.creator = 'BOW fidelity progress R29'; book.properties.title = 'B.O.W. Protocol — corps et images natives R29'
    UTIL.sheet(book, SHEETS[0], ['Indicateur', 'Valeur', 'Portée'], [[row['label'], row['value'], row['scope']] for row in value['statistics']]
        + [['Recette R29 complète', value.get('validation_complete', False), 'Les rapports manquants restent explicites dans la recette.'], ['Fidélité globale 1:1 certifiée', 0, 'Critères encore ouverts.']])
    UTIL.sheet(book, SHEETS[1], ['ID', 'Chantier', 'Statut', 'Acquis R29', 'Critères ouverts', 'Preuves'],
        [[row['id'], row['title'], row['status'], '\n'.join(row['advances']), '\n'.join(row['remaining']), '\n'.join(proof['label'] + ': ' + proof['href'] for proof in row['evidence'])] for row in value['tasks']])
    by_uid = {row['id']: row for row in value['identity_updates']}
    UTIL.sheet(book, SHEETS[2], list(catalogue[0]) + list(FIELDS), [list(row) + [by_uid[row[0]][key] for key in FIELDS] for row in catalogue[1:]])
    UTIL.sheet(book, SHEETS[3], ['Chantier', 'UID ou portée', 'Nature', 'Fichier', 'Limites', 'SHA ou citation'], value['source_rows'])
    UTIL.sheet(book, SHEETS[4], ['Contrôle', 'Résultat', 'Portée', 'Rapport', 'SHA'], [[row['label'], row['result'], row['scope'], row['path'], row['sha256']] for row in value['validation']]
        + [['Rapport en attente ou en échec', 'NON VALIDÉ', 'Version provisoire, résultat non compté comme réussi.', path, ''] for path in value.get('validation_gaps', [])])
    UTIL.sheet(book, SHEETS[5], ['UID', 'Scène', 'Clip', 'Rôle', 'Édition source', 'Image', 'SHA', 'Ticks natifs', 'Secondes', 'Angle', 'Type', 'Geste', 'Cycle de combat admis', 'Image du jeu original'], value['native_rows'])
    UTIL.sheet(book, SHEETS[6], ['UID', 'Modèle', 'Édition', 'Publication', 'Vues', 'Variante/pose', 'Équipement', 'Image originale certifiée'], value['body_rows'])
    book.active = book.sheetnames.index('Chantiers R29'); destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=destination.parent, suffix='.xlsx', delete=False) as handle:
        temporary = Path(handle.name)
    try:
        book.save(temporary); book.close(); OLD.canonical_zip(baseline, temporary)
        with ZipFile(baseline) as before, ZipFile(temporary) as after:
            if any(before.read(name) != after.read(name) for name in before.namelist() if name.startswith('xl/worksheets/')):
                raise ValueError('Historical worksheet parts changed')
        check = load_workbook(temporary)
        try:
            if check['Catalogue R29'].max_row != 390 or check['Chantiers R29'].max_row != 18 or any(cell.data_type == 'f' for title in SHEETS for row in check[title] for cell in row):
                raise ValueError('New workbook has invented progress or formula cells')
        finally:
            check.close()
        os.replace(temporary, destination); write(progress_path, progress(value))
    finally:
        temporary.unlink(missing_ok=True)
    return {'workbook': destination.relative_to(root).as_posix(), 'sha256': sha(destination),
        'historical_worksheets_preserved': len(original_names), 'worksheets': len(original_names) + len(SHEETS),
        'identities': 389, 'open_workstreams': 17, 'source_rows': len(value['source_rows']), 'native_rows': len(value['native_rows']),
        'trackingSha256': sha(root / 'docs/R29_TRACKING.json'), 'progressSha256': sha(progress_path),
        'baselineSha256': sha(baseline), 'canonical_fidelity_certified': 0}


def overview(value, checks):
    stats = {row['label']: row['value'] for row in value['statistics']}; total = checks['totals']
    return f'''# R29 — Corps, poses et images natives

Le [classeur R29](BOW_FIDELITE_R29.xlsx) ajoute sept feuilles aux 42 feuilles R28, conservées avec leurs relations. Le [suivi central](../modules/fidelity-r29/index.html) expose les 17 chantiers ouverts et les 389 identités.

R29 compte {stats['versions remplacées en R29']} versions corporelles remplacées et {stats['poses requalifiées en R29']} poses existantes requalifiées. Le cumul reste de 363 identités avec PNG actif et 22 refus conservés. {stats['corps anatomiques entiers revus']} corps anatomiques sont revus et {stats['attentes statiques admissibles']} attentes statiques sont admissibles. L’équipement et les autres phases gardent leurs critères propres.

{stats['variantes documentaires d’autre édition']} versions documentaires appartiennent à une autre édition et ne remplacent pas le PNG actif. Les interprétations HD/PC2013 des formes Ooze restent distinctes des incarnations Nintendo3DS2012 ; elles ne ferment aucun corps cible de cette édition.

Les anciennes interprétations R28 de Wall Blister et Scagdead conservent leurs pixels et leurs observations corporelles. Leurs notes R29 précisent désormais la source HD/PC (2013) et l’équivalence non vérifiée avec Nintendo 3DS (2012) ; aucun idle ni nouvelle revue de pose n’en est déduit.

Les modèles sources donnent {stats['images des clés natives']} images de clés et {stats['vues de pose native']} vues de pose. Les éditions et limites de rendu sont explicites. Les codes de clips restent non attribués ; ces observations ne certifient aucun cycle de combat jouable, timing officiel ou image originale du moteur du jeu. Les sources candidates de mauvaise phase ou édition restent hors production.

Les Dossiers relient uniquement les nouvelles preuves du même UID, sans charger les PNG automatiquement. Le lecteur permet de comparer les images natives et leurs références. Les commandes, évolutions, sons et sauvegardes R28 conservent leur périmètre historique.

La recette R29 rassemble {total['python_tests']} tests Python, {total['targeted_node_groups']} groupes Node et {total['browser_groups']} groupes navigateur réussis. État des rapports : {'complet' if checks['complete'] else 'provisoire, avec rapports manquants ou en échec explicites'}. Les résultats historiques ne sont pas recomptés comme exécutés en R29.

{stats['corps de combattants encore ouverts']} corps de combattants restent ouverts. Les rencontres contextuelles, les refus, les accessoires et les cycles complets restent séparés dans le suivi. La fidélité globale 1:1 reste non certifiée.

Preuves : [poses](../modules/fidelity-r29/sprites/index.html), [images natives](../modules/fidelity-r29/motion/index.html), [recette](R29_VALIDATION.json), [suivi structuré](R29_TRACKING.json), [reçu d’export](R29_EXPORT.json), [reproduction](R29_REPRODUCTION.json).

```sh
python3 scripts/fidelity-progress-r29.py build
python3 scripts/fidelity-progress-r29.py check
python3 scripts/fidelity-progress-r29.py export
python3 scripts/fidelity-progress-r29.py reproduce
```
'''.encode()


def reproduce(value, root=ROOT):
    root = Path(root).resolve(); checks = []
    try:
        for seed in ('17', '73'):
            result = subprocess.run([sys.executable, 'scripts/fidelity-progress-r29.py', 'check'], cwd=root,
                env={**os.environ, 'PYTHONHASHSEED': seed}, capture_output=True, text=True)
            if result.returncode:
                raise ValueError(result.stderr)
        checks.append({'name': 'Tracking is stable across independent Python hash seeds', 'pass': True})
        folder = root / 'qa-r29-reproduction'; folder.mkdir(exist_ok=True)
        baseline = root / 'docs/BOW_FIDELITE_R28.xlsx'; before = sha(baseline)
        with tempfile.TemporaryDirectory(dir=folder) as name:
            temporary = Path(name); result = export(value, root, temporary / 'reproduced.xlsx', temporary / 'progress.json')
            if result['sha256'] != sha(root / 'docs/BOW_FIDELITE_R29.xlsx'):
                raise ValueError('Actual workbook bytes do not reproduce')
            checks.append({'name': 'Actual R29 workbook reproduces exactly', 'pass': True, 'sha256': result['sha256']})
            if (temporary / 'progress.json').read_bytes() != (root / MODULE / 'progress.json').read_bytes() or sha(baseline) != before:
                raise ValueError('Progress or historical baseline changed')
            checks.append({'name': 'Actual progress matches and baseline bytes remain unchanged', 'pass': True})
    except Exception as error:
        checks.append({'name': 'Reproduction failure', 'pass': False, 'error': str(error)})
    return {'schema': 're.bow.fidelity-reproduction-r29/1', 'passed': sum(row['pass'] for row in checks),
            'failed': sum(not row['pass'] for row in checks), 'checks': checks}


def main():
    parser = argparse.ArgumentParser(description=__doc__); parser.add_argument('action', choices=('build', 'check', 'export', 'reproduce'))
    parser.add_argument('--allow-pending', action='store_true'); args = parser.parse_args()
    if args.action in ('build', 'check'):
        checks, value = build(allow_pending=args.allow_pending)
        outputs = {'docs/R29_VALIDATION.json': raw(checks), 'docs/R29_TRACKING.json': raw(value),
            'docs/R29_WORKSTREAMS.json': raw({key: value[key] for key in ('release', 'summary', 'tasks', 'generation_preference')}),
            (MODULE / 'progress.json').as_posix(): raw(progress(value)), 'docs/FIDELITE_R29.md': overview(value, checks)}
        for path, content in outputs.items():
            if args.action == 'check':
                if (ROOT / path).read_bytes() != content:
                    raise ValueError('R29 output differs from its evidence: ' + path)
            else:
                write(ROOT / path, content)
        print(json.dumps({'complete': checks['complete'], 'totals': checks['totals'], 'sourceRows': len(value['source_rows'])}))
    elif args.action == 'export':
        value = read('docs/R29_TRACKING.json'); result = export(value); write(ROOT / 'docs/R29_EXPORT.json', result); print(json.dumps(result))
    else:
        result = reproduce(read('docs/R29_TRACKING.json')); write(ROOT / 'docs/R29_REPRODUCTION.json', result); print(json.dumps(result)); return 1 if result['failed'] else 0
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
