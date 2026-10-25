#!/usr/bin/env python3
"""Export collected evidence to Excel without changing the historical tracker.

Uses openpyxl from the cloud tooling Python environment. This is an evidence
workbook, not a declaration of completed sprites or canonical fidelity.
"""
import argparse
import hashlib
import json
import os
from pathlib import Path
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
CREATOR = 'BOW collected source evidence R21'


def read(path):
    return json.loads(Path(path).read_text())


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def safe(value):
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def sheet(book, name, headers, rows):
    page = book.create_sheet(name)
    page.append(headers)
    for row in rows:
        page.append([safe(value) for value in row])
    page.freeze_panes = 'B2'
    page.auto_filter.ref = page.dimensions
    for cell in page[1]:
        cell.fill = PatternFill('solid', fgColor='343840')
        cell.font = Font(color='FFFFFF', bold=True)
    for row in page.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, str) and cell.value.startswith('https://') and '\n' not in cell.value:
                cell.hyperlink = cell.value
                cell.style = 'Hyperlink'
    for column in range(1, len(headers) + 1):
        page.column_dimensions[get_column_letter(column)].width = 35 if column < 4 else 58
    return page


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--evidence-root', type=Path, default=ROOT / 'qa-source-research-r21')
    parser.add_argument('--pack', type=Path, default=ROOT / 'qa-source-references-r21')
    parser.add_argument('--motion', type=Path)
    parser.add_argument('--additional', type=Path)
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    pages = read(args.evidence_root / 'catalog-source-evidence.json')
    images = read(args.evidence_root / 'catalog-image-evidence.json')
    pack = read(args.pack / 'manifest.json')
    if pages['schema'] != 'bow.catalog-source-evidence/1' or pack['schema'] != 're.bow.source-reference-pack/1':
        raise ValueError('Unexpected evidence or pack schema')
    records = {item['id']: item for item in pages['identities']}
    if len(records) != 389 or len(records) != len(pages['identities']):
        raise ValueError('Expected all 389 distinct catalogue identities')
    roster = read(ROOT / 'data/roster.json')
    tracking = read(ROOT / 'docs/R20_TRACKING.json')
    old_art = {item['id'] for item in read(ROOT / 'assets/sprites/bow/manifest.json')['sprites']}
    imported = {item['id']: item for item in pack['entries']}
    if not set(imported).issubset(records):
        raise ValueError('Imported reference outside the catalogue')
    for item in imported.values():
        image = (args.pack / item['image']).resolve(strict=True)
        if not image.is_relative_to(args.pack.resolve()) or sha(image) != item['sha256']:
            raise ValueError('Imported reference bytes differ from the pack manifest')
    motion = read(args.motion) if args.motion else {}
    additional = read(args.additional) if args.additional else {}
    extra_by_id = {item['id']: item.get('additional_candidates', []) for item in additional.get('identities', [])}
    reviewed = {item['id']: item for item in motion.get('subjects', [])}
    book = Workbook()
    book.remove(book.active)
    book.properties.creator = CREATOR
    book.properties.title = 'B.O.W. Protocol — preuves collectées R21'
    metrics = [
        ('Portée', 'Collecte et import de références statiques', 'Aucun chantier artistique déclaré terminé.'),
        ('Identités documentées', len(records), 'Une page source collectée par entrée ; pas une validation canonique.'),
        ('Pages du catalogue collectées', pages['summary']['usable_documents'], 'Documents HTTP 200, empreintes conservées.'),
        ('Images sources téléchargées', images['summary']['images_downloaded_and_decoded'], 'Décodages et empreintes contrôlés.'),
        ('Références statiques importées', len(imported), 'Pixels conservés, sans miroir ni animation inventée.'),
        ('Références de corps', pack['counts']['fighter_body_references'], 'Ne prouve pas que le corps soit complet ou latéral.'),
        ('Références de contextes', pack['counts']['context_scene_references'], 'Scènes distinctes des corps combattants.'),
        ('PNG R20 restant à produire', tracking['spriteProduction']['remainingIdentitiesWithoutR20Png'], 'Compteur historique conservé ; une référence importée ne clôture pas ce travail.'),
        ('Lots globaux clos', 0, 'Les 15 lots globaux restent ouverts.'),
        ('Fidélité 1:1 certifiée', 0, 'Autres vues, rigs, gestes, sons et revue par incarnation restent requis.'),
        ('Références croisées examinées', len(reviewed), 'Dossiers individuels ; ne pas confondre avec les mentions candidates.'),
        ('Origine du suivi', 'docs/R20_TRACKING.json', sha(ROOT / 'docs/R20_TRACKING.json')),
    ]
    sheet(book, 'Synthèse', ['Indicateur', 'Valeur', 'Portée'], metrics)
    rows = []
    for identity in roster['catalog']:
        uid = identity['id']
        page, asset = records[uid], imported.get(uid)
        rows.append([
            uid, identity['name'], identity['game'], page['source_page'], page['http_status'],
            page['page_title'], page['body_sha256'],
            'IMPORTÉE — RÉFÉRENCE STATIQUE' if asset else 'ATTRIBUTION À REVOIR',
            asset['reference_kind'] if asset else '', asset['source_image']['url'] if asset else '',
            asset['image'] if asset else '', asset['sha256'] if asset else '',
            'PRODUIT R20' if uid in old_art else 'À PRODUIRE R20',
            len(reviewed.get(uid, {}).get('verifiedPages', [])), len(extra_by_id.get(uid, [])),
            'NON CERTIFIÉ — rig, autres vues, mouvements et sons à produire',
        ])
    sheet(book, 'Sources 389', ['UID', 'Identité', 'Jeu', 'Page collectée', 'HTTP', 'Titre', 'SHA-256 page',
                                'Import source', 'Nature', 'Image source', 'Fichier du pack', 'SHA-256 image',
                                'Statut PNG historique', 'Pages du dossier examiné', 'Mentions candidates',
                                'Reste à vérifier'], rows)
    cross = []
    for subject in motion.get('subjects', []):
        for source in subject['verifiedPages']:
            cross.append([subject['id'], subject['game'], source['url'], source['title'], source['httpStatus'],
                          source['snapshotSha256'], source.get('excerpt', ''), 'Source consultée ; fidélité finale non certifiée'])
    sheet(book, 'Dossiers examinés', ['UID', 'Jeu', 'URL', 'Titre', 'HTTP', 'SHA-256', 'Extrait', 'Limite'], cross)
    candidate_rows = []
    for identity in additional.get('identities', []):
        for candidate in identity.get('additional_candidates', []):
            candidate_rows.append([identity['id'], identity['game'], candidate['url'], candidate['title'],
                                   candidate['body_sha256'], candidate.get('match_level', ''),
                                   candidate.get('article_excerpt', ''), 'Mention documentaire ; phase et incarnation à revoir'])
    sheet(book, 'Pistes documentaires', ['UID', 'Jeu', 'URL', 'Titre', 'SHA-256', 'Rapprochement', 'Extrait', 'Limite'], candidate_rows)
    footage = motion.get('downloadedNativeMotion')
    footage_rows = []
    if footage:
        footage_rows.append(['re3r-nemesis', footage['url'], footage['file'], footage['sha256'],
                             footage['format'], footage['frames'], footage['durationMs'],
                             'Infection d’un zombie ; décor et caméra présents, sans audio. Pas un sprite de combat.'])
    sheet(book, 'Séquence originale', ['UID', 'URL', 'Fichier', 'SHA-256', 'Format', 'Images', 'Durée ms', 'Portée'], footage_rows)
    sheet(book, 'Travaux ouverts', ['ID', 'Chantier', 'Statut historique', 'Reste à produire'],
          [[item['id'], item['title'], item['status'], item['remaining']]
           for item in tracking['tasks'] if item['status'] != 'VALIDÉ'])
    output = args.output or args.pack / 'BOW_REFERENCES_R21.xlsx'
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        existing = load_workbook(output, read_only=True)
        creator = existing.properties.creator
        existing.close()
        if creator != CREATOR:
            raise ValueError('Refusing to overwrite a workbook created outside this exporter')
    with tempfile.NamedTemporaryFile(dir=output.parent, suffix='.xlsx', delete=False) as temporary:
        temporary_name = Path(temporary.name)
    try:
        book.save(temporary_name)
        check = load_workbook(temporary_name, read_only=True, data_only=True)
        assert check['Sources 389'].max_row == 390
        assert check['Travaux ouverts'].max_row == 18
        imported_rows = check['Sources 389'].iter_rows(min_row=2, min_col=8, max_col=8, values_only=True)
        assert sum(row[0].startswith('IMPORTÉE') for row in imported_rows) == len(imported)
        check.close()
        os.replace(temporary_name, output)
    finally:
        temporary_name.unlink(missing_ok=True)
    print(json.dumps({'workbook': str(output), 'identities': len(records), 'imported_static_references': len(imported),
                      'documentary_candidates': len(candidate_rows), 'unfinished_tasks': 17,
                      'canonical_fidelity_certified': 0}, ensure_ascii=False))


if __name__ == '__main__':
    main()
