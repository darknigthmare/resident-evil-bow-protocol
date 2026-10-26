"""R28 export rejects invented scope, preserves history and quotes spreadsheet text."""
import copy
import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('export_r28', ROOT / 'scripts/export-fidelity-progress-r28.py')
EXPORT = importlib.util.module_from_spec(SPEC); SPEC.loader.exec_module(EXPORT)


class ExportR28Tests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(); self.addCleanup(temporary.cleanup); self.root = Path(temporary.name)
        (self.root / 'docs').mkdir(); (self.root / 'data').mkdir()
        self.proof = self.root / 'docs/proof.json'; self.proof.write_text('{"observed":true}\n')
        self.baseline = self.root / 'docs/baseline.xlsx'
        book = Workbook(); book.properties.creator = 'BOW fidelity progress R27'
        book.active.title = 'Literal history'; book.active.append(['Existing formula', '=2+2']); book.active['A1'].hyperlink = 'https://example.invalid/old'
        page = book.create_sheet('Catalogue R27'); page.append(['UID', 'Original name'])
        for index in range(389): page.append([str(index), 'Identity ' + str(index)])
        book.save(self.baseline); book.close()
        tasks = [{'id': 'TASK-' + str(index), 'title': 'Workstream ' + str(index), 'status': 'EN COURS', 'category': 'sources',
                  'advances': ['Source checked'], 'remaining': ['Original animations open'],
                  'evidence': [{'href': '../../docs/proof.json', 'label': 'Proof'}]} for index in range(17)]
        (self.root / 'docs/R27_TRACKING.json').write_text(json.dumps({'tasks': tasks}))
        (self.root / 'data/roster.json').write_text(json.dumps({'catalog': [{'id': str(index)} for index in range(389)]}))
        updates = []
        for index in range(389):
            item = {key: '' for key in EXPORT.FIELDS}; item.update(id=str(index), sprite_generated=False,
                sprite_full_body_reviewed=False, sprite_combat_eligible=False, sprite_available_states=[],
                r28_equipment_complete=False, r28_known_appearance_gaps=[])
            updates.append(item)
        self.tracking = {'schema': 're.bow.parallel-workstreams/3', 'release': 'R28', 'canonical_fidelity_certified': 0,
            'summary': 'Still open', 'statistics': [{'label': 'Open', 'value': 17, 'scope': 'No certification'}], 'tasks': tasks,
            'resources': [{'href': '../../docs/proof.json', 'label': 'Proof'}], 'identity_updates': updates,
            'source_rows': [['TASK-0', '0', 'Quote', ' \t=HYPERLINK("https://bad.invalid")', 'Open', '@external-command']],
            'validation': [{'label': 'Fixture', 'result': 'Pass', 'scope': 'Synthetic evidence', 'path': 'docs/proof.json', 'sha256': self.sha(self.proof)}],
            'inputs': []}
        self.bind(self.proof); self.bind(self.baseline)

    def sha(self, path): return hashlib.sha256(path.read_bytes()).hexdigest()

    def bind(self, path):
        relative = path.relative_to(self.root).as_posix()
        self.tracking['inputs'] = [item for item in self.tracking['inputs'] if item['path'] != relative]
        self.tracking['inputs'].append({'path': relative, 'sha256': self.sha(path)})

    def export(self, destination='docs/output.xlsx'):
        tracking = self.root / 'docs/tracking.json'; tracking.write_text(json.dumps(self.tracking))
        return EXPORT.export(tracking, self.baseline, self.root / destination, self.root / 'modules/fidelity-r28/progress.json', self.root)

    def test_preserves_worksheet_xml_relationships_and_quotes_new_formulas(self):
        original_bytes = self.baseline.read_bytes(); result = self.export(); output = self.root / result['workbook']
        self.assertEqual(self.baseline.read_bytes(), original_bytes)
        with ZipFile(self.baseline) as original, ZipFile(output) as current:
            for name in original.namelist():
                if name.startswith('xl/worksheets/'): self.assertEqual(original.read(name), current.read(name), name)
        book = load_workbook(output)
        self.assertEqual(book['Literal history']['B1'].value, '=2+2'); self.assertEqual(book['Literal history']['A1'].hyperlink.target, 'https://example.invalid/old')
        self.assertEqual(book['Catalogue R28'].max_row, 390); self.assertEqual(book['Chantiers R28'].max_row, 18)
        self.assertTrue(book['Sources R28']['D2'].value.startswith("' \t=")); self.assertTrue(book['Sources R28']['F2'].value.startswith("'@"))
        self.assertEqual(book['Sources R28']['D2'].data_type, 's'); self.assertEqual(len(book.sheetnames), 9); book.close()

    def test_stale_input_is_rejected_before_creating_outputs(self):
        self.proof.write_text('{"observed":false}\n')
        with self.assertRaisesRegex(ValueError, 'Evidence hash mismatch'): self.export()
        self.assertFalse((self.root / 'docs/output.xlsx').exists())

    def test_closed_workstream_or_missing_identity_is_rejected(self):
        self.tracking['tasks'][0]['status'] = 'VALIDÉ'
        with self.assertRaisesRegex(ValueError, 'Open workstream'): self.export()
        self.tracking['tasks'][0]['status'] = 'EN COURS'; self.tracking['identity_updates'].pop()
        with self.assertRaisesRegex(ValueError, '389 identities'): self.export()

    def test_historical_and_foreign_workbooks_are_not_overwritten(self):
        with self.assertRaisesRegex(ValueError, 'overwrite historical'): self.export('docs/baseline.xlsx')
        book = Workbook(); book.properties.creator = 'Another author'; book.save(self.root / 'docs/output.xlsx'); book.close()
        with self.assertRaisesRegex(ValueError, 'another author'): self.export()

    def test_exact_outputs_are_reproducible(self):
        first = self.export('docs/first.xlsx'); second = self.export('docs/second.xlsx')
        self.assertEqual((self.root / first['workbook']).read_bytes(), (self.root / second['workbook']).read_bytes())

    def test_binding_pose_or_tool_refusal_cannot_gain_combat_idle(self):
        folder = self.root / 'modules/fidelity-r28/sprites'; folder.mkdir(parents=True)
        jobs_folder = self.root / 'modules/fidelity-r27/sprites'; jobs_folder.mkdir(parents=True)
        (jobs_folder / 'jobs.json').write_text(json.dumps({'jobs': [{'id': str(i)} for i in range(385)]}))
        image = folder / 'fixture.png'; image.write_bytes(b'synthetic file integrity fixture')
        sprite = {'id': '0', 'image': image.relative_to(self.root).as_posix(), 'bytes': image.stat().st_size, 'sha256': self.sha(image),
            'reviewStatus': 'visually-reviewed', 'fullBodyReviewed': True, 'completeBodyReferenceVerified': True,
            'scope': 'combatant-neutral-pose', 'sourceFacing': 'front', 'poseKind': 'source-bind-pose', 'availableStates': [],
            'canMirror': False, 'canonicalFidelityCertified': False, 'combatIntegrated': False}
        pending = [{'id': str(i), 'generationStatus': 'BLOCKED_BY_TOOL', 'reason': 'Recorded refusal'} for i in range(1, 385)]
        pack = {'schema': 're.bow.generated-combat-sprites/1', 'sprites': [sprite], 'pending': pending,
            'counts': {'targetIdentities': 385, 'generatedSprites': 1, 'pendingSprites': 384, 'visuallyReviewedSprites': 1,
                'fullBodyReviewedSprites': 1, 'neutralBaseCombatEligibleSprites': 0, 'generationBlockedSprites': 384,
                'originalGameSprites': 0, 'completeCombatAnimationSets': 0, 'canonicalFidelityCertified': 0}}
        path = folder / 'manifest.json'; path.write_text(json.dumps(pack)); self.bind(path); self.tracking['sprite_manifest'] = path.relative_to(self.root).as_posix()
        item = self.tracking['identity_updates'][0]
        item.update(visual=sprite['image'], sprite_generated=True, sprite_review_status=sprite['reviewStatus'], sprite_full_body_reviewed=True,
            sprite_reference_full_body_verified=True, sprite_scope=sprite['scope'], sprite_facing='front', sprite_pose_kind='source-bind-pose', sprite_generation_status='GENERATED')
        for update in self.tracking['identity_updates'][1:385]: update.update(sprite_generation_status='BLOCKED_BY_TOOL', sprite_generation_reason='Recorded refusal')
        EXPORT.validate(copy.deepcopy(self.tracking), self.root)
        item['sprite_combat_eligible'] = True
        with self.assertRaisesRegex(ValueError, 'sprite qualification'): EXPORT.validate(self.tracking, self.root)
        item['sprite_combat_eligible'] = False; self.tracking['identity_updates'][1]['sprite_generated'] = True
        with self.assertRaisesRegex(ValueError, 'refusal cannot become'): EXPORT.validate(self.tracking, self.root)


if __name__ == '__main__':
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(ExportR28Tests)
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    paths = ['scripts/export-fidelity-progress-r28.py', 'tests/test_fidelity_export_r28.py']
    report = {'schema': 're.bow.export-tests-r28/1', 'passed': result.testsRun - len(result.failures) - len(result.errors),
              'failed': len(result.failures) + len(result.errors), 'testsRun': result.testsRun,
              'sourceFiles': [{'path': path, 'sha256': hashlib.sha256((ROOT / path).read_bytes()).hexdigest()} for path in paths],
              'scope': 'Synthetic integrity fixtures test historical preservation, SHA drift, formula text, refusal/pose gates and reproducible exports.'}
    (ROOT / 'docs/R28_EXPORT_TESTS.json').write_text(json.dumps(report, indent=2) + '\n')
    raise SystemExit(0 if result.wasSuccessful() else 1)
