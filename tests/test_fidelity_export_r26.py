"""R26 export: preserve history, verify evidence, and keep native clocks distinct."""
import copy
import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('export_r26', ROOT / 'scripts/export-fidelity-progress-r26.py')
EXPORT = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(EXPORT)


class FidelityExportR26Tests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        for folder in ('docs', 'data', 'modules/fidelity-r26/motion'):
            (self.root / folder).mkdir(parents=True)
        self.proof = self.root / 'docs/proof.json'
        self.proof.write_text('{"observed": true}\n')
        self.baseline = self.root / 'docs/baseline.xlsx'
        book = Workbook()
        book.properties.creator = 'BOW fidelity progress R25'
        page = book.active
        page.title = 'Historical controls'
        page.append(['Literal historic formula', '=2+2'])
        page['A1'].font = Font(name='Arial', bold=True, color='BB3311')
        page['A1'].fill = PatternFill('solid', fgColor='ECE0D0')
        page['A2'] = 'Existing hyperlink'
        page['A2'].hyperlink = 'https://example.invalid/historical'
        page.freeze_panes = 'B2'
        for title in ('Catalogue 389', 'Catalogue R25'):
            page = book.create_sheet(title)
            page.append(['UID', 'Name', 'R25 evidence'])
            for i in range(389):
                page.append([str(i), 'Historical identity ' + str(i), 'Preserved R25'])
        book.save(self.baseline)
        book.close()
        tasks = [{'id': 'TASK-' + str(i), 'title': 'Open workstream ' + str(i), 'status': 'EN COURS',
                  'category': 'sources', 'advances': ['Reviewed evidence'], 'remaining': ['Production still open'],
                  'evidence': [{'href': '../../docs/proof.json', 'label': 'Proof'}]} for i in range(17)]
        (self.root / 'docs/R20_TRACKING.json').write_text(json.dumps({'tasks': tasks}))
        (self.root / 'data/roster.json').write_text(json.dumps({'catalog': [{'id': str(i)} for i in range(389)]}))
        self.tracking = {
            'schema': 're.bow.parallel-workstreams/2', 'canonical_fidelity_certified': 0,
            'summary': 'All production criteria remain open.', 'statistics': [{'label': 'Open workstreams', 'value': 17}],
            'resources': [{'href': '../../docs/proof.json', 'title': 'Proof', 'label': 'Read evidence'}],
            'tasks': tasks,
            'identity_updates': [{'id': str(i), 'visual': '', 'rig': '', 'motion_clips': 0, 'audio': 0,
                                  'claims': 0, 'limitations': ['Not certified']} for i in range(389)],
            'source_rows': [['TASK-0', '0', 'Quote', ' \t=HYPERLINK("https://bad.invalid")', 'Unverified', '@external-command']],
            'validation': [{'label': 'Synthetic export check', 'result': 'Passed', 'scope': 'Test fixture only', 'path': 'docs/proof.json'}],
            'inputs': [],
        }
        self.record_hash(self.proof)
        self.record_hash(self.baseline)

    def record_hash(self, path):
        relative = path.relative_to(self.root).as_posix()
        self.tracking['inputs'] = [item for item in self.tracking['inputs'] if item['path'] != relative]
        self.tracking['inputs'].append({'path': relative, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})

    def export(self, output=None):
        tracking_path = self.root / 'docs/tracking.json'
        tracking_path.write_text(json.dumps(self.tracking))
        return EXPORT.export(tracking_path, self.baseline, output or self.root / 'docs/output.xlsx',
                             self.root / 'modules/fidelity-r26/progress.json', self.root)

    def add_motion(self):
        clip = {'id': 'native-clip', 'member': 'Animations/native.FBX', 'sha256': 'a' * 64,
                'keyTimeStart': '-1', 'keyTimeEnd': '99', 'ticksPerSecond': 100,
                'sourceTimeSpanSeconds': 1.0, 'nativePreRollSeconds': 0.01, 'lastNativeTimeSeconds': 0.99,
                'varyingAnimationCurveCount': 3, 'varyingNonnegativeAnimationCurveCount': 2,
                'positiveTimePoseOnly': False, 'playbackVariationScope': 'nonnegative-time-motion',
                'semanticActionVerified': False}
        manifest = {'schema': 're.bow.rigged-motion-pack/1', 'summary': {'sourceAnimationFiles': 1,
                    'positiveTimeMotionFiles': 1, 'positiveTimePoseFiles': 0, 'canonicalFidelityCertified': 0},
                    'entries': [{'id': '0', 'game': 'Synthetic metadata fixture', 'rigProductionFinalized': False,
                                 'canonicalFidelityCertified': False, 'clips': [clip]}]}
        path = self.root / 'modules/fidelity-r26/motion/manifest.json'
        path.write_text(json.dumps(manifest))
        self.tracking['motion_manifest'] = path.relative_to(self.root).as_posix()
        self.record_hash(path)
        return path, manifest

    def test_export_preserves_all_old_sheet_bytes_and_quotes_new_formula_text(self):
        before = self.baseline.read_bytes()
        result = self.export()
        self.assertEqual(self.baseline.read_bytes(), before)
        with ZipFile(self.baseline) as original, ZipFile(result['workbook']) as output:
            for name in original.namelist():
                if name.startswith('xl/worksheets/'):
                    self.assertEqual(original.read(name), output.read(name), name)
        book = load_workbook(result['workbook'])
        self.assertEqual(book['Historical controls']['B1'].value, '=2+2')
        self.assertEqual(book['Historical controls']['A1'].font.color.rgb, '00BB3311')
        self.assertEqual(book['Catalogue R25']['C2'].value, 'Preserved R25')
        self.assertEqual(book['Catalogue R26']['C2'].value, 'Preserved R25')
        self.assertEqual(book['Catalogue R26'].max_row, 390)
        self.assertIn('Clips audio documentaires R26', [cell.value for cell in book['Catalogue R26'][1]])
        self.assertEqual(book['Chantiers R26'].max_row, 18)
        self.assertEqual(book['Sources R26']['D2'].data_type, 's')
        self.assertTrue(book['Sources R26']['D2'].value.startswith("' \t="))
        self.assertTrue(book['Sources R26']['F2'].value.startswith("'@"))
        self.assertEqual(book['Catalogue R26']['I2'].value, '["Not certified"]')
        book.close()
        progress = json.loads(Path(result['progress']).read_text())
        self.assertEqual(progress['schema'], 're.bow.fidelity-progress/1')
        self.assertEqual(len(progress['tasks']), 17)

    def test_changed_evidence_is_rejected_before_any_delivery(self):
        self.proof.write_text('{"observed": false}\n')
        with self.assertRaisesRegex(ValueError, 'hash mismatch'):
            self.export()
        self.assertFalse((self.root / 'docs/output.xlsx').exists())
        self.assertFalse((self.root / 'modules/fidelity-r26/progress.json').exists())

    def test_missing_or_duplicate_workstream_and_false_closure_are_rejected(self):
        original = copy.deepcopy(self.tracking)
        self.tracking['tasks'].pop()
        with self.assertRaisesRegex(ValueError, 'workstreams'):
            self.export()
        self.tracking = copy.deepcopy(original)
        self.tracking['tasks'][-1] = copy.deepcopy(self.tracking['tasks'][0])
        with self.assertRaisesRegex(ValueError, 'workstreams'):
            self.export()
        self.tracking = original
        self.tracking['tasks'][0]['status'] = 'VALIDÉ'
        with self.assertRaisesRegex(ValueError, 'closed'):
            self.export()

    def test_unknown_missing_or_duplicate_identity_cannot_gain_progress(self):
        original = copy.deepcopy(self.tracking)
        for mutation in ('unknown', 'missing', 'duplicate'):
            self.tracking = copy.deepcopy(original)
            if mutation == 'unknown':
                self.tracking['identity_updates'][0]['id'] = 'unknown-incarnation'
            elif mutation == 'missing':
                self.tracking['identity_updates'].pop()
            else:
                self.tracking['identity_updates'][-1] = copy.deepcopy(self.tracking['identity_updates'][0])
            with self.assertRaisesRegex(ValueError, 'identity|identities'):
                self.export()

    def test_global_certification_and_unrecognized_identity_fields_are_rejected(self):
        self.tracking['canonical_fidelity_certified'] = 1
        with self.assertRaisesRegex(ValueError, 'certification'):
            self.export()
        self.tracking['canonical_fidelity_certified'] = 0
        self.tracking['identity_updates'][0]['rig_final_certified'] = True
        with self.assertRaisesRegex(ValueError, 'Unknown identity progress field'):
            self.export()

    def test_evidence_links_and_output_paths_cannot_escape_checkout_or_overwrite_history(self):
        self.tracking['tasks'][0]['evidence'][0]['href'] = '../../../../outside.json'
        with self.assertRaisesRegex(ValueError, 'inside the checkout'):
            self.export()
        self.tracking['tasks'][0]['evidence'][0]['href'] = '../../docs/proof.json'
        (self.root / 'tests').mkdir()
        (self.root / 'tests/proof.mjs').write_text('// Local code is not a review resource')
        self.tracking['tasks'][0]['evidence'][0]['href'] = '../../tests/proof.mjs'
        with self.assertRaisesRegex(ValueError, 'modules or docs'):
            self.export()
        self.tracking['tasks'][0]['evidence'][0]['href'] = '../../docs/proof.json'
        with self.assertRaisesRegex(ValueError, 'inside the checkout'):
            self.export(self.root.parent / 'escaped.xlsx')
        with self.assertRaisesRegex(ValueError, 'overwrite historical'):
            self.export(self.baseline)

    def test_baseline_requires_the_expected_author_and_a_recorded_hash(self):
        self.tracking['inputs'] = [item for item in self.tracking['inputs'] if item['path'] != 'docs/baseline.xlsx']
        with self.assertRaisesRegex(ValueError, 'recorded evidence hash'):
            self.export()
        book = load_workbook(self.baseline)
        book.properties.creator = 'Unrelated workbook'
        book.save(self.baseline)
        book.close()
        self.record_hash(self.baseline)
        with self.assertRaisesRegex(ValueError, 'historical R25'):
            self.export()

    def test_native_min_to_max_span_preroll_and_zero_to_max_playback_remain_distinct(self):
        self.add_motion()
        result = self.export()
        self.assertEqual(result['source_clips'], 1)
        book = load_workbook(result['workbook'])
        page = book['Clips sources R26']
        self.assertEqual(page['G2'].value, -0.01)
        self.assertEqual(page['H2'].value, 0.99)
        self.assertEqual(page['I2'].value, 1)
        self.assertEqual(page['J2'].value, 0.01)
        self.assertEqual(page['K2'].value, 0.99)
        self.assertEqual(page['L2'].value, 3)
        self.assertEqual(page['M2'].value, 2)
        self.assertEqual(page['N2'].value, 'MOUVEMENT APRÈS ZÉRO')
        self.assertEqual(page['P2'].value, 'ACTION NON ATTRIBUÉE')
        book.close()

    def test_associated_object_remains_a_distinct_resource_under_its_documented_subject(self):
        path, manifest = self.add_motion()
        entry = manifest['entries'][0]
        entry.update(id='source-trap-reference', subjectId='0', modelRole='associated-object')
        path.write_text(json.dumps(manifest))
        self.record_hash(path)
        result = self.export()
        book = load_workbook(result['workbook'])
        page = book['Clips sources R26']
        self.assertEqual(page['A2'].value, '0')
        self.assertEqual(page['R2'].value, 'source-trap-reference')
        self.assertEqual(page['S2'].value, 'associated-object')
        book.close()
        for changes in ({'subjectId': 'unknown'}, {'modelRole': 'body'}, {'modelRole': 'invented-role'}):
            original = copy.deepcopy(entry)
            entry.update(changes)
            path.write_text(json.dumps(manifest))
            self.record_hash(path)
            with self.assertRaisesRegex(ValueError, 'Unknown motion identity'):
                self.export()
            entry.clear()
            entry.update(original)

    def test_preroll_variation_does_not_turn_a_positive_time_pose_into_a_movement(self):
        path, manifest = self.add_motion()
        pose = copy.deepcopy(manifest['entries'][0]['clips'][0])
        pose.update(id='pose-after-zero', varyingNonnegativeAnimationCurveCount=0,
                    positiveTimePoseOnly=True, playbackVariationScope='pose-only-after-time-zero')
        manifest['entries'][0]['clips'].append(pose)
        manifest['summary'].update(sourceAnimationFiles=2, positiveTimeMotionFiles=1, positiveTimePoseFiles=1)
        path.write_text(json.dumps(manifest))
        self.record_hash(path)
        result = self.export()
        self.assertEqual(result['source_clips'], 2)
        self.assertEqual(result['positive_time_motion_files'], 1)
        self.assertEqual(result['positive_time_pose_files'], 1)
        book = load_workbook(result['workbook'])
        page = book['Clips sources R26']
        self.assertEqual(page['L3'].value, 3)
        self.assertEqual(page['M3'].value, 0)
        self.assertEqual(page['N3'].value, 'POSE FIXE APRÈS ZÉRO')
        self.assertEqual(page['O3'].value, 'pose-only-after-time-zero')
        self.assertEqual(page['J3'].value, 0.01)
        self.assertEqual(page['K3'].value, 0.99)
        self.assertIn('même si le pré-roll varie', page['Q3'].value)
        book.close()

    def test_contradictory_pose_flags_scope_curve_counts_and_summary_are_rejected(self):
        path, original = self.add_motion()
        for mutation in ('pose_flag', 'scope', 'too_many_curves', 'summary'):
            with self.subTest(mutation=mutation):
                manifest = copy.deepcopy(original)
                clip = manifest['entries'][0]['clips'][0]
                if mutation == 'pose_flag':
                    clip['positiveTimePoseOnly'] = True
                elif mutation == 'scope':
                    clip['playbackVariationScope'] = 'pose-only-after-time-zero'
                elif mutation == 'too_many_curves':
                    clip['varyingNonnegativeAnimationCurveCount'] = clip['varyingAnimationCurveCount'] + 1
                else:
                    manifest['summary'].update(positiveTimeMotionFiles=0, positiveTimePoseFiles=1)
                path.write_text(json.dumps(manifest))
                self.record_hash(path)
                with self.assertRaisesRegex(ValueError, 'classification|scope|totals'):
                    self.export()

    def test_inconsistent_native_clock_is_rejected_even_with_a_valid_manifest_hash(self):
        path, manifest = self.add_motion()
        manifest['entries'][0]['clips'][0]['sourceTimeSpanSeconds'] = 0.99
        path.write_text(json.dumps(manifest))
        self.record_hash(path)
        with self.assertRaisesRegex(ValueError, 'Native animation times disagree'):
            self.export()


if __name__ == '__main__':
    unittest.main()
