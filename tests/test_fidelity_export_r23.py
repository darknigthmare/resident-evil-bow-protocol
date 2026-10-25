"""R23 export trust boundaries: stale evidence, dropped scope and spreadsheet formulas."""
import copy
import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('r23_export', ROOT / 'scripts/export-fidelity-progress-r23.py')
EXPORT = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(EXPORT)


class FidelityExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        for directory in ('docs', 'data', 'modules/fidelity-r23'):
            (self.root / directory).mkdir(parents=True)
        self.evidence = self.root / 'docs/proof.json'
        self.evidence.write_text('{"measured": true}\n')
        tasks = [{'id': 'TASK-' + str(i), 'title': 'Open ' + str(i), 'status': 'EN COURS',
                  'category': 'sources', 'advances': ['Documented reference'], 'remaining': ['Production still open'],
                  'evidence': [{'label': 'Proof', 'href': '../../docs/proof.json'}]} for i in range(17)]
        (self.root / 'docs/R20_TRACKING.json').write_text(json.dumps({'tasks': tasks}))
        (self.root / 'data/roster.json').write_text(json.dumps({'catalog': [{'id': str(i)} for i in range(389)]}))
        self.tracking = {'schema': 're.bow.parallel-workstreams/1', 'canonical_fidelity_certified': 0,
                         'summary': 'Open production', 'tasks': tasks, 'statistics': [{'label': 'Open', 'value': 17}],
                         'resources': [{'label': 'Proof', 'title': 'Proof', 'href': '../../docs/proof.json'}],
                         'source_rows': [['TASK-0', '0', 'Quote', '=HYPERLINK("https://bad.example")', 'Unverified']],
                         'identity_updates': [{'id': '0', 'claims': 1}], 'validation': [],
                         'inputs': [{'path': 'docs/proof.json', 'sha256': hashlib.sha256(self.evidence.read_bytes()).hexdigest()}]}
        self.baseline = self.root / 'docs/baseline.xlsx'
        book = Workbook()
        book.properties.creator = 'BOW fidelity progress R22'
        page = book.create_sheet('Catalogue 389')
        page.append(['UID', 'Name'])
        for i in range(389):
            page.append([str(i), 'Historical identity'])
        book.save(self.baseline)

    def export(self):
        source = self.root / 'docs/tracking.json'
        source.write_text(json.dumps(self.tracking))
        return EXPORT.export(source, self.baseline, self.root / 'docs/output.xlsx',
                             self.root / 'modules/fidelity-r23/progress.json', self.root)

    def test_export_retains_baseline_and_quotes_untrusted_spreadsheet_text(self):
        before = self.baseline.read_bytes()
        result = self.export()
        book = load_workbook(result['workbook'])
        self.assertEqual(book['Catalogue 389']['B2'].value, 'Historical identity')
        self.assertEqual(book['Catalogue R23'].max_row, 390)
        self.assertEqual(book['Chantiers R23'].max_row, 18)
        self.assertEqual(book['Sources R23']['D2'].data_type, 's')
        self.assertTrue(book['Sources R23']['D2'].value.startswith("'="))
        self.assertEqual(self.baseline.read_bytes(), before)
        progress = json.loads(Path(result['progress']).read_text())
        self.assertEqual(progress['schema'], 're.bow.fidelity-progress/1')
        self.assertEqual(len(progress['tasks']), 17)

    def test_changed_evidence_is_rejected_before_publishing(self):
        self.evidence.write_text('{"measured": false}\n')
        with self.assertRaisesRegex(ValueError, 'hash mismatch'):
            self.export()
        self.assertFalse((self.root / 'docs/output.xlsx').exists())

    def test_missing_task_or_false_closure_is_rejected(self):
        original = copy.deepcopy(self.tracking)
        self.tracking['tasks'].pop()
        with self.assertRaisesRegex(ValueError, 'workstreams'):
            self.export()
        self.tracking = original
        self.tracking['tasks'][0]['status'] = 'VALIDÉ'
        with self.assertRaisesRegex(ValueError, 'closed'):
            self.export()

    def test_source_link_cannot_escape_checkout(self):
        self.tracking['tasks'][0]['evidence'][0]['href'] = '../../../../outside.json'
        with self.assertRaisesRegex(ValueError, 'inside the checkout'):
            self.export()

    def test_unknown_identity_cannot_gain_progress(self):
        self.tracking['identity_updates'][0]['id'] = 'unknown-incarnation'
        with self.assertRaisesRegex(ValueError, 'identity'):
            self.export()


if __name__ == '__main__':
    unittest.main()
