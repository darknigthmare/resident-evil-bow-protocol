"""Exact UID coverage, honest motion admission and byte-stable R29 history."""
import copy
import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('r30_progress_test', ROOT / 'scripts/fidelity-progress-r30.py')
EXPORT = importlib.util.module_from_spec(SPEC); SPEC.loader.exec_module(EXPORT)


class FidelityR30Tests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(); self.addCleanup(temporary.cleanup); self.root = Path(temporary.name)
        self.put('data/roster.json', {'catalog': [{'id': 'fixture-' + str(index)} for index in range(389)]})
        self.put('docs/R29_VALIDATION.json', {'historical': True})
        self.put(EXPORT.CONTEXT_GRAPH, {'contextOnlyUids': []})
        self.put('docs/R29_TRACKING.json', {
            'identity_updates': [{'id': 'fixture-' + str(index), 'sprite_generation_status': 'BLOCKED_BY_TOOL' if index < 22 else 'GENERATED'} for index in range(389)],
            'tasks': [{'id': 'task-' + str(index), 'title': 'Historical task', 'category': 'sources',
                       'remaining': ['Unverified criterion'], 'evidence': [{'href': 'index.html', 'label': 'History'}]} for index in range(17)],
            'generation_preference': 'Keep refusals; fidelity first'})
        self.put('modules/fidelity-r29/sprites/manifest.json', {'sprites': list(range(363)), 'pending': list(range(22))})
        (self.root / 'modules/fidelity-r29/index.html').write_text('Retained historical overview')
        for path in ('scripts/fidelity-progress-r30.py', 'modules/fidelity-r30/index.html', 'modules/fidelity-r30/dossier-link.js'):
            target = self.root / path; target.parent.mkdir(parents=True, exist_ok=True); target.write_bytes((ROOT / path).read_bytes())
        book = Workbook(); book.remove(book.active)
        for index in range(49):
            page = book.create_sheet('Catalogue R29' if index == 0 else 'History ' + str(index))
            if index == 0:
                page.append(['UID', 'Historical value'])
                for uid in range(389):
                    page.append(['fixture-' + str(uid), 'Retained historical detail'])
            elif index == 1:
                page['A1'] = 'Keep exact hyperlink'; page['A1'].hyperlink = 'https://example.invalid/history'
                page['B1'] = '=2+2'; page['B1'].data_type = 's'
        book.properties.creator = 'BOW fidelity progress R29'
        book.save(self.root / EXPORT.BASELINE); book.close()
        self.entry = {'id': 'fixture-native-pose', 'subjectId': 'fixture-0', 'kind': 'native-static-pose',
            'sourceEdition': 'Fixture correct edition', 'admittedTargetEdition': True, 'fullBodyReviewed': True,
            'fightingPoseReviewed': True, 'sourceFacing': 'right', 'canMirror': False,
            'actionIdentity': 'neutral-stance', 'actionVerified': False, 'continuityVerified': False,
            'admittedCombatStates': ['idle'], 'knownAppearanceGaps': [], 'canonicalFidelityCertified': False,
            'originalGameFrames': False, 'frames': [self.image('pose.png')]}
        self.lot = 'modules/fidelity-r30/native/other-games/manifest.json'
        self.install_native(self.entry)
        avatar = {'id': 'fixture-0', 'label': 'Fixture', 'game': 'Fixture game',
            'image': self.entry['frames'][0]['path'], 'bytes': self.entry['frames'][0]['bytes'],
            'sha256': self.entry['frames'][0]['sha256'], 'width': 2, 'height': 2,
            'sourceKind': 'secondary-reference', 'sourceEdition': 'Fixture correct edition',
            'canonicalFidelityCertified': False}
        self.put(EXPORT.AVATARS, {'schema': 're.bow.selection-avatars-r30/1', 'entries': [avatar],
            'missing': [{'id': 'fixture-' + str(index)} for index in range(1, 389)]})
        self.proof = 'docs/fixture-test.json'
        self.put(self.proof, {'passed': 1, 'failed': 0, 'inputs': [self.binding('data/roster.json')]})
        self.put(EXPORT.REPORT_INDEX, {'reports': [{'path': self.proof, 'label': 'Fixture recipe', 'kind': 'fixture_checks'}]})

    def put(self, path, value):
        target = self.root / path; target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(json.dumps(value, ensure_ascii=False) + '\n')

    def binding(self, path):
        target = self.root / path
        return {'path': path, 'bytes': target.stat().st_size, 'sha256': hashlib.sha256(target.read_bytes()).hexdigest()}

    def image(self, name):
        path = 'modules/fidelity-r30/native/other-games/' + name; target = self.root / path
        target.parent.mkdir(parents=True, exist_ok=True); target.write_bytes(b'Synthetic integrity fixture: ' + name.encode())
        return {**self.binding(path), 'width': 2, 'height': 2}

    def install_native(self, entry):
        self.put(self.lot, {'schema': 're.bow.native-2d-lot-r30/1', 'release': 'R30', 'entries': [entry]})
        self.put(EXPORT.NATIVE, {'schema': 're.bow.native-2d-index-r30/1', 'release': 'R30', 'lots': [self.binding(self.lot)]})

    def build(self):
        return EXPORT.build(self.root)

    def test_static_pose_does_not_fill_full_idle_gap(self):
        _, value = self.build(); row = value['identity_updates'][0]
        self.assertEqual(row['r30_native_combat_stances'], 1)
        self.assertEqual(row['r30_native_verified_idle_cycles'], 0)
        self.assertFalse(row['r30_native_combat_pending']); self.assertTrue(row['r30_full_idle_pending'])
        self.assertEqual(row['sprite_generation_status'], 'BLOCKED_BY_TOOL')

    def test_smooth_unassigned_wrong_edition_scene_stays_documentary(self):
        entry = copy.deepcopy(self.entry)
        entry.update(kind='native-sequence', actionIdentity='unassigned', admittedTargetEdition=False,
            admittedCombatStates=[], fullNativeDuration=True, sampleRate=60, durationSeconds=1,
            frames=[self.image('frame-' + str(index) + '.png') for index in range(60)])
        self.install_native(entry); _, value = self.build(); row = value['identity_updates'][0]
        self.assertEqual(row['r30_native_smooth_sequences'], 1)
        self.assertEqual(row['r30_native_verified_idle_cycles'], 0)
        self.assertTrue(row['r30_native_combat_pending']); self.assertEqual(row['r30_native_edition_variants'], 1)
        entry['admittedCombatStates'] = ['idle']; self.install_native(entry)
        with self.assertRaisesRegex(ValueError, 'documentary, unverified or sparse'):
            self.build()

    def test_true_idle_requires_action_continuity_and_full_source_duration(self):
        entry = copy.deepcopy(self.entry)
        entry.update(kind='native-sequence', actionIdentity='idle', actionVerified=True,
            continuityVerified=True, fullNativeDuration=True, sampleRate=60, durationSeconds=1,
            frames=[self.image('idle-' + str(index) + '.png') for index in range(60)])
        self.install_native(entry); _, value = self.build()
        self.assertEqual(value['identity_updates'][0]['r30_native_verified_idle_cycles'], 1)
        for key in ('actionVerified', 'continuityVerified', 'fullNativeDuration'):
            bad = copy.deepcopy(entry); bad[key] = False; self.install_native(bad)
            with self.assertRaisesRegex(ValueError, 'unverified or sparse'):
                self.build()
        bad = copy.deepcopy(entry); bad['frames'] = bad['frames'][:16]; self.install_native(bad)
        with self.assertRaisesRegex(ValueError, 'unverified or sparse'):
            self.build()
        bad = copy.deepcopy(entry); bad['durationSeconds'] = 10; self.install_native(bad)
        with self.assertRaisesRegex(ValueError, 'unverified or sparse'):
            self.build()

    def test_optional_observed_edition_loop_is_separate_from_target_idle(self):
        entry = copy.deepcopy(self.entry)
        entry.update(kind='native-sequence', actionIdentity='unassigned', admittedTargetEdition=False,
            admittedCombatStates=[], sourceActionCertified=False, observedNeutralLoopReviewed=True,
            continuityVerified=True, selectablePresentationVariant=True, variantCombatStates=['idle'],
            variantFamilyId='fixture-source-edition', fullNativeDuration=True, sampleRate=60, durationSeconds=1,
            frames=[self.image('observed-' + str(index) + '.png') for index in range(60)])
        self.install_native(entry); _, value = self.build(); row = value['identity_updates'][0]
        self.assertEqual(row['r30_native_optional_neutral_variants'], 1)
        self.assertEqual(row['r30_native_playable_optional_neutral_variants'], 1)
        self.assertEqual(row['r30_native_verified_idle_cycles'], 0)
        self.assertTrue(row['r30_full_idle_pending']); self.assertTrue(row['r30_native_combat_pending'])
        self.put(EXPORT.CONTEXT_GRAPH, {'contextOnlyUids': ['fixture-0']})
        _, held = self.build(); context = held['identity_updates'][0]
        self.assertEqual(context['r30_native_optional_neutral_variants'], 1)
        self.assertEqual(context['r30_native_playable_optional_neutral_variants'], 0)
        self.assertTrue(context['r30_context_only'])
        entry['sourceActionCertified'] = True; self.install_native(entry)
        with self.assertRaisesRegex(ValueError, 'optional source-edition loop'):
            self.build()

    def test_changed_pixels_and_unknown_native_identity_are_rejected(self):
        target = self.root / self.entry['frames'][0]['path']; target.write_bytes(b'Changed source-derived bytes')
        with self.assertRaisesRegex(ValueError, 'hash or size mismatch'):
            self.build()
        self.entry['frames'] = [self.image('replacement.png')]; self.entry['subjectId'] = 'another-incarnation'
        self.install_native(self.entry)
        with self.assertRaisesRegex(ValueError, 'identity or scope'):
            self.build()

    def test_missing_and_duplicate_avatar_coverage_cannot_be_silently_filled(self):
        pack = EXPORT.read(EXPORT.AVATARS, self.root); pack['missing'].pop(); self.put(EXPORT.AVATARS, pack)
        with self.assertRaisesRegex(ValueError, 'All 389 identities'):
            self.build()
        pack['missing'].append({'id': 'fixture-388'}); pack['entries'].append(copy.deepcopy(pack['entries'][0])); self.put(EXPORT.AVATARS, pack)
        with self.assertRaisesRegex(ValueError, 'unique exact identity'):
            self.build()

    def test_all_49_historical_sheets_and_relations_stay_byte_exact(self):
        _, value = self.build(); baseline = (self.root / EXPORT.BASELINE).read_bytes()
        value['source_rows'].append(['G-PREUVES', 'fixture-0', 'Boundary fixture', ' \t=2+2', 'source', '@name'])
        result = EXPORT.export(value, self.root); self.assertEqual(result['worksheets'], 56)
        self.assertEqual((self.root / EXPORT.BASELINE).read_bytes(), baseline)
        with ZipFile(self.root / EXPORT.BASELINE) as before, ZipFile(self.root / result['workbook']) as after:
            for name in before.namelist():
                if name.startswith('xl/worksheets/'):
                    self.assertEqual(before.read(name), after.read(name), name)
        book = load_workbook(self.root / result['workbook'])
        self.assertEqual(book['History 1']['A1'].hyperlink.target, 'https://example.invalid/history')
        self.assertEqual(book['History 1']['B1'].value, '=2+2')
        self.assertEqual(book['Catalogue R30'].max_row, 390); self.assertEqual(book['Chantiers R30'].max_row, 18)
        last = book['Sources R30'].max_row
        self.assertTrue(book['Sources R30'].cell(last, 4).value.startswith("' \t="))
        self.assertTrue(book['Sources R30'].cell(last, 6).value.startswith("'@")); book.close()

    def test_independent_fresh_exports_have_identical_bytes(self):
        _, value = self.build()
        first = EXPORT.export(value, self.root, self.root / 'docs/first.xlsx', self.root / 'docs/first.json')
        second = EXPORT.export(value, self.root, self.root / 'docs/second.xlsx', self.root / 'docs/second.json')
        self.assertEqual(first['sha256'], second['sha256'])

    def test_export_rejects_invented_idle_closed_workstreams_and_stale_inputs(self):
        _, value = self.build(); invented = copy.deepcopy(value)
        invented['identity_updates'][0]['r30_native_verified_idle_cycles'] = 1
        with self.assertRaisesRegex(ValueError, 'coverage differs'):
            EXPORT.export(invented, self.root)
        bad = copy.deepcopy(value); bad['tasks'][0]['status'] = 'VALIDÉ'
        with self.assertRaisesRegex(ValueError, 'seventeen open workstreams'):
            EXPORT.export(bad, self.root)
        self.put(self.proof, {'passed': 1, 'failed': 0, 'changed': True})
        with self.assertRaisesRegex(ValueError, 'hash or size mismatch'):
            EXPORT.export(value, self.root)

    def test_stale_executed_report_is_excluded_from_current_totals(self):
        roster = EXPORT.read('data/roster.json', self.root); roster['historical'] = 'changed'; self.put('data/roster.json', roster)
        with self.assertRaisesRegex(ValueError, 'validation is unfinished'):
            self.build()
        checks = EXPORT.validation(self.root, allow_pending=True)
        self.assertFalse(checks['complete']); self.assertEqual(checks['totals']['fixture_checks'], 0)
        self.assertEqual(checks['failed_reports'], [self.proof])

    def test_output_cannot_replace_history_or_another_authors_workbook(self):
        _, value = self.build()
        with self.assertRaisesRegex(ValueError, 'historical inputs'):
            EXPORT.export(value, self.root, self.root / EXPORT.BASELINE)
        book = Workbook(); book.properties.creator = 'Other author'; book.save(self.root / 'docs/output.xlsx'); book.close()
        with self.assertRaisesRegex(ValueError, 'another author'):
            EXPORT.export(value, self.root, self.root / 'docs/output.xlsx')


if __name__ == '__main__':
    unittest.main()
