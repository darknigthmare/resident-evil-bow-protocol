"""R29 history preservation, source-frame scope and exact-UID edition boundaries."""
import copy
import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]


def module(name, relative):
    spec = importlib.util.spec_from_file_location(name, ROOT / relative)
    result = importlib.util.module_from_spec(spec); spec.loader.exec_module(result); return result


EXPORT = module('r29_progress_test', 'scripts/fidelity-progress-r29.py')
DOSSIER = module('r29_dossier_test', 'scripts/build-dossier-additions-r29.py')


class FidelityR29Tests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(); self.addCleanup(temporary.cleanup); self.root = Path(temporary.name)
        for name in ('docs', 'data', 'modules/fidelity-r28/sprites', 'modules/fidelity-r29/sprites'):
            (self.root / name).mkdir(parents=True)
        self.proof = self.root / 'docs/proof.json'; self.proof.write_text('{"fixture":true}\n')
        self.baseline = self.root / 'docs/BOW_FIDELITE_R28.xlsx'
        book = Workbook(); book.properties.creator = 'BOW fidelity progress R28'; book.active.title = 'Literal history'
        book.active.append(['Historical formula', '=2+2']); book.active['A1'].hyperlink = 'https://example.invalid/old'
        page = book.create_sheet('Catalogue R28'); page.append(['UID', 'Historical identity'])
        for i in range(389): page.append([str(i), 'Identity ' + str(i)])
        for i in range(40): book.create_sheet('Historical ' + str(i))
        book.save(self.baseline); book.close()
        tasks = [{'id': 'TASK-' + str(i), 'title': 'Open ' + str(i), 'category': 'sources', 'status': 'EN COURS',
            'advances': ['Observed source'], 'remaining': ['Unverified original cycle'],
            'evidence': [{'href': '../../docs/proof.json', 'label': 'Local proof'}]} for i in range(17)]
        prior = self.root / 'docs/R28_TRACKING.json'; prior.write_text(json.dumps({'tasks': tasks}))
        (self.root / 'data/roster.json').write_text(json.dumps({'catalog': [{'id': str(i)} for i in range(389)]}))
        updates = []
        for i in range(389):
            row = {key: '' for key in EXPORT.FIELDS}; row.update(id=str(i), sprite_available_states=[], sprite_combat_eligible=False,
                r29_native_clips=0, r29_native_frames=0, r29_native_combat_states=[], r29_documentary_variants=[], r29_documentary_source_editions=[])
            updates.append(row)
        self.tracking = {'schema': 're.bow.parallel-workstreams/3', 'release': 'R29', 'canonical_fidelity_certified': 0,
            'summary': 'Synthetic integrity fixture', 'statistics': [{'label': 'Open', 'value': 17, 'scope': 'No certification'}],
            'tasks': tasks, 'identity_updates': updates, 'resources': [{'href': '../../docs/proof.json', 'title': 'Proof', 'label': 'Read'}],
            'source_rows': [['TASK-0', '0', 'Quote', ' \t=HYPERLINK("https://bad.invalid")', 'Open', '@external']],
            'native_rows': [], 'body_rows': [], 'validation': [], 'inputs': []}
        for path in (self.baseline, prior, self.proof): self.bind(path)

    def bind(self, path):
        self.tracking['inputs'].append({'path': path.relative_to(self.root).as_posix(), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})

    def export(self, destination='docs/output.xlsx'):
        (self.root / 'docs/R29_TRACKING.json').write_text(json.dumps(self.tracking))
        return EXPORT.export(self.tracking, self.root, self.root / destination, self.root / 'modules/fidelity-r29/progress.json')

    def dossier_fixture(self):
        sprites = [{'id': str(i), 'sha256': 'a' * 64, 'fullBodyReviewed': False, 'completeBodyReferenceVerified': False,
            'scope': 'combatant-neutral-pose', 'sourceFacing': 'front', 'poseKind': 'source-presentation', 'availableStates': [],
            'knownAppearanceGaps': [], 'equipmentCompletenessReviewed': False, 'reviewStatus': 'visually-reviewed', 'canonicalFidelityCertified': False} for i in range(363)]
        previous = {'sprites': sprites, 'pending': [{'id': str(i), 'generationStatus': 'BLOCKED_BY_TOOL', 'reason': 'Preserved refusal'} for i in range(363, 385)]}
        current = copy.deepcopy(previous)
        (self.root / 'modules/fidelity-r28/sprites/manifest.json').write_text(json.dumps(previous))
        return current

    def write_dossier(self, current):
        (self.root / 'modules/fidelity-r29/sprites/manifest.json').write_text(json.dumps(current))
        return DOSSIER.build(self.root)

    def test_all_42_historical_sheets_and_relationships_stay_byte_exact(self):
        before = self.baseline.read_bytes(); result = self.export()
        self.assertEqual(self.baseline.read_bytes(), before); self.assertEqual(result['worksheets'], 49)
        with ZipFile(self.baseline) as old, ZipFile(self.root / result['workbook']) as new:
            for name in old.namelist():
                if name.startswith('xl/worksheets/'): self.assertEqual(old.read(name), new.read(name), name)
        book = load_workbook(self.root / result['workbook'])
        self.assertEqual(book['Literal history']['B1'].value, '=2+2'); self.assertEqual(book['Literal history']['A1'].hyperlink.target, 'https://example.invalid/old')
        self.assertEqual(book['Catalogue R29'].max_row, 390); self.assertEqual(book['Chantiers R29'].max_row, 18)
        self.assertTrue(book['Sources R29']['D2'].value.startswith("' \t=")); self.assertTrue(book['Sources R29']['F2'].value.startswith("'@")); book.close()

    def test_fresh_exports_have_identical_bytes(self):
        a = self.export('docs/first.xlsx'); b = self.export('docs/second.xlsx')
        self.assertEqual((self.root / a['workbook']).read_bytes(), (self.root / b['workbook']).read_bytes())

    def test_stale_evidence_and_closed_workstreams_cannot_be_exported(self):
        self.proof.write_text('{"fixture":false}\n')
        with self.assertRaisesRegex(ValueError, 'hash mismatch'): self.export()
        self.assertFalse((self.root / 'docs/output.xlsx').exists())
        self.proof.write_text('{"fixture":true}\n'); self.tracking['tasks'][0]['status'] = 'VALIDÉ'
        with self.assertRaisesRegex(ValueError, 'seventeen'): self.export()

    def test_missing_identity_and_unsafe_output_are_rejected(self):
        self.tracking['identity_updates'].pop()
        with self.assertRaisesRegex(ValueError, '389 identities'): self.export()
        self.setUp()
        with self.assertRaisesRegex(ValueError, 'historical inputs'): self.export('docs/BOW_FIDELITE_R28.xlsx')
        book = Workbook(); book.properties.creator = 'Another author'; book.save(self.root / 'docs/output.xlsx'); book.close()
        with self.assertRaisesRegex(ValueError, 'another author'): self.export()

    def test_dossier_preserves_refusals_and_rejects_known_appearance_gaps_with_idle(self):
        current = self.dossier_fixture(); current['pending'][0]['reason'] = 'Removed refusal reason'
        with self.assertRaisesRegex(ValueError, '22 tool refusals'): self.write_dossier(current)
        current = self.dossier_fixture(); row = current['sprites'][0]
        row.update(fullBodyReviewed=True, completeBodyReferenceVerified=True, sourceFacing='right', poseKind='neutral-base',
                   knownAppearanceGaps=['Cropped holder'], availableStates=['idle'])
        with self.assertRaisesRegex(ValueError, 'appearance gaps'): self.write_dossier(current)
        row['availableStates'] = []; result = self.write_dossier(current)
        self.assertEqual(result['identities'][0]['spriteChange']['poseKind'], 'neutral-base')
        self.assertEqual(result['identities'][0]['spriteChange']['availableStates'], [])

    def test_body_reference_sha_and_repository_path_base_are_verified(self):
        current = self.dossier_fixture(); folder = self.root / 'modules/fidelity-r29/sources/body-rasters'; folder.mkdir(parents=True)
        image = folder / 'fixture.png'; image.write_bytes(b'Synthetic source integrity fixture')
        view = {'angle': 0, 'path': 'fixture.png', 'sha256': hashlib.sha256(image.read_bytes()).hexdigest(), 'bytes': image.stat().st_size}
        entry = {'id': '0', 'modelId': 'source-0', 'sourceEdition': 'Fixture edition', 'sourceVariantScope': 'Base pose', 'equipmentScope': 'No equipment certification', 'views': [view]}
        manifest = folder / 'manifest.json'; manifest.write_text(json.dumps({'entries': [entry]}))
        result = self.write_dossier(current); self.assertEqual(result['identities'][0]['bodyModels'], 1); self.assertEqual(result['identities'][0]['bodyViews'], 1)
        view['sha256'] = '0' * 64; manifest.write_text(json.dumps({'entries': [entry]}))
        with self.assertRaisesRegex(ValueError, 'differs'): self.write_dossier(current)
        view.update(path='../../../docs/proof.json', pathBase='repository'); manifest.write_text(json.dumps({'entries': [entry]}))
        with self.assertRaisesRegex(ValueError, 'local source PNG'): self.write_dossier(current)

    def test_native_images_cannot_be_assigned_as_combat_cycles(self):
        current = self.dossier_fixture(); folder = self.root / 'modules/fidelity-r29/motion'; folder.mkdir()
        entry = {'id': 'source-clip-0', 'subjectId': '0', 'clipId': 'BA-opaque', 'frameCount': 1, 'frames': [{}],
            'sourceEdition': 'HD port; original platform not certified', 'actionIdentity': 'unassigned', 'availableCombatStates': [],
            'originalGameFrames': False, 'canonicalFidelityCertified': False}
        path = folder / 'manifest.json'; path.write_text(json.dumps({'schema': 're.bow.native-frame-library-r29/1', 'entries': [entry]}))
        result = self.write_dossier(current); self.assertEqual(result['summary']['nativeSourceFrames'], 1); self.assertEqual(result['summary']['assignedCombatCycles'], 0)
        entry['availableCombatStates'] = ['idle']; path.write_text(json.dumps({'schema': 're.bow.native-frame-library-r29/1', 'entries': [entry]}))
        with self.assertRaisesRegex(ValueError, 'assigned combat action'): self.write_dossier(current)

    def test_frozen_job_flag_stays_false_while_retained_manual_hashes_qualify_source(self):
        current = self.dossier_fixture(); row = current['sprites'][0]
        reference = self.root / 'docs/reference.png'; reference.write_bytes(b'Synthetic source integrity bytes')
        contact = self.root / 'docs/comparison.png'; contact.write_bytes(b'Synthetic visual comparison integrity bytes')
        image = self.root / 'modules/fidelity-r29/sprites/active-fixture.png'; image.write_bytes(b'Synthetic generated integrity bytes')
        image_sha = hashlib.sha256(image.read_bytes()).hexdigest(); reference_sha = hashlib.sha256(reference.read_bytes()).hexdigest()
        observation = {'schema': 're.bow.source-body-observation-r27/1', 'uid': '0', 'source_complete_body_observed': True,
            'generated_complete_body_observed': True, 'exact_identity_declared': True, 'image_sha256': image_sha,
            'observed_facing': 'left', 'canonical_fidelity_certified': False, 'original_game_pose_certified': False,
            'reference': {'path': 'docs/reference.png', 'sha256': reference_sha},
            'visual_evidence': {'path': 'docs/comparison.png', 'sha256': hashlib.sha256(contact.read_bytes()).hexdigest()}}
        proof = self.root / 'docs/observation.json'; proof.write_text(json.dumps(observation))
        row.update(sha256=image_sha, image=image.relative_to(self.root).as_posix(), fullBodyReviewed=True,
            sourceFacing='left', poseKind='neutral-base', availableStates=['idle'],
            provenance={'references': [{'path': 'docs/reference.png', 'sha256': reference_sha}]},
            sourceCompletenessEvidence={'path': 'docs/observation.json', 'sha256': hashlib.sha256(proof.read_bytes()).hexdigest()})
        result = self.write_dossier(current); pose = result['identities'][0]['spriteChange']
        self.assertFalse(pose['completeBodyReferenceVerified']); self.assertTrue(pose['sourceBodyObserved'])
        self.assertTrue(result['identities'][0]['activeSourceBodyObserved']); self.assertTrue(EXPORT.eligible(row, self.root))
        reference.write_bytes(b'Changed source bytes')
        with self.assertRaisesRegex(ValueError, 'source or visual comparison bytes changed'): self.write_dossier(current)

    def test_documentary_edition_variant_never_completes_or_replaces_target(self):
        current = self.dossier_fixture(); image = self.root / 'modules/fidelity-r29/sprites/fixture.png'; image.write_bytes(b'Synthetic edition variant integrity fixture')
        variant = {'id': '0', 'variantId': '0-hd-pc-2013', 'sourceEdition': 'HD/PC2013', 'targetIncarnation': 'Original3DS2012',
            'scope': 'documentary-edition-variant', 'admittedTargetEdition': False, 'targetFullBodyReviewed': False,
            'comparesAgainstSha256': 'a' * 64, 'availableStates': [], 'image': image.relative_to(self.root).as_posix(),
            'sha256': hashlib.sha256(image.read_bytes()).hexdigest(), 'bytes': image.stat().st_size}
        current['documentaryVariants'] = [variant]; result = self.write_dossier(current)
        self.assertEqual(result['summary']['documentaryEditionVariants'], 1); self.assertEqual(result['summary']['pixelRevisions'], 0)
        self.assertIsNone(result['identities'][0]['spriteChange'])
        for key, bad in [('admittedTargetEdition', True), ('targetFullBodyReviewed', True), ('availableStates', ['idle'])]:
            saved = variant[key]; variant[key] = bad
            with self.assertRaisesRegex(ValueError, 'cannot supersede'): self.write_dossier(current)
            variant[key] = saved
        current['sprites'][0]['sha256'] = 'b' * 64
        with self.assertRaisesRegex(ValueError, 'cannot supersede'): self.write_dossier(current)

    def test_legacy_edition_note_changes_information_without_pixels_or_pose_review(self):
        current = self.dossier_fixture(); row = current['sprites'][0]
        row.update(sourceEdition='HD/PC2013', sourceVariantLabel='Original3DS2012 equivalence unverified', knownAppearanceGaps=['Edition equivalence remains open'])
        baseline = self.root / 'modules/fidelity-r28/sprites/manifest.json'
        note = {'id': '0', 'imageSha256': row['sha256'], 'sourceEdition': row['sourceEdition'],
            'sourceVariantLabel': row['sourceVariantLabel'], 'editionEquivalenceVerified': False, 'knownAppearanceGaps': row['knownAppearanceGaps']}
        document = {'schema': 're.bow.legacy-sprite-source-notes-r29/1', 'manualPngEdits': False, 'newPoseReviews': 0, 'newTargetIdentities': 0,
            'baselineManifest': {'path': baseline.relative_to(self.root).as_posix(), 'sha256': hashlib.sha256(baseline.read_bytes()).hexdigest()}, 'entries': [note]}
        path = self.root / 'modules/fidelity-r29/sprites/legacy-source-notes.json'; path.write_text(json.dumps(document))
        row['r29LegacySourceNote'] = {'path': path.relative_to(self.root).as_posix(), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
        result = self.write_dossier(current); pose = result['identities'][0]['spriteChange']
        self.assertEqual(pose['sourceEdition'], 'HD/PC2013'); self.assertFalse(pose['pixelRevision']); self.assertEqual(pose['poseKind'], 'source-presentation')
        self.assertEqual(result['summary']['pixelRevisions'], 0); self.assertEqual(result['summary']['documentaryEditionVariants'], 0)
        note['imageSha256'] = 'b' * 64; path.write_text(json.dumps(document)); row['r29LegacySourceNote']['sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
        with self.assertRaisesRegex(ValueError, 'does not bind'): self.write_dossier(current)


if __name__ == '__main__':
    result = unittest.TextTestRunner(verbosity=2).run(unittest.defaultTestLoader.loadTestsFromTestCase(FidelityR29Tests))
    files = ['scripts/fidelity-progress-r29.py', 'scripts/build-dossier-additions-r29.py', 'tests/test_fidelity_r29.py']
    report = {'schema': 're.bow.export-tests-r29/1', 'passed': result.testsRun - len(result.failures) - len(result.errors),
        'failed': len(result.failures) + len(result.errors), 'testsRun': result.testsRun,
        'sourceFiles': [{'path': path, 'sha256': hashlib.sha256((ROOT / path).read_bytes()).hexdigest()} for path in files],
        'scope': 'Synthetic integrity fixtures; preserves 42 old sheets, exact UID/refusals, native action quarantine and documentary edition boundaries.'}
    (ROOT / 'docs/R29_EXPORT_TESTS.json').write_text(json.dumps(report, indent=2) + '\n')
    raise SystemExit(0 if result.wasSuccessful() else 1)
