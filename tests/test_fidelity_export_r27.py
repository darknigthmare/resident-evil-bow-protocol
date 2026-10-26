"""R27 export: preserve history, verify evidence, and keep native clocks distinct."""
import copy
import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('export_r27', ROOT / 'scripts/export-fidelity-progress-r27.py')
EXPORT = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(EXPORT)


class FidelityExportR27Tests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        for folder in ('docs', 'data', 'modules/fidelity-r27/motion'):
            (self.root / folder).mkdir(parents=True)
        self.proof = self.root / 'docs/proof.json'
        self.proof.write_text('{"observed": true}\n')
        self.baseline = self.root / 'docs/baseline.xlsx'
        book = Workbook()
        book.properties.creator = 'BOW fidelity progress R26'
        page = book.active
        page.title = 'Historical controls'
        page.append(['Literal historic formula', '=2+2'])
        page['A1'].font = Font(name='Arial', bold=True, color='BB3311')
        page['A1'].fill = PatternFill('solid', fgColor='ECE0D0')
        page['A2'] = 'Existing hyperlink'
        page['A2'].hyperlink = 'https://example.invalid/historical'
        page.freeze_panes = 'B2'
        for title in ('Catalogue 389', 'Catalogue R26'):
            page = book.create_sheet(title)
            page.append(['UID', 'Name', 'R26 evidence'])
            for i in range(389):
                page.append([str(i), 'Historical identity ' + str(i), 'Preserved R26'])
        book.save(self.baseline)
        book.close()
        tasks = [{'id': 'TASK-' + str(i), 'title': 'Open workstream ' + str(i), 'status': 'EN COURS',
                  'category': 'sources', 'advances': ['Reviewed evidence'], 'remaining': ['Production still open'],
                  'evidence': [{'href': '../../docs/proof.json', 'label': 'Proof'}]} for i in range(17)]
        (self.root / 'docs/R20_TRACKING.json').write_text(json.dumps({'tasks': tasks}))
        (self.root / 'data/roster.json').write_text(json.dumps({'catalog': [{'id': str(i)} for i in range(389)]}))
        self.tracking = {
            'schema': 're.bow.parallel-workstreams/2', 'canonical_fidelity_certified': 0,
            'summary': 'All production criteria remain open.', 'statistics': [{'label': 'Open workstreams', 'value': 17}],
            'resources': [{'href': '../../docs/proof.json', 'title': 'Proof', 'label': 'Read evidence'}],
            'tasks': tasks,
            'identity_updates': [{'id': str(i), 'visual': '', 'rig': '', 'motion_clips': 0, 'audio': 0,
                                  'claims': 0, 'limitations': ['Not certified']} for i in range(389)],
            'source_rows': [['TASK-0', '0', 'Quote', ' \t=HYPERLINK("https://bad.invalid")', 'Unverified', '@external-command']],
            'validation': [{'label': 'Synthetic export check', 'result': 'Passed', 'scope': 'Test fixture only', 'path': 'docs/proof.json'}],
            'inputs': [],
        }
        self.record_hash(self.proof)
        self.record_hash(self.baseline)

    def record_hash(self, path):
        relative = path.relative_to(self.root).as_posix()
        self.tracking['inputs'] = [item for item in self.tracking['inputs'] if item['path'] != relative]
        self.tracking['inputs'].append({'path': relative, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})

    def export(self, output=None):
        tracking_path = self.root / 'docs/tracking.json'
        tracking_path.write_text(json.dumps(self.tracking))
        return EXPORT.export(tracking_path, self.baseline, output or self.root / 'docs/output.xlsx',
                             self.root / 'modules/fidelity-r27/progress.json', self.root)

    def add_motion(self):
        clip = {'id': 'native-clip', 'member': 'Animations/native.FBX', 'sha256': 'a' * 64,
                'keyTimeStart': '-1', 'keyTimeEnd': '99', 'ticksPerSecond': 100,
                'sourceTimeSpanSeconds': 1.0, 'nativePreRollSeconds': 0.01, 'lastNativeTimeSeconds': 0.99,
                'varyingAnimationCurveCount': 3, 'varyingNonnegativeAnimationCurveCount': 2,
                'positiveTimePoseOnly': False, 'playbackVariationScope': 'nonnegative-time-motion',
                'semanticActionVerified': False}
        manifest = {'schema': 're.bow.rigged-motion-pack/1', 'summary': {'sourceAnimationFiles': 1,
                    'positiveTimeMotionFiles': 1, 'positiveTimePoseFiles': 0, 'canonicalFidelityCertified': 0},
                    'entries': [{'id': '0', 'game': 'Synthetic metadata fixture', 'rigProductionFinalized': False,
                                 'canonicalFidelityCertified': False, 'clips': [clip]}]}
        path = self.root / 'modules/fidelity-r27/motion/manifest.json'
        path.write_text(json.dumps(manifest))
        self.tracking['motion_manifest'] = path.relative_to(self.root).as_posix()
        self.record_hash(path)
        return path, manifest

    def test_export_preserves_all_old_sheet_bytes_and_quotes_new_formula_text(self):
        before = self.baseline.read_bytes()
        result = self.export()
        self.assertEqual(self.baseline.read_bytes(), before)
        with ZipFile(self.baseline) as original, ZipFile(result['workbook']) as output:
            for name in original.namelist():
                if name.startswith('xl/worksheets/'):
                    self.assertEqual(original.read(name), output.read(name), name)
        book = load_workbook(result['workbook'])
        self.assertEqual(book['Historical controls']['B1'].value, '=2+2')
        self.assertEqual(book['Historical controls']['A1'].font.color.rgb, '00BB3311')
        self.assertEqual(book['Catalogue R26']['C2'].value, 'Preserved R26')
        self.assertEqual(book['Catalogue R27']['C2'].value, 'Preserved R26')
        self.assertEqual(book['Catalogue R27'].max_row, 390)
        self.assertIn('Clips audio documentaires R27', [cell.value for cell in book['Catalogue R27'][1]])
        self.assertEqual(book['Chantiers R27'].max_row, 18)
        self.assertEqual(book['Sources R27']['D2'].data_type, 's')
        self.assertTrue(book['Sources R27']['D2'].value.startswith("' \t="))
        self.assertTrue(book['Sources R27']['F2'].value.startswith("'@"))
        self.assertEqual(book['Catalogue R27']['I2'].value, '["Not certified"]')
        book.close()
        progress = json.loads(Path(result['progress']).read_text())
        self.assertEqual(progress['schema'], 're.bow.fidelity-progress/1')
        self.assertEqual(len(progress['tasks']), 17)

    def test_changed_evidence_is_rejected_before_any_delivery(self):
        self.proof.write_text('{"observed": false}\n')
        with self.assertRaisesRegex(ValueError, 'hash mismatch'):
            self.export()
        self.assertFalse((self.root / 'docs/output.xlsx').exists())
        self.assertFalse((self.root / 'modules/fidelity-r27/progress.json').exists())

    def test_missing_or_duplicate_workstream_and_false_closure_are_rejected(self):
        original = copy.deepcopy(self.tracking)
        self.tracking['tasks'].pop()
        with self.assertRaisesRegex(ValueError, 'workstreams'):
            self.export()
        self.tracking = copy.deepcopy(original)
        self.tracking['tasks'][-1] = copy.deepcopy(self.tracking['tasks'][0])
        with self.assertRaisesRegex(ValueError, 'workstreams'):
            self.export()
        self.tracking = original
        self.tracking['tasks'][0]['status'] = 'VALIDÉ'
        with self.assertRaisesRegex(ValueError, 'closed'):
            self.export()

    def test_unknown_missing_or_duplicate_identity_cannot_gain_progress(self):
        original = copy.deepcopy(self.tracking)
        for mutation in ('unknown', 'missing', 'duplicate'):
            self.tracking = copy.deepcopy(original)
            if mutation == 'unknown':
                self.tracking['identity_updates'][0]['id'] = 'unknown-incarnation'
            elif mutation == 'missing':
                self.tracking['identity_updates'].pop()
            else:
                self.tracking['identity_updates'][-1] = copy.deepcopy(self.tracking['identity_updates'][0])
            with self.assertRaisesRegex(ValueError, 'identity|identities'):
                self.export()

    def test_global_certification_and_unrecognized_identity_fields_are_rejected(self):
        self.tracking['canonical_fidelity_certified'] = 1
        with self.assertRaisesRegex(ValueError, 'certification'):
            self.export()
        self.tracking['canonical_fidelity_certified'] = 0
        self.tracking['identity_updates'][0]['rig_final_certified'] = True
        with self.assertRaisesRegex(ValueError, 'Unknown identity progress field'):
            self.export()

    def test_evidence_links_and_output_paths_cannot_escape_checkout_or_overwrite_history(self):
        self.tracking['tasks'][0]['evidence'][0]['href'] = '../../../../outside.json'
        with self.assertRaisesRegex(ValueError, 'inside the checkout'):
            self.export()
        self.tracking['tasks'][0]['evidence'][0]['href'] = '../../docs/proof.json'
        (self.root / 'tests').mkdir()
        (self.root / 'tests/proof.mjs').write_text('// Local code is not a review resource')
        self.tracking['tasks'][0]['evidence'][0]['href'] = '../../tests/proof.mjs'
        with self.assertRaisesRegex(ValueError, 'modules or docs'):
            self.export()
        self.tracking['tasks'][0]['evidence'][0]['href'] = '../../docs/proof.json'
        with self.assertRaisesRegex(ValueError, 'inside the checkout'):
            self.export(self.root.parent / 'escaped.xlsx')
        with self.assertRaisesRegex(ValueError, 'overwrite historical'):
            self.export(self.baseline)

    def test_baseline_requires_the_expected_author_and_a_recorded_hash(self):
        self.tracking['inputs'] = [item for item in self.tracking['inputs'] if item['path'] != 'docs/baseline.xlsx']
        with self.assertRaisesRegex(ValueError, 'recorded evidence hash'):
            self.export()
        book = load_workbook(self.baseline)
        book.properties.creator = 'Unrelated workbook'
        book.save(self.baseline)
        book.close()
        self.record_hash(self.baseline)
        with self.assertRaisesRegex(ValueError, 'historical R26'):
            self.export()

    def test_native_min_to_max_span_preroll_and_zero_to_max_playback_remain_distinct(self):
        self.add_motion()
        result = self.export()
        self.assertEqual(result['source_clips'], 1)
        book = load_workbook(result['workbook'])
        page = book['Clips sources R27']
        self.assertEqual(page['G2'].value, -0.01)
        self.assertEqual(page['H2'].value, 0.99)
        self.assertEqual(page['I2'].value, 1)
        self.assertEqual(page['J2'].value, 0.01)
        self.assertEqual(page['K2'].value, 0.99)
        self.assertEqual(page['L2'].value, 3)
        self.assertEqual(page['M2'].value, 2)
        self.assertEqual(page['N2'].value, 'MOUVEMENT APRÈS ZÉRO')
        self.assertEqual(page['P2'].value, 'ACTION NON ATTRIBUÉE')
        book.close()

    def test_associated_object_remains_a_distinct_resource_under_its_documented_subject(self):
        path, manifest = self.add_motion()
        entry = manifest['entries'][0]
        entry.update(id='source-trap-reference', subjectId='0', modelRole='associated-object')
        path.write_text(json.dumps(manifest))
        self.record_hash(path)
        result = self.export()
        book = load_workbook(result['workbook'])
        page = book['Clips sources R27']
        self.assertEqual(page['A2'].value, '0')
        self.assertEqual(page['R2'].value, 'source-trap-reference')
        self.assertEqual(page['S2'].value, 'associated-object')
        book.close()
        for changes in ({'subjectId': 'unknown'}, {'modelRole': 'body'}, {'modelRole': 'invented-role'}):
            original = copy.deepcopy(entry)
            entry.update(changes)
            path.write_text(json.dumps(manifest))
            self.record_hash(path)
            with self.assertRaisesRegex(ValueError, 'Unknown motion identity'):
                self.export()
            entry.clear()
            entry.update(original)

    def test_preroll_variation_does_not_turn_a_positive_time_pose_into_a_movement(self):
        path, manifest = self.add_motion()
        pose = copy.deepcopy(manifest['entries'][0]['clips'][0])
        pose.update(id='pose-after-zero', varyingNonnegativeAnimationCurveCount=0,
                    positiveTimePoseOnly=True, playbackVariationScope='pose-only-after-time-zero')
        manifest['entries'][0]['clips'].append(pose)
        manifest['summary'].update(sourceAnimationFiles=2, positiveTimeMotionFiles=1, positiveTimePoseFiles=1)
        path.write_text(json.dumps(manifest))
        self.record_hash(path)
        result = self.export()
        self.assertEqual(result['source_clips'], 2)
        self.assertEqual(result['positive_time_motion_files'], 1)
        self.assertEqual(result['positive_time_pose_files'], 1)
        book = load_workbook(result['workbook'])
        page = book['Clips sources R27']
        self.assertEqual(page['L3'].value, 3)
        self.assertEqual(page['M3'].value, 0)
        self.assertEqual(page['N3'].value, 'POSE FIXE APRÈS ZÉRO')
        self.assertEqual(page['O3'].value, 'pose-only-after-time-zero')
        self.assertEqual(page['J3'].value, 0.01)
        self.assertEqual(page['K3'].value, 0.99)
        self.assertIn('même si le pré-roll varie', page['Q3'].value)
        book.close()

    def test_contradictory_pose_flags_scope_curve_counts_and_summary_are_rejected(self):
        path, original = self.add_motion()
        for mutation in ('pose_flag', 'scope', 'too_many_curves', 'summary'):
            with self.subTest(mutation=mutation):
                manifest = copy.deepcopy(original)
                clip = manifest['entries'][0]['clips'][0]
                if mutation == 'pose_flag':
                    clip['positiveTimePoseOnly'] = True
                elif mutation == 'scope':
                    clip['playbackVariationScope'] = 'pose-only-after-time-zero'
                elif mutation == 'too_many_curves':
                    clip['varyingNonnegativeAnimationCurveCount'] = clip['varyingAnimationCurveCount'] + 1
                else:
                    manifest['summary'].update(positiveTimeMotionFiles=0, positiveTimePoseFiles=1)
                path.write_text(json.dumps(manifest))
                self.record_hash(path)
                with self.assertRaisesRegex(ValueError, 'classification|scope|totals'):
                    self.export()

    def test_inconsistent_native_clock_is_rejected_even_with_a_valid_manifest_hash(self):
        path, manifest = self.add_motion()
        manifest['entries'][0]['clips'][0]['sourceTimeSpanSeconds'] = 0.99
        path.write_text(json.dumps(manifest))
        self.record_hash(path)
        with self.assertRaisesRegex(ValueError, 'Native animation times disagree'):
            self.export()

    def add_sprite_pack(self, scope='combatant-neutral-pose', reviewed=True, pose_kind='neutral-base'):
        folder = self.root / 'assets/sprites/bow'
        folder.mkdir(parents=True)
        (folder / 'manifest.json').write_text(json.dumps({'sprites': [{'id': str(i)} for i in range(4)]}))
        path = self.root / 'modules/fidelity-r27/sprites/generated/4/idle-right.png'
        path.parent.mkdir(parents=True)
        path.write_bytes(b'PNG bytes used only by the export integrity fixture')
        sprite = {'id': '4', 'image': path.relative_to(self.root).as_posix(), 'bytes': path.stat().st_size,
                  'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'reviewStatus': 'visually-reviewed' if reviewed else 'unreviewed',
                  'fullBodyReviewed': reviewed, 'completeBodyReferenceVerified': reviewed, 'scope': scope,
                  'sourceFacing': 'right', 'poseKind': pose_kind, 'canMirror': False,
                  'availableStates': ['idle'] if reviewed and scope == 'combatant-neutral-pose' and pose_kind == 'neutral-base' else [], 'combatIntegrated': False,
                  'canonicalFidelityCertified': False}
        pack = {'schema': 're.bow.generated-combat-sprites/1', 'sprites': [sprite],
                'pending': [{'id': str(i)} for i in range(5, 389)],
                'counts': {'targetIdentities': 385, 'generatedSprites': 1, 'pendingSprites': 384,
                           'visuallyReviewedSprites': int(reviewed), 'fullBodyReviewedSprites': int(reviewed),
                           'neutralBaseSprites': int(pose_kind == 'neutral-base'), 'sourceBindPoseSprites': int(pose_kind == 'source-bind-pose'),
                           'sourcePresentationSprites': int(pose_kind == 'source-presentation'), 'unknownPoseSprites': int(pose_kind == 'unknown'),
                           'neutralBaseCombatEligibleSprites': int(bool(sprite['availableStates'])),
                           'originalGameSprites': 0, 'completeCombatAnimationSets': 0, 'canonicalFidelityCertified': 0}}
        manifest = self.root / 'modules/fidelity-r27/sprites/manifest.json'
        manifest.write_text(json.dumps(pack))
        self.tracking['sprite_manifest'] = manifest.relative_to(self.root).as_posix()
        self.record_hash(manifest)
        jobs = self.root / 'modules/fidelity-r27/sprites/jobs.json'
        jobs.write_text(json.dumps({'schema': 're.bow.sprite-generation-jobs/1',
                                    'jobs': [{'id': str(i), 'scope': scope if i == 4 else 'combatant-neutral-pose',
                                              'outputPath': f'modules/fidelity-r27/sprites/generated/{i}/idle-right.png'} for i in range(4, 389)]}))
        self.record_hash(jobs)
        for update in self.tracking['identity_updates']:
            update.update(sprite_generated=False, sprite_review_status='pending', sprite_full_body_reviewed=False,
                          sprite_reference_full_body_verified=False, sprite_scope='combatant-neutral-pose', sprite_facing='',
                          sprite_pose_kind='neutral-base' if int(update['id']) < 4 else 'unknown',
                          sprite_available_states=[], sprite_combat_eligible=False, sprite_combat_integrated=False,
                          sprite_generation_status='HISTORICAL_R20' if int(update['id']) < 4 else 'AWAITING_OUTPUT', sprite_generation_reason='')
        self.tracking['identity_updates'][4].update(sprite_generated=True, sprite_review_status=sprite['reviewStatus'],
                                                  sprite_full_body_reviewed=reviewed, sprite_reference_full_body_verified=reviewed,
                                                  sprite_scope=scope, sprite_facing='right', sprite_pose_kind=pose_kind, sprite_available_states=sprite['availableStates'],
                                                  sprite_combat_eligible=bool(sprite['availableStates']), sprite_generation_status='GENERATED')
        return manifest, pack, path

    def test_generated_pending_and_reviewed_sprites_remain_distinct_in_new_columns(self):
        self.add_sprite_pack()
        result = self.export()
        book = load_workbook(result['workbook'])
        page = book['Catalogue R27']; headers = [cell.value for cell in page[1]]
        made = headers.index('Nouveau sprite généré') + 1
        eligible = headers.index('Admissible au rendu idle de combat') + 1
        integrated = headers.index('Intégration de combat certifiée') + 1
        self.assertIs(page.cell(6, made).value, True)
        self.assertIs(page.cell(7, made).value, False)
        self.assertIs(page.cell(6, eligible).value, True)
        self.assertIs(page.cell(6, integrated).value, False)
        self.assertEqual(len([name for name in book.sheetnames if name.endswith('R27')]), 6)
        book.close()

    def test_queued_assets_and_falsified_generated_counts_or_bytes_are_rejected(self):
        manifest, original, path = self.add_sprite_pack()
        for mutation in ('missing_pending', 'duplicate', 'false_count', 'false_completion'):
            changed = copy.deepcopy(original)
            if mutation == 'missing_pending':
                changed['pending'].pop()
            elif mutation == 'duplicate':
                changed['pending'][0]['id'] = '4'
            elif mutation == 'false_count':
                changed['counts']['generatedSprites'] = 385
            else:
                changed['counts']['completeCombatAnimationSets'] = 1
            manifest.write_text(json.dumps(changed)); self.record_hash(manifest)
            with self.subTest(mutation=mutation), self.assertRaisesRegex(ValueError, '385|counts'):
                self.export()
        manifest.write_text(json.dumps(original)); self.record_hash(manifest)
        path.write_bytes(b'changed generated sprite')
        with self.assertRaisesRegex(ValueError, 'sprite bytes'):
            self.export()

    def test_contexts_or_unreviewed_poses_cannot_be_promoted_into_combat_scope(self):
        manifest, pack, _ = self.add_sprite_pack(scope='context-encounter', reviewed=True)
        self.export()
        self.tracking['identity_updates'][4]['sprite_combat_eligible'] = True
        with self.assertRaisesRegex(ValueError, 'identity progress'):
            self.export()
        # Even a rewritten manifest cannot relabel a known context as a combatant.
        pack['sprites'][0]['scope'] = 'combatant-neutral-pose'
        self.tracking['identity_updates'][4]['sprite_scope'] = 'combatant-neutral-pose'
        manifest.write_text(json.dumps(pack)); self.record_hash(manifest)
        with self.assertRaisesRegex(ValueError, 'frozen generation job'):
            self.export()

    def test_tool_refusals_stay_visible_and_cannot_be_flattened_to_generic_pending(self):
        manifest, pack, _ = self.add_sprite_pack()
        pack['pending'][0].update(generationStatus='BLOCKED_BY_TOOL', reason='The image tool returned its explicit test refusal.')
        manifest.write_text(json.dumps(pack)); self.record_hash(manifest)
        update = self.tracking['identity_updates'][5]
        with self.assertRaisesRegex(ValueError, 'tool refusal reasons'):
            self.export()
        update.update(sprite_generation_status='BLOCKED_BY_TOOL', sprite_generation_reason=pack['pending'][0]['reason'])
        result = self.export()
        book = load_workbook(result['workbook']); page = book['Catalogue R27']
        headers = [cell.value for cell in page[1]]
        self.assertEqual(page.cell(7, headers.index('État de la génération') + 1).value, 'BLOCKED_BY_TOOL')
        self.assertEqual(page.cell(7, headers.index('Motif d’attente ou de refus') + 1).value, pack['pending'][0]['reason'])
        book.close()

    def test_reviewed_whole_body_bind_pose_cannot_become_combat_idle(self):
        manifest, pack, _ = self.add_sprite_pack(pose_kind='source-bind-pose')
        result = self.export()
        book = load_workbook(result['workbook']); page = book['Catalogue R27']
        headers = [cell.value for cell in page[1]]
        self.assertIs(page.cell(6, headers.index('Corps entier revu') + 1).value, True)
        self.assertEqual(page.cell(6, headers.index('Type de pose observé') + 1).value, 'source-bind-pose')
        self.assertIs(page.cell(6, headers.index('Admissible au rendu idle de combat') + 1).value, False)
        book.close()
        self.tracking['identity_updates'][4]['sprite_combat_eligible'] = True
        with self.assertRaisesRegex(ValueError, 'identity progress'):
            self.export()
        self.tracking['identity_updates'][4]['sprite_combat_eligible'] = False
        pack['sprites'][0]['availableStates'] = ['idle']
        self.tracking['identity_updates'][4]['sprite_available_states'] = ['idle']
        manifest.write_text(json.dumps(pack)); self.record_hash(manifest)
        with self.assertRaisesRegex(ValueError, 'neutral poses'):
            self.export()


if __name__ == '__main__':
    unittest.main()
